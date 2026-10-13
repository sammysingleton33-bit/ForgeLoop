"""Tests for tools/content/content.py (game-001 content library). Run: python -m pytest -q tools/tests

Most tests run on the mini library in tools/tests/fixtures/content/valid, so they do not depend on the
starter content. The tests named test_real_* always run against games/game-001/content: it is on main, so
a missing or renamed CSV must fail them (never skip them; review finding L4).
"""
import csv
import importlib.util
import io
import json
import os
import re
import shutil
import subprocess
import sys
from pathlib import Path

import pytest

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
SPEC = importlib.util.spec_from_file_location("content_tool", ROOT / "tools" / "content" / "content.py")
content = importlib.util.module_from_spec(SPEC)
sys.modules["content_tool"] = content
SPEC.loader.exec_module(content)

FIXTURE = HERE / "fixtures" / "content" / "valid"
BROKEN_CASES = json.loads((HERE / "fixtures" / "content" / "broken_cases.json").read_text(encoding="utf-8"))
REAL = ROOT / "games" / "game-001" / "content"
REAL_OUT = ROOT / "games" / "game-001" / "src" / "shared" / "Content"
# Workbook tests need openpyxl. CI must install it (pip install pytest openpyxl) or these are skipped.
needs_openpyxl = pytest.mark.skipif(importlib.util.find_spec("openpyxl") is None,
                                    reason="openpyxl not installed (python3 -m pip install openpyxl)")
CSV_PATHS = [s.path for s in content.SPECS]


def run(capsys, *argv):
    code = content.main([str(a) for a in argv])
    return code, capsys.readouterr().out


def read_rows(path: Path) -> list[list[str]]:
    return list(csv.reader(io.StringIO(path.read_text(encoding="utf-8"), newline="")))


def write_rows(path: Path, rows: list[list[str]]) -> None:
    buf = io.StringIO()
    csv.writer(buf, lineterminator="\n").writerows(rows)
    path.write_text(buf.getvalue(), encoding="utf-8", newline="")


def set_cell(lib: Path, rel: str, row: int, column: str, value: str) -> None:
    """row is the spreadsheet row (header = 1)."""
    path = lib / rel
    rows = read_rows(path)
    rows[row - 1][rows[0].index(column)] = value
    write_rows(path, rows)


def snapshot(directory: Path) -> dict[str, bytes]:
    return {p.relative_to(directory).as_posix(): p.read_bytes() for p in sorted(directory.rglob("*")) if p.is_file()}


def csv_bytes(lib: Path) -> dict[str, bytes]:
    return {rel: (lib / rel).read_bytes() for rel in CSV_PATHS}


@pytest.fixture
def lib(tmp_path):
    dst = tmp_path / "content"
    shutil.copytree(FIXTURE, dst)
    return dst


def find_tool(name: str) -> str | None:
    extra = [str(Path.home() / ".rokit" / "bin"), str(Path.home() / ".cargo" / "bin"), "/root/.cargo/bin"]
    return shutil.which(name, path=os.pathsep.join([os.environ.get("PATH", ""), *extra]))


# ----------------------------------------------------------------------------- validate
def test_fixture_library_validates(capsys):
    code, out = run(capsys, "validate", "--content", FIXTURE)
    assert code == 0, out
    assert "3 scenarios, 7 responses, 7 decisions, 6 consequences, 6 traits, 101 names, 4 occupations, " \
           "6 categories, 4 actions — 0 errors" in out


def test_real_content_validates(capsys):
    code, out = run(capsys, "validate")
    assert code == 0, out
    assert out.strip().endswith("— 0 errors")


def test_real_generated_luau_is_up_to_date(capsys):
    code, out = run(capsys, "check")
    assert code == 0, out + f"\n(run: python3 tools/content/content.py build)"


def test_real_content_gates_can_never_be_skipped():
    """Review finding L4: a skipif on the real library would silently drop the staleness gate if a CSV were
    renamed. The real-content tests carry no skip marker, and a renamed CSV fails `check`."""
    real_tests = [v for k, v in globals().items() if k.startswith("test_real_") and callable(v)]
    assert test_real_generated_luau_is_up_to_date in real_tests and len(real_tests) >= 5
    for fn in real_tests:
        marks = [m.name for m in getattr(fn, "pytestmark", [])]
        assert "skipif" not in marks and "skip" not in marks, f"{fn.__name__} must always run"


def test_check_fails_when_a_csv_is_renamed(lib, tmp_path, capsys):
    out = tmp_path / "out"
    assert run(capsys, "build", "--content", lib, "--out", out)[0] == 0
    (lib / "reputation" / "actions.csv").rename(lib / "reputation" / "actions_old.csv")
    code, text = run(capsys, "check", "--content", lib, "--out", out)
    assert code == 1 and "reputation/actions.csv: missing file" in text


@pytest.mark.parametrize("case", BROKEN_CASES, ids=[c["name"] for c in BROKEN_CASES])
def test_broken_variant_fails_with_location(lib, tmp_path, capsys, case):
    set_cell(lib, case["file"], case["row"], case["column"], case["value"])
    for p in case.get("patch", []):
        set_cell(lib, p["file"], p["row"], p["column"], p["value"])
    code, out = run(capsys, "validate", "--content", lib)
    assert code == 1, out
    for expected in case["expect"]:
        assert expected in out, f"missing {expected!r} in:\n{out}"
    # build refuses to write anything
    code, out = run(capsys, "build", "--content", lib, "--out", tmp_path / "out")
    assert code == 1
    assert not (tmp_path / "out").exists()


def test_missing_file_and_csv_syntax_errors(lib, capsys):
    (lib / "citizens" / "names.csv").unlink()
    (lib / "dialogue" / "responses.csv").write_text('scenario_id,response_id\n"unclosed,x\n', encoding="utf-8")
    code, out = run(capsys, "validate", "--content", lib)
    assert code == 1
    assert "citizens/names.csv: missing file" in out
    assert "dialogue/responses.csv:2: a quoted cell is never closed" in out


def test_non_utf8_file_names_the_row_and_byte(lib, capsys):
    """Review finding L10: the error gives the byte offset and the spreadsheet row of the bad byte."""
    path = lib / "dialogue" / "scenarios.csv"
    rows = read_rows(path)
    rows[1][-1] = "line one\nline two, with a comma"  # a quoted multi-line cell: row numbers count records
    write_rows(path, rows)
    data = path.read_bytes().replace(b"My hammer is cracked", b"My hammer is caf\xe9 cracked")
    path.write_bytes(data)
    offset = data.index(b"\xe9")
    code, out = run(capsys, "validate", "--content", lib)
    assert code == 1
    assert (f"dialogue/scenarios.csv:3: not UTF-8 text: byte 0xE9 at byte offset {offset} (column 8) is not "
            "valid UTF-8; save the CSV as 'CSV UTF-8'") in out


def test_non_utf8_position_treats_a_quote_inside_an_unquoted_cell_as_text(lib, capsys):
    """Review finding L10: a literal '"' in an unquoted cell (it validates) shifted every later row and column."""
    path = lib / "dialogue" / "scenarios.csv"
    data = path.read_bytes()
    lines = data.split(b"\n")
    assert lines[2].endswith(b",")  # row 3 has an empty, unquoted notes cell
    lines[2] += b'a 5" plank'
    path.write_bytes(b"\n".join(lines))
    code, out = run(capsys, "validate", "--content", lib)
    assert code == 0, out  # the stray quote is a literal character
    rows = read_rows(path)
    assert rows[2][-1] == 'a 5" plank'
    data = path.read_bytes().replace(b"We bled at the wall", b"We bled at the caf\xe9 wall")
    path.write_bytes(data)
    offset = data.index(b"\xe9")
    code, out = run(capsys, "validate", "--content", lib)
    assert code == 1
    assert f"dialogue/scenarios.csv:4: not UTF-8 text: byte 0xE9 at byte offset {offset} (column 8)" in out
    # quoting is followed exactly like the parser: an escaped "" stays inside a quoted cell
    assert content._csv_position('a,"x "" y\nz",b') == (1, 3)  # the line break is inside the quoted cell
    assert content._csv_position('a,"x" ,b\nc') == (2, 1)
    assert content._csv_position('a,b"c,d\ne,') == (2, 2)
    assert content._csv_position('\ufeff"a\nb",c') == (1, 2)


@pytest.mark.parametrize("text, ok", [
    ("3", True), ("+3", True), ("-0.5", True), (".5", True), ("5.", True), ("1e3", True), ("3.0", True),
    ("1_0", False), ("\u0661\u0662", False), ("\uff11\uff12", False), ("\u0663.\u0665", False), ("nan", False),
    ("inf", False), ("0x10", False), ("1,5", False), ("--1", False), ("", False), ("\u00b2", False),
])
def test_numbers_are_ascii_decimal_only(text, ok):
    """Review finding L3: float() alone took '1_0', Arabic-Indic and full-width digits."""
    if ok:
        assert content._num(text) == float(text)
    else:
        with pytest.raises(ValueError):
            content._num(text)


def test_names_minimum_counts(lib, capsys):
    rows = read_rows(lib / "citizens" / "names.csv")
    write_rows(lib / "citizens" / "names.csv", [r for r in rows if r[0] != "given_n"])
    code, out = run(capsys, "validate", "--content", lib)
    assert code == 1
    assert "only 41 given names; at least 60 are needed" in out


def test_action_rows_share_anti_farming_settings(lib, capsys):
    set_cell(lib, "reputation/actions.csv", 3, "repeat_decay", "0.9")
    code, out = run(capsys, "validate", "--content", lib)
    assert code == 1
    assert "reputation/actions.csv:3:6: [repeat_decay] differs from row 2 of action 'fund_barracks'" in out


# ----------------------------------------------------------------------------- condition language
def vocab():
    return content.Vocab({"farmer", "soldier"}, {"greedy", "kind"}, {"d_one"})


@pytest.mark.parametrize("text, expected", [
    ("", []),
    ("kingdom.fed_ratio < 0.6", [{"var": "kingdom.fed_ratio", "op": "<", "value": 0.6}]),
    ("kingdom.count.farm>=2", [{"var": "kingdom.count.farm", "op": ">=", "value": 2}]),
    ("citizen.trait.greedy & !citizen.has_rival", [
        {"var": "citizen.trait.greedy", "op": "=", "value": True},
        {"var": "citizen.has_rival", "op": "=", "value": False}]),
    ("decision.d_one != true", [{"var": "decision.d_one", "op": "=", "value": False}]),
    ("phase in Founding, Expansion", [{"var": "phase", "op": "in", "value": ["Founding", "Expansion"]}]),
    ("occupation = farmer & rep.fear >= 40", [
        {"var": "occupation", "op": "=", "value": "farmer"}, {"var": "rep.fear", "op": ">=", "value": 40}]),
    ("citizen.memory.rulerVisited & citizen.chemistry > -5", [
        {"var": "citizen.memory.rulerVisited", "op": "=", "value": True},
        {"var": "citizen.chemistry", "op": ">", "value": -5}]),
    ("citizen.has_friend", [{"var": "citizen.has_friend", "op": "=", "value": True}]),
])
def test_condition_parses_into_typed_clauses(text, expected):
    clauses, _ = content.parse_condition(text, vocab(), "citizen")
    assert clauses == expected
    for c in clauses:
        assert type(c["value"]) is type(next(e for e in expected if e["var"] == c["var"])["value"])


@pytest.mark.parametrize("text, message", [
    ("kingdom.food <", "malformed"),
    ("kingdom.food < lots", "'lots' is not a number"),
    ("kingdom.food", "is a number; compare it"),
    ("phase = founding", "unknown phase 'founding' (did you mean 'Founding'?)"),
    ("phase > Founding", "use = != or in"),
    ("rep.fear in 1,2", "'in' only works with"),
    ("citizen.trait.greedy = maybe", "is true/false"),
    ("rep.mercy > 5", "unknown variable 'rep.mercy'"),
    ("event.flood", "unknown variable 'event.flood'"),
    ("kingdom.count.castle > 1", "unknown variable 'kingdom.count.castle'"),
    ("citizen.chemistry > 25", "outside the range of 'citizen.chemistry' (-20..20)"),
    ("rep.fear > 5 && rep.honor < 5", "single '&'"),
    ("rep.fear > 5 & ", "clause 2 is empty"),
    ("occupation in farmer,farmer", "appears twice"),
    # hints for the mistakes non-programmers make
    ("kingdom.food < 50 and citizen.hunger > 0.5", "join clauses with ' & ' (not 'and')"),
    ("Kingdom.Food < 50", "unknown variable 'Kingdom.Food' (did you mean 'kingdom.food'?)"),
    ("kingdom.food =< 50", "write '<=' instead of '=<'"),
    ("kingdom.food => 50", "write '>=' instead of '=>'"),
    ("kingdom.food ≤ 50", "write '<=' instead of '≤'"),
    ("not event.raid", "write '!event.raid' instead of 'not event.raid'"),
])
def test_condition_rejects_with_precise_message(text, message):
    with pytest.raises(content.CondError) as err:
        content.parse_condition(text, vocab(), "citizen")
    assert message in str(err.value)


def test_kingdom_scope_rejects_citizen_variables():
    with pytest.raises(content.CondError, match="describes a citizen"):
        content.parse_condition("occupation = farmer", vocab(), "kingdom")


# ----------------------------------------------------------------------------- build / check
def test_build_is_deterministic_and_check_passes(lib, tmp_path, capsys):
    out1, out2 = tmp_path / "out1", tmp_path / "out2"
    assert run(capsys, "build", "--content", lib, "--out", out1)[0] == 0
    assert run(capsys, "build", "--content", lib, "--out", out2)[0] == 0
    first = snapshot(out1)
    assert first == snapshot(out2)
    assert sorted(first) == ["Consequences.luau", "Decisions.luau", "Dialogue.luau", "Names.luau",
                             "Occupations.luau", "Reputation.luau", "Traits.luau"]
    code, out = run(capsys, "build", "--content", lib, "--out", out1)
    assert code == 0 and "(0 changed)" in out
    assert snapshot(out1) == first
    assert run(capsys, "check", "--content", lib, "--out", out1)[0] == 0
    for text in first.values():
        lines = text.decode("utf-8").splitlines()
        assert lines[0] == "--!strict"
        assert lines[1].startswith("-- GENERATED by tools/content/content.py from ")
        assert lines[1].endswith(" — do not edit")


def test_real_header_names_the_canonical_source():
    _, model, _ = content.load_and_validate(FIXTURE)
    files = content.generate(model, "games/game-001/content")
    assert files["Dialogue.luau"].splitlines()[1] == \
        "-- GENERATED by tools/content/content.py from games/game-001/content — do not edit"


def test_check_detects_stale_and_missing_files(lib, tmp_path, capsys):
    out = tmp_path / "out"
    run(capsys, "build", "--content", lib, "--out", out)
    (out / "Traits.luau").write_text("-- edited by hand\n", encoding="utf-8")
    code, text = run(capsys, "check", "--content", lib, "--out", out)
    assert code == 1 and "Traits.luau (differs)" in text
    (out / "Names.luau").unlink()
    code, text = run(capsys, "check", "--content", lib, "--out", out)
    assert code == 1 and "Names.luau (missing)" in text


def test_generated_shapes(lib):
    _, model, report = content.load_and_validate(lib)
    assert not report.errors
    files = content.generate(model, "x")
    dialogue = files["Dialogue.luau"]
    assert '{ var = "phase", op = "in", value = { "Founding", "Expansion" } },' in dialogue
    assert '{ var = "event.raid", op = "=", value = false },' in dialogue
    assert "'We bled at the wall, {ruler}. The men say \"a crown that cannot pay cannot rule.\"'" in dialogue
    decisions = files["Decisions.luau"]
    assert "rep = { compassion = 3, honor = 1, competence = 0, courage = 0, authority = 0, fear = 0 }," in decisions
    assert 'consequences = { "c_remember_kindness", "c_cheer" },' in decisions
    assert 'memory_tag = "fed_by_ruler",' in decisions
    assert "notes" not in decisions  # notes never reach the game
    assert 'farm_key = "d_grain_gift",' in decisions  # empty farm_key falls back to the decision id
    consequences = files["Consequences.luau"]
    assert 'param = "shown_kindness",' in consequences
    assert 'target = "occupation:farmer",' in consequences
    reputation = files["Reputation.luau"]
    assert '["BuildingCompleted:barracks"] = { 1, 2 },' in reputation
    assert 'audience_condition = { { var = "citizen.hunger", op = ">=", value = 0.5 } },' in reputation
    names = files["Names.luau"]
    assert '"D\'Arcy-Mae",' in names
    # every module is deeply frozen at require time (review finding L6)
    for name, text in files.items():
        assert "local function deepFreeze<T>(value: T): T" in text, name
        assert "\nreturn deepFreeze({\n" in text and text.endswith("\n})\n"), name
    # presentation (CONTRACTS G13.3) and unlock gating (SPEC §5.4)
    assert dialogue.count("locked = true,") == 1  # s_smith_tools, opened by c_unlock_followup
    assert 'icon = "request",' in dialogue and 'tone = "neutral",' in dialogue


def test_generated_luau_passes_stylua_and_selene(lib, tmp_path, capsys):
    stylua, selene = find_tool("stylua"), find_tool("selene")
    if not stylua or not selene:
        pytest.skip("stylua/selene not installed")
    # awkward strings: both quote kinds, a backslash, non-ASCII letters, a Luau keyword as a trait id
    set_cell(lib, "dialogue/responses.csv", 2, "reaction_text", "It's \"fine\", he says \\ shrugs. Ça va.")
    set_cell(lib, "citizens/names.csv", 2, "name", "Ébrel")
    rows = read_rows(lib / "citizens" / "traits.csv")
    rows.append(["end", "Final", "Thinks every day is the last.", "", "0", "0", "0", "1", "0", "0", "0", "1", "1",
                 "1", "1"])
    write_rows(lib / "citizens" / "traits.csv", rows)
    out = tmp_path / "out"
    code, text = run(capsys, "build", "--content", lib, "--out", out)
    assert code == 0, text
    assert '["end"] = {' in (out / "Traits.luau").read_text(encoding="utf-8")
    files = sorted(str(p) for p in out.glob("*.luau"))
    res = subprocess.run([stylua, "--config-path", str(ROOT / "stylua.toml"), "--check", *files],
                         capture_output=True, text=True)
    assert res.returncode == 0, res.stdout + res.stderr
    res = subprocess.run([selene, str(out)], capture_output=True, text=True, cwd=ROOT)
    assert res.returncode == 0 and "0 errors" in res.stdout, res.stdout + res.stderr


# ----------------------------------------------------------------------------- lint
def test_lint_reports_design_warnings(lib, capsys):
    code, out = run(capsys, "lint", "--content", lib)
    assert code == 0
    assert "[farm_key] no farm_key: repeats of 'd_grain_gift'" in out
    # free decision with a reputation sum above 6 (fear does not count as a gain)
    set_cell(lib, "dialogue/decisions.csv", 5, "cost", "")
    set_cell(lib, "dialogue/decisions.csv", 5, "rep_competence", "7")
    # a response that dominates: d_threaten becomes d_pay_bonus without the cost
    rows = read_rows(lib / "dialogue" / "decisions.csv")
    pay = next(r for r in rows if r[0] == "d_pay_bonus")
    idx = next(i for i, r in enumerate(rows) if r[0] == "d_threaten")
    rows[idx] = ["d_threaten", "threats", *pay[2:4], "", *pay[5:23], "", "", ""]
    write_rows(lib / "dialogue" / "decisions.csv", rows)
    code, out = run(capsys, "lint", "--content", lib)
    assert code == 0, out
    assert "free decision 'd_buy_tools' gives +7 reputation in total, fear aside (> 6)" in out
    assert "response 'r_sp_threat' (d_threaten) is at least as good as every other answer of 's_soldier_pay'" in out
    assert "consequence 'c_desert_risk' is not used by any decision" in out
    set_cell(lib, "dialogue/decisions.csv", 8, "rep_fear", "9")  # a big fear gain is not "better"
    code, out = run(capsys, "lint", "--content", lib)
    assert "response 'r_sp_threat' (d_threaten) is at least as good" in out


def test_lint_scores_answers_with_the_personality_reaction(lib, capsys):
    """The explicit columns alone hide a free answer that most citizens prefer (review finding: drill/inspect)."""
    code, out = run(capsys, "lint", "--content", lib)
    assert "free answer" not in out
    set_cell(lib, "dialogue/decisions.csv", 6, "rep_honor", "0")  # d_promise_later: free, ...
    set_cell(lib, "dialogue/decisions.csv", 6, "rep_competence", "3")  # ... as much reputation as buying tools
    set_cell(lib, "dialogue/decisions.csv", 6, "speaker_loyalty", "8")
    code, out = run(capsys, "lint", "--content", lib)
    assert ("free answer 'r_st_later' (d_promise_later) is the best speaker-loyalty choice for 100% of sampled "
            "citizens of 's_smith_tools'") in out
    # deterministic: the sample is seeded by the scenario id
    assert run(capsys, "lint", "--content", lib)[1] == out


def test_lint_contract_memory_cue_and_number_warnings(lib, capsys):
    set_cell(lib, "dialogue/consequences.csv", 2, "notice_text", "Farmers work 20% slower for 2 minutes.")
    code, out = run(capsys, "lint", "--content", lib)
    assert code == 0
    # RaidDefended is a contracted CONTRACTS G13.4 event: no warning (review finding L5)
    assert "RaidDefended" not in out and "is not a contracted event" not in out
    set_cell(lib, "reputation/actions.csv", 5, "event", "RaidSuffered")
    code, out = run(capsys, "lint", "--content", lib)
    # one warning per event name that is neither a G9 bus event nor a G13.4 reputation event
    assert ("event 'RaidSuffered' (rows 5) is not a contracted event (CONTRACTS G9 bus or G13.4 reputation "
            "events)") in out
    assert out.count("is not a contracted event") == 1
    assert "memory 'fed_by_ruler' is set but no trigger reads it" in out
    assert "memory 'promised_tools' is set but" not in out  # s_smith_tools reads it
    assert "cue 'cue_grumble' is not in the cue vocabulary" in out
    assert "consequences.csv:2:10: [notice_text] the notice types a number" in out


@pytest.mark.parametrize("event", content.REPUTATION_EVENTS + ("UnitsRecruited", "MatchPhaseChanged:Conflict"))
def test_lint_knows_every_contracted_event(lib, capsys, event):
    """Review finding L5: every CONTRACTS G9 and G13.4 event (VS-1 and VS-2) is known to lint."""
    set_cell(lib, "reputation/actions.csv", 5, "event", event)
    code, out = run(capsys, "lint", "--content", lib)
    name = event.split(":")[0]
    assert code == 0 and f"event '{name}'" not in out and "[event]" not in out, out


def test_lint_flags_a_locked_scenario_that_nothing_unlocks(lib, capsys):
    set_cell(lib, "dialogue/decisions.csv", 5, "consequences", "")  # d_buy_tools no longer unlocks s_smith_tools
    code, out = run(capsys, "lint", "--content", lib)
    assert "scenario 's_smith_tools' is locked, but no decision uses a consequence that unlocks it" in out


def test_notes_may_hold_tabs_and_line_breaks(lib, capsys):
    set_cell(lib, "dialogue/decisions.csv", 2, "notes", "line one\nline\ttwo")
    code, out = run(capsys, "validate", "--content", lib)
    assert code == 0, out


@pytest.mark.parametrize("value, ok", [
    ("+housing capacity", True), ("-gold, +military readiness", True), ("- see decisions.csv (row 4)", True),
    ("plain = fine", True), ("=1+1", False), ("@SUM(1+1)", False), ("  =cmd|' /C calc'!A0", False),
    ("+HYPERLINK(\"http://evil.example\")", False), ("-2+3+cmd|' /C calc'!A0", False), ("+=1", False),
    ("- @x", False),
])
def test_notes_may_start_with_plus_or_minus_but_never_hold_a_formula(lib, capsys, value, ok):
    """Review finding L1: notes and actions.direct are opened in spreadsheets too (CSV injection)."""
    set_cell(lib, "reputation/actions.csv", 2, "direct", value)
    set_cell(lib, "dialogue/scenarios.csv", 2, "notes", value)
    code, out = run(capsys, "validate", "--content", lib)
    assert (code == 0) == ok, out
    if not ok:
        assert "reputation/actions.csv:2:4: [direct] a spreadsheet would run this note as a formula" in out
        assert "dialogue/scenarios.csv:2:9: [notes] a spreadsheet would run this note as a formula" in out


def test_bubble_texts_fit_the_bubble_with_the_longest_names(lib, capsys):
    """Review finding L7: the raw limit counts '{ruler}' as 7 characters; the bubble gets the display name."""
    assert content.PLACEHOLDER_MAX == {"name": 16, "family": 16, "friend": 16, "rival": 16, "ruler": 20,
                                       "occupation": 24, "kingdom": 20}
    assert content._expanded_length("{ruler}, {name}!") == 20 + 2 + 16 + 1
    assert content._expanded_length("{unknown} text") == len("{unknown} text")  # validation reports it
    set_cell(lib, "dialogue/responses.csv", 2, "reaction_text", "{ruler}, {name} thanks you. Truly.")
    set_cell(lib, "dialogue/scenarios.csv", 2, "thought_text", "{ruler}, {name} is hungry.")
    code, out = run(capsys, "validate", "--content", lib)
    assert code == 0, out
    over = "{ruler} and {kingdom}: {name} the {occupation} says no"  # 54 raw, 100 filled in
    set_cell(lib, "dialogue/responses.csv", 2, "reaction_text", over)
    set_cell(lib, "dialogue/scenarios.csv", 2, "thought_text", "{ruler}{kingdom}{name}{occupation}{name}")
    code, out = run(capsys, "validate", "--content", lib)
    assert code == 1
    assert ("dialogue/responses.csv:2:5: [reaction_text] with the longest names filled in ({ruler} 20, {kingdom} "
            "20, {name} 16, {occupation} 24 characters) this text can reach 100 characters, but the bubble shows "
            "80 and cuts the rest: shorten it") in out
    assert "dialogue/scenarios.csv:2:7: [thought_text] with the longest names filled in" in out
    # other texts keep their raw limit only (the question is a panel, not a bubble)
    set_cell(lib, "dialogue/responses.csv", 2, "reaction_text", "Bless you!")
    set_cell(lib, "dialogue/scenarios.csv", 2, "thought_text", "Hungry.")
    set_cell(lib, "dialogue/scenarios.csv", 2, "question_text", "{ruler} " * 10 + "{kingdom}")
    code, out = run(capsys, "validate", "--content", lib)
    assert code == 0, out


def test_risk_notices_describe_the_outcome_and_never_repeat_the_chance(lib, capsys):
    """FORGE-00 ruling on PR #67 review #2: the conversation panel shows a risk's chance itself, so no notice
    may write {chance}; a risk (chance < 1) still needs a notice that describes the outcome."""
    rows = read_rows(lib / "dialogue" / "consequences.csv")
    assert rows[5][0] == "c_desert_risk" and "{chance}" not in rows[5][rows[0].index("notice_text")]
    set_cell(lib, "dialogue/consequences.csv", 6, "notice_text", "Some soldiers may desert for {duration}.")
    code, out = run(capsys, "validate", "--content", lib)
    assert code == 0, out
    for row, text in ((6, "{chance}: some soldiers may desert."), (2, "Farmers ({chance}) work slower.")):
        set_cell(lib, "dialogue/consequences.csv", row, "notice_text", text)
        code, out = run(capsys, "validate", "--content", lib)
        assert code == 1
        assert (f"dialogue/consequences.csv:{row}:10: [notice_text] do not write {{chance}}: the conversation panel "
                "already shows a risk's chance next to its text") in out, out
        assert "allowed here" not in out  # one clear message, not the generic placeholder error
        set_cell(lib, "dialogue/consequences.csv", row, "notice_text", "Some soldiers may desert.")
    # the other columns keep the generic message ({chance} never meant anything there)
    set_cell(lib, "reputation/actions.csv", 2, "notice_text", "Odds {chance}.")
    code, out = run(capsys, "validate", "--content", lib)
    assert "[notice_text] placeholder '{chance}' cannot be used in this column" in out
    # lint no longer tells writers to type {chance}
    set_cell(lib, "reputation/actions.csv", 2, "notice_text", "The soldiers appreciate the new barracks.")
    set_cell(lib, "dialogue/consequences.csv", 6, "notice_text", "Half of the soldiers may desert in 5 minutes.")
    code, out = run(capsys, "lint", "--content", lib)
    assert "consequences.csv:6:10: [notice_text] the notice types a number" in out
    assert "write {magnitude}, {magnitude_pct} or {duration}" in out and "write {chance}" not in out
    # the workbook header note lists the placeholders without {chance} and says why
    col = content.SPEC_BY_KEY["consequences"].col("notice_text")
    note = content._column_comment(content.SPEC_BY_KEY["consequences"], col)
    assert "{chance}" not in note.split("Placeholders:")[1].split(".")[0]
    assert "the game shows a risk's chance next to this text" in note
    assert content._column_placeholders("consequences", "notice_text") == content.CONSEQUENCE_NOTICE_PLACEHOLDERS
    assert "chance" not in content.CONSEQUENCE_NOTICE_PLACEHOLDERS


@pytest.mark.parametrize("text", ["Ruler,\u2003we starve.", "A\u2009thin space", "Wait\u00a0here."])
def test_typographic_spaces_inside_a_text_are_fine(lib, capsys, text):
    """Ruling #10: U+2000-U+200A (and other spaces) are refused only as a text's whole content."""
    set_cell(lib, "dialogue/scenarios.csv", 2, "question_text", text)
    code, out = run(capsys, "validate", "--content", lib)
    assert code == 0, out


def test_blank_and_invisible_character_sets():
    for ch in ("\u2800", "\u034f", "\U0001D159", "\u115f", "\u1160", "\u3164", "\uffa0", "\u180e"):
        assert content._invisible_problem(f"a{ch}b", "text"), f"U+{ord(ch):04X} must be refused"
        assert content._invisible_problem(f"a{ch}b", "notes"), f"U+{ord(ch):04X} must be refused in notes too"
    for o in range(0x2000, 0x200B):
        assert content._invisible_problem(f"a{chr(o)}b", "text") is None  # a visible space between words
        assert content._blank_text(chr(o)) is not None  # ... but never a whole text
    assert content._blank_text("x") is None and content._blank_text(" \u2800x ") is None
    assert content._blank_text("\u200b \u3000") == "U+200B zero-width space"


DEFAULT_IGNORABLE_MARKS = ["\ufe00", "\ufe0f", "\U000e0100", "\U000e01ef", "\u17b4", "\u17b5", "\u180b", "\u180c",
                           "\u180d", "\u180f", "\ufff0", "\ufff8", "\U000e0000", "\U000e0002", "\u2065"]


@pytest.mark.parametrize("ch", DEFAULT_IGNORABLE_MARKS, ids=[f"U+{ord(c):04X}" for c in DEFAULT_IGNORABLE_MARKS])
def test_default_ignorable_characters_are_a_blank_text(lib, capsys, ch):
    """Ruling #10 follow-up: a TEXT holding only default-ignorable code points (variation selectors, Khmer inherent
    vowels, Mongolian free variation selectors, reserved ones) shows as an empty line, so it is refused as blank."""
    assert content._blank_text(ch) == f"U+{ord(ch):04X} {content._char_name(ch)}"
    assert content._blank_text(f"{ch} {ch}") is not None
    set_cell(lib, "dialogue/responses.csv", 2, "response_text", ch)
    code, out = run(capsys, "validate", "--content", lib)
    assert code == 1, out
    assert f"dialogue/responses.csv:2:4: [response_text] text is blank: it holds only spaces or invisible characters " \
           f"(U+{ord(ch):04X} " in out


@pytest.mark.parametrize("text", ["Yes \u2764\ufe0f", "Hm\u2026 \u263a\ufe0f fine", "Ok \u2764\U000e0100"])
def test_variation_selector_after_a_character_is_fine(lib, capsys, text):
    """Ruling #10 follow-up: a variation selector is refused only as a text's whole content, never after an emoji."""
    assert content._invisible_problem(text, "text") is None and content._blank_text(text) is None
    set_cell(lib, "dialogue/responses.csv", 2, "response_text", text)
    code, out = run(capsys, "validate", "--content", lib)
    assert code == 0, out


def test_default_ignorable_table_is_the_unicode_property():
    """_DEFAULT_IGNORABLE is DerivedCoreProperties Default_Ignorable_Code_Point: 4174 code points (Unicode 15.1)."""
    import unicodedata
    counts: dict[str, int] = {}
    for lo, hi in content._DEFAULT_IGNORABLE:
        for o in range(lo, hi + 1):
            cat = unicodedata.category(chr(o))
            counts[cat] = counts.get(cat, 0) + 1
            if cat == "Cn":
                assert content._char_name(chr(o)) == "unassigned character"
    assert sum(counts.values()) == 4174
    assert counts["Mn"] == 263 and counts["Lo"] == 4  # the marks and fillers category Cf does not cover


def test_thought_icon_and_reaction_tone_columns(lib, tmp_path, capsys):
    rows = read_rows(lib / "dialogue" / "scenarios.csv")
    rows[0].insert(6, "thought_icon")
    for r, icon in zip(rows[1:], ("concern", "", "idea")):
        r.insert(6, icon)
    write_rows(lib / "dialogue" / "scenarios.csv", rows)
    rows = read_rows(lib / "dialogue" / "responses.csv")
    rows[0].insert(5, "reaction_tone")
    for r in rows[1:]:
        r.insert(5, "happy" if r[1] == "r_hf_share" else "")
    write_rows(lib / "dialogue" / "responses.csv", rows)
    out = tmp_path / "out"
    code, text = run(capsys, "build", "--content", lib, "--out", out)
    assert code == 0, text
    dialogue = (out / "Dialogue.luau").read_text(encoding="utf-8")
    assert 'icon = "concern",' in dialogue and 'icon = "idea",' in dialogue and 'icon = "request",' in dialogue
    assert 'tone = "happy",' in dialogue and 'tone = "neutral",' in dialogue
    set_cell(lib, "dialogue/scenarios.csv", 2, "thought_icon", "shout")
    code, text = run(capsys, "validate", "--content", lib)
    assert code == 1 and "dialogue/scenarios.csv:2:7: [thought_icon] 'shout' is not allowed" in text


# ----------------------------------------------------------------------------- CSV round trip helpers
def test_csv_parser_keeps_raw_bytes_for_unchanged_rows():
    text = '﻿a,b\r\n"x, y","say ""hi"""\r\n"quoted",plain\r\n'
    doc = content.parse_csv(text)
    assert [r.fields for r in doc.records] == [["a", "b"], ["x, y", 'say "hi"'], ["quoted", "plain"]]
    assert doc.newline == "\r\n" and doc.bom and doc.trailing_newline
    rows = [list(r.fields) for r in doc.records]
    assert content.render_csv(rows, doc) == text
    rows[2][1] = "changed"
    assert content.render_csv(rows, doc) == '﻿a,b\r\n"x, y","say ""hi"""\r\nquoted,changed\r\n'


# ----------------------------------------------------------------------------- workbook
def _xlsx_roundtrip(lib: Path, tmp_path: Path, capsys) -> None:
    before = csv_bytes(lib)
    book = tmp_path / "book.xlsx"
    code, out = run(capsys, "export-xlsx", "--content", lib, book)
    assert code == 0, out
    code, out = run(capsys, "import-xlsx", "--content", lib, book)
    assert code == 0, out
    assert "0 CSV file(s) updated" in out
    assert csv_bytes(lib) == before


@needs_openpyxl
def test_xlsx_roundtrip_is_byte_identical(lib, tmp_path, capsys):
    _xlsx_roundtrip(lib, tmp_path, capsys)


@needs_openpyxl
def test_xlsx_roundtrip_keeps_crlf_and_extra_quoting(lib, tmp_path, capsys):
    path = lib / "dialogue" / "scenarios.csv"
    buf = io.StringIO()
    csv.writer(buf, lineterminator="\r\n", quoting=csv.QUOTE_ALL).writerows(read_rows(path))
    path.write_text(buf.getvalue(), encoding="utf-8", newline="")
    _xlsx_roundtrip(lib, tmp_path, capsys)


@needs_openpyxl
def test_xlsx_roundtrip_real_content(tmp_path, capsys):
    lib = tmp_path / "content"
    for spec in content.SPECS:
        (lib / spec.path).parent.mkdir(parents=True, exist_ok=True)
        shutil.copyfile(REAL / spec.path, lib / spec.path)
    _xlsx_roundtrip(lib, tmp_path, capsys)


@needs_openpyxl
def test_workbook_layout(lib, tmp_path, capsys):
    from openpyxl import load_workbook

    book = tmp_path / "book.xlsx"
    run(capsys, "export-xlsx", "--content", lib, book)
    wb = load_workbook(book)
    assert wb.sheetnames == ["README", "lists", *[s.key for s in content.SPECS]]
    assert wb["lists"].sheet_state == "hidden"
    ws = wb["scenarios"]
    assert ws.freeze_panes == "B2"
    header = [c.value for c in ws[1]]
    q = header.index("question_text") + 1
    trig = header.index("trigger") + 1
    assert ws.cell(1, q).comment.text.startswith("EDIT")
    assert "160" in ws.cell(1, q).comment.text
    assert ws.cell(2, q).fill.fgColor.rgb.endswith(content.FILL["text"])
    assert ws.cell(1, trig).comment.text.startswith("MECHANICS")
    assert ws.cell(2, trig).fill.fgColor.rgb.endswith(content.FILL["mech"])
    assert ws.cell(2, trig).number_format == "@"
    kinds = {(dv.type, str(dv.formula1)) for dv in ws.data_validations.dataValidation}
    assert ("textLength", "160") in kinds and ("textLength", "48") in kinds
    assert any(t == "list" and f.startswith("lists!") for t, f in kinds)  # occupation dropdown
    rv = {(dv.type, str(dv.formula1)) for dv in wb["responses"].data_validations.dataValidation}
    assert ("list", "decisions!$A$2:$A$2000") in rv  # decision_id dropdown, live from the decisions sheet
    assert ("textLength", "80") in rv  # reaction_text: FORGE-04's bubble maxChars (decision L7)
    rheader = [c.value for c in wb["responses"][1]]
    assert "Max 80 characters" in wb["responses"].cell(1, rheader.index("reaction_text") + 1).comment.text
    stamps = [c.value for c in wb["lists"][1]]
    assert stamps[-4:] == ["stamp_file", "stamp_sha256", "stamp_cell", "stamp_cell_sha256"]
    cv = {str(dv.formula1) for dv in wb["consequences"].data_validations.dataValidation if dv.type == "list"}
    assert len(cv) == 2  # kind and target


@needs_openpyxl
def test_editing_text_in_the_workbook_changes_only_that_text(lib, tmp_path, capsys):
    from openpyxl import load_workbook

    out1, out2 = tmp_path / "out1", tmp_path / "out2"
    assert run(capsys, "build", "--content", lib, "--out", out1)[0] == 0
    before_csv = csv_bytes(lib)
    book = tmp_path / "book.xlsx"
    run(capsys, "export-xlsx", "--content", lib, book)
    wb = load_workbook(book)
    ws = wb["scenarios"]
    header = [c.value for c in ws[1]]
    new_text = 'Pay us, {ruler}, or the "men" will walk.'
    ws.cell(4, header.index("question_text") + 1).value = new_text
    wb.save(book)
    code, out = run(capsys, "import-xlsx", "--content", lib, "--out", out2, book, "--build")
    assert code == 0, out
    after_csv = csv_bytes(lib)
    changed = [rel for rel in CSV_PATHS if before_csv[rel] != after_csv[rel]]
    assert changed == ["dialogue/scenarios.csv"]
    old_lines = before_csv[changed[0]].decode().splitlines()
    new_lines = after_csv[changed[0]].decode().splitlines()
    assert len(old_lines) == len(new_lines)
    diff = [i for i, (a, b) in enumerate(zip(old_lines, new_lines)) if a != b]
    assert diff == [3]
    old_row = next(csv.reader([old_lines[3]]))
    new_row = next(csv.reader([new_lines[3]]))
    q = header.index("question_text")
    assert new_row[q] == new_text
    assert [v for i, v in enumerate(old_row) if i != q] == [v for i, v in enumerate(new_row) if i != q]
    # generated Luau: only the question line of Dialogue.luau changed
    a, b = snapshot(out1), snapshot(out2)
    assert [k for k in a if a[k] != b[k]] == ["Dialogue.luau"]
    la, lb = a["Dialogue.luau"].decode().splitlines(), b["Dialogue.luau"].decode().splitlines()
    assert len(la) == len(lb)
    changed_lines = [(x, y) for x, y in zip(la, lb) if x != y]
    assert len(changed_lines) == 1
    assert changed_lines[0][1].strip() == "question = 'Pay us, {ruler}, or the \"men\" will walk.',"


@needs_openpyxl
def test_import_refuses_unknown_sheets_and_columns(lib, tmp_path, capsys):
    from openpyxl import load_workbook

    before = csv_bytes(lib)
    book = tmp_path / "book.xlsx"
    run(capsys, "export-xlsx", "--content", lib, book)
    wb = load_workbook(book)
    wb.create_sheet("Sheet2")
    wb["decisions"].cell(1, 1).value = "decision"
    wb.save(book)
    code, out = run(capsys, "import-xlsx", "--content", lib, book)
    assert code == 1
    assert "unknown sheet 'Sheet2'" in out
    assert "decisions!A1 [decision]: unknown column 'decision' (did you mean 'decision_id'?)" in out
    assert "no CSV was changed" in out
    assert csv_bytes(lib) == before


@needs_openpyxl
def test_import_validates_before_writing_and_names_the_cell(lib, tmp_path, capsys):
    from openpyxl import load_workbook

    before = csv_bytes(lib)
    book = tmp_path / "book.xlsx"
    run(capsys, "export-xlsx", "--content", lib, book)
    wb = load_workbook(book)
    ws = wb["responses"]
    header = [c.value for c in ws[1]]
    ws.cell(5, header.index("response_text") + 1).value = "x" * 71
    ws.cell(9, 1).value = None  # a blank row at the end is ignored
    wb.save(book)
    code, out = run(capsys, "import-xlsx", "--content", lib, book)
    assert code == 1
    assert "book.xlsx › responses!D5 [response_text]: text is 71 characters; the limit is 70" in out
    assert csv_bytes(lib) == before


@needs_openpyxl
def test_import_converts_typed_cells_back_to_text(lib, tmp_path, capsys):
    from openpyxl import load_workbook

    book = tmp_path / "book.xlsx"
    run(capsys, "export-xlsx", "--content", lib, book)
    wb = load_workbook(book)
    ws = wb["scenarios"]
    header = [c.value for c in ws[1]]
    ws.cell(2, header.index("priority") + 1).value = 81  # typed as a number, as Excel would
    ws.cell(2, header.index("once_per_citizen") + 1).value = True
    wb.save(book)
    code, out = run(capsys, "import-xlsx", "--content", lib, book)
    assert code == 0, out
    row = read_rows(lib / "dialogue" / "scenarios.csv")[1]
    assert row[3] == "81" and row[5] == "true"


def _sheet_cell(wb, sheet: str, row: int, column: str):
    ws = wb[sheet]
    header = [c.value for c in ws[1]]
    return ws.cell(row, header.index(column) + 1)


@needs_openpyxl
def test_import_refuses_mechanics_from_a_stale_workbook(lib, tmp_path, capsys):
    """Review finding: an old workbook silently reverted a cost changed in the CSV after the export."""
    from openpyxl import load_workbook

    book = tmp_path / "book.xlsx"
    run(capsys, "export-xlsx", "--content", lib, book)
    set_cell(lib, "dialogue/decisions.csv", 2, "cost", "food:30")  # a developer edit after the export
    after_dev = csv_bytes(lib)
    wb = load_workbook(book)  # the editor only rewrote two lines of text
    _sheet_cell(wb, "responses", 2, "reaction_text").value = "Bless you, truly!"
    _sheet_cell(wb, "scenarios", 3, "thought_text").value = "My hammer is cracked."
    wb.save(book)
    code, out = run(capsys, "import-xlsx", "--content", lib, book)
    assert code == 1, out
    assert "decisions!E2 [cost]: decisions › d_share_grain [cost]: 'food:30' → 'food:20'" in out
    assert "dialogue/decisions.csv changed after this workbook was exported" in out
    assert "--text-only" in out and "no CSV was changed" in out
    assert csv_bytes(lib) == after_dev
    # --text-only takes the two texts and keeps the developer's cost
    code, out = run(capsys, "import-xlsx", "--content", lib, book, "--text-only")
    assert code == 0, out
    assert "2 change(s) from the workbook: TEXT 2, MECHANICS 0, NOTES 0" in out
    assert "1 mechanics change(s) in the workbook were NOT imported" in out
    assert "ignored   decisions › d_share_grain [cost]: 'food:30' → 'food:20'" in out
    assert read_rows(lib / "dialogue" / "decisions.csv")[1][4] == "food:30"
    assert read_rows(lib / "dialogue" / "responses.csv")[1][4] == "Bless you, truly!"
    now = csv_bytes(lib)
    assert [rel for rel in CSV_PATHS if now[rel] != after_dev[rel]] == ["dialogue/scenarios.csv",
                                                                         "dialogue/responses.csv"]


@needs_openpyxl
def test_import_lists_every_change_and_takes_mechanics_from_a_fresh_workbook(lib, tmp_path, capsys):
    from openpyxl import load_workbook

    book = tmp_path / "book.xlsx"
    run(capsys, "export-xlsx", "--content", lib, book)
    wb = load_workbook(book)
    _sheet_cell(wb, "scenarios", 2, "priority").value = "81"
    _sheet_cell(wb, "scenarios", 2, "notes").value = "my note"
    _sheet_cell(wb, "responses", 3, "response_text").value = "Rations for all. Me too."
    wb.save(book)
    code, out = run(capsys, "import-xlsx", "--content", lib, book)
    assert code == 0, out
    assert "3 change(s) from the workbook: TEXT 1, MECHANICS 1, NOTES 1" in out
    assert "MECHANICS scenarios › s_hungry_farmer [priority]: '80' → '81'" in out
    assert "NOTES     scenarios › s_hungry_farmer [notes]: \"\" → 'my note'" in out
    assert "TEXT      responses › r_hf_ration [response_text]:" in out
    assert read_rows(lib / "dialogue" / "scenarios.csv")[1][3] == "81"


@needs_openpyxl
def test_reimporting_one_workbook_after_more_text_edits(lib, tmp_path, capsys):
    from openpyxl import load_workbook

    book = tmp_path / "book.xlsx"
    run(capsys, "export-xlsx", "--content", lib, book)
    wb = load_workbook(book)
    _sheet_cell(wb, "responses", 2, "reaction_text").value = "First pass."
    wb.save(book)
    assert run(capsys, "import-xlsx", "--content", lib, book)[0] == 0
    wb = load_workbook(book)  # the same, now older, workbook: more text only
    _sheet_cell(wb, "responses", 3, "reaction_text").value = "Second pass."
    wb.save(book)
    code, out = run(capsys, "import-xlsx", "--content", lib, book)
    assert code == 0, out
    assert "TEXT 1, MECHANICS 0" in out
    rows = read_rows(lib / "dialogue" / "responses.csv")
    assert rows[1][4] == "First pass." and rows[2][4] == "Second pass."


@needs_openpyxl
def test_import_without_an_export_stamp_takes_text_but_not_mechanics(lib, tmp_path, capsys):
    from openpyxl import load_workbook

    book = tmp_path / "book.xlsx"
    run(capsys, "export-xlsx", "--content", lib, book)
    wb = load_workbook(book)
    lists = wb["lists"]
    header = [c.value for c in lists[1]]
    lists.delete_cols(header.index("stamp_file") + 1, 2)
    _sheet_cell(wb, "scenarios", 2, "priority").value = "81"
    wb.save(book)
    code, out = run(capsys, "import-xlsx", "--content", lib, book)
    assert code == 1 and "this workbook has no export stamp" in out
    wb = load_workbook(book)
    _sheet_cell(wb, "scenarios", 2, "priority").value = "80"
    _sheet_cell(wb, "scenarios", 2, "thought_text").value = "Still hungry."
    wb.save(book)
    code, out = run(capsys, "import-xlsx", "--content", lib, book)
    assert code == 0, out
    assert read_rows(lib / "dialogue" / "scenarios.csv")[1][6] == "Still hungry."


@needs_openpyxl
def test_sorting_the_workbook_never_reorders_the_csv(lib, tmp_path, capsys):
    """Review finding: sorting the scenarios sheet changed Dialogue.order, the priority tie-break."""
    from openpyxl import load_workbook

    out1, out2 = tmp_path / "out1", tmp_path / "out2"
    assert run(capsys, "build", "--content", lib, "--out", out1)[0] == 0
    before = csv_bytes(lib)
    book = tmp_path / "book.xlsx"
    run(capsys, "export-xlsx", "--content", lib, book)
    wb = load_workbook(book)
    for sheet in ("scenarios", "responses", "names", "actions"):  # keyed and keyless tables
        ws = wb[sheet]
        values = [[c.value for c in row] for row in ws.iter_rows(min_row=2)]
        for r, row in enumerate(reversed(values), 2):
            for c, v in enumerate(row, 1):
                ws.cell(r, c).value = v
    wb.save(book)
    code, out = run(capsys, "import-xlsx", "--content", lib, "--out", out2, book, "--build")
    assert code == 0, out
    assert "0 CSV file(s) updated" in out
    assert csv_bytes(lib) == before
    assert snapshot(out1) == snapshot(out2)
    # sorted and edited: only the edited text changes
    wb = load_workbook(book)
    ws = wb["scenarios"]
    header = [c.value for c in ws[1]]
    row = next(r for r in range(2, ws.max_row + 1) if ws.cell(r, 1).value == "s_smith_tools")
    ws.cell(row, header.index("thought_text") + 1).value = "My tools are worn to nothing."
    wb.save(book)
    assert run(capsys, "import-xlsx", "--content", lib, book)[0] == 0
    old, new = before["dialogue/scenarios.csv"].decode().splitlines(), csv_bytes(lib)["dialogue/scenarios.csv"].decode()
    assert [i for i, (a, b) in enumerate(zip(old, new.splitlines())) if a != b] == [2]


@needs_openpyxl
def test_text_only_refuses_rows_moved_in_a_table_without_a_key(lib, tmp_path, capsys):
    from openpyxl import load_workbook

    book = tmp_path / "book.xlsx"
    run(capsys, "export-xlsx", "--content", lib, book)
    wb = load_workbook(book)
    wb["names"].delete_rows(2)
    wb.save(book)
    before = csv_bytes(lib)
    code, out = run(capsys, "import-xlsx", "--content", lib, book, "--text-only")
    assert code == 1
    assert "rows were added, removed or moved in sheet 'names'" in out
    assert csv_bytes(lib) == before


@needs_openpyxl
def test_responses_sheet_has_read_only_help_columns(lib, tmp_path, capsys):
    from openpyxl import load_workbook

    book = tmp_path / "book.xlsx"
    run(capsys, "export-xlsx", "--content", lib, book)
    wb = load_workbook(book)
    ws = wb["responses"]
    header = [c.value for c in ws[1]]
    assert header[-2:] == ["info: question", "info: effects"]
    assert ws.cell(1, len(header)).comment.text.startswith("READ ONLY")
    assert ws.cell(2, len(header) - 1).value.startswith("{ruler}, the granary is low")
    effects = ws.cell(2, len(header)).value
    assert effects.startswith("cost food:20 | rep compassion +3, honor +1 | speaker +4")
    assert "then c_remember_kindness, c_cheer" in effects
    ws.cell(2, len(header)).value = "edited help text"  # ignored on import
    wb.save(book)
    before = csv_bytes(lib)
    code, out = run(capsys, "import-xlsx", "--content", lib, book)
    assert code == 0 and "0 CSV file(s) updated" in out, out
    assert csv_bytes(lib) == before


@needs_openpyxl
def test_import_refuses_formula_cells(lib, tmp_path, capsys):
    """Review finding L1: data_only=True blanked a formula without a cached value (here an optional notice, so
    the import 'succeeded' and lost the text) and silently took the cached value of any other formula."""
    from openpyxl import load_workbook

    before = csv_bytes(lib)
    book = tmp_path / "book.xlsx"
    run(capsys, "export-xlsx", "--content", lib, book)
    wb = load_workbook(book)
    cell = _sheet_cell(wb, "consequences", 2, "notice_text")
    cell.value = '="Farmers work slower."'  # openpyxl stores this as a formula (data_type 'f'), no cached value
    assert cell.data_type == "f"
    _sheet_cell(wb, "scenarios", 2, "priority").value = "=80+1"
    wb.save(book)
    for flags in ((), ("--text-only",)):
        code, out = run(capsys, "import-xlsx", "--content", lib, book, *flags)
        assert code == 1, out
        assert ("book.xlsx: consequences!J2: this cell is a formula ('=\"Farmers work slower.\"'); formulas are "
                "never imported: type the value itself, as text") in out
        assert "book.xlsx: scenarios!D2: this cell is a formula ('=80+1')" in out
        assert csv_bytes(lib) == before


@needs_openpyxl
@pytest.mark.parametrize("value", ["=cmd|' /C calc'!A0", "+5 rations for all", "-2+3+cmd|' /C calc'!A0",
                                   "@SUM(1+1)"])
def test_import_refuses_text_that_a_spreadsheet_would_run(lib, tmp_path, capsys, value):
    """Review finding L1 (CSV injection): a TEXT cell typed as text but starting with = + - @ is refused."""
    from openpyxl import load_workbook

    before = csv_bytes(lib)
    book = tmp_path / "book.xlsx"
    run(capsys, "export-xlsx", "--content", lib, book)
    wb = load_workbook(book)
    cell = _sheet_cell(wb, "responses", 2, "reaction_text")
    cell.value = value
    cell.data_type = "s"  # a string, as Excel stores text typed into a Text cell
    wb.save(book)
    code, out = run(capsys, "import-xlsx", "--content", lib, book, "--text-only")
    assert code == 1, out
    assert f"book.xlsx › responses!E2 [reaction_text]: text may not start with '{value[0]}'" in out
    assert csv_bytes(lib) == before


@needs_openpyxl
@pytest.mark.parametrize("text_only", [False, True], ids=["full", "text-only"])
def test_import_keeps_text_that_changed_in_the_csv_after_the_export(lib, tmp_path, capsys, text_only):
    """Review finding M4: a CSV text edited after the export was silently reverted by the old workbook,
    even with --text-only. Now a cell the writer did not touch keeps the newer CSV text (and is listed)."""
    from openpyxl import load_workbook

    book = tmp_path / "book.xlsx"
    run(capsys, "export-xlsx", "--content", lib, book)
    set_cell(lib, "dialogue/responses.csv", 2, "reaction_text", "Dev fixed this line.")  # after the export
    set_cell(lib, "dialogue/scenarios.csv", 2, "notes", "dev note")
    wb = load_workbook(book)
    _sheet_cell(wb, "responses", 3, "reaction_text").value = "The writer's new line."
    wb.save(book)
    code, out = run(capsys, "import-xlsx", "--content", lib, book, *(["--text-only"] if text_only else []))
    assert code == 0, out
    rows = read_rows(lib / "dialogue" / "responses.csv")
    assert rows[1][4] == "Dev fixed this line."
    assert rows[2][4] == "The writer's new line."
    assert read_rows(lib / "dialogue" / "scenarios.csv")[1][-1] == "dev note"
    assert "1 change(s) from the workbook: TEXT 1, MECHANICS 0, NOTES 0" in out
    assert "import-xlsx: 2 cell(s) kept as in the CSV (changed there after your export" in out
    assert ("kept      responses › r_hf_share [reaction_text]: 'Bless you, {ruler}! We will not forget this.' → "
            "'Dev fixed this line.'") in out
    assert "kept      scenarios › s_hungry_farmer [notes]" in out


@needs_openpyxl
@pytest.mark.parametrize("text_only", [False, True], ids=["full", "text-only"])
def test_import_refuses_text_changed_in_both_the_csv_and_the_workbook(lib, tmp_path, capsys, text_only):
    from openpyxl import load_workbook

    book = tmp_path / "book.xlsx"
    run(capsys, "export-xlsx", "--content", lib, book)
    set_cell(lib, "dialogue/responses.csv", 2, "reaction_text", "Dev fixed this line.")
    after_dev = csv_bytes(lib)
    wb = load_workbook(book)
    _sheet_cell(wb, "responses", 2, "reaction_text").value = "The writer rewrote it too."
    _sheet_cell(wb, "responses", 3, "reaction_text").value = "An unrelated edit."
    wb.save(book)
    code, out = run(capsys, "import-xlsx", "--content", lib, book, *(["--text-only"] if text_only else []))
    assert code == 1, out
    assert ("book.xlsx › responses!E2 [reaction_text]: responses › r_hf_share [reaction_text]: refused: it was "
            "also changed in dialogue/responses.csv after this workbook was exported, possibly by an earlier import "
            "of this same workbook (CSV now 'Dev fixed this line.', workbook 'The writer rewrote it too.'). Export a "
            "fresh workbook and redo this edit there") in out
    assert "no CSV was changed" in out
    assert csv_bytes(lib) == after_dev


@needs_openpyxl
def test_import_without_cell_stamps_refuses_text_only_where_the_csv_changed(lib, tmp_path, capsys):
    """A workbook from the older tool has file stamps only: text is taken from tables whose CSV is unchanged,
    and refused where the CSV changed (the import cannot tell which cells moved)."""
    from openpyxl import load_workbook

    book = tmp_path / "book.xlsx"
    run(capsys, "export-xlsx", "--content", lib, book)
    wb = load_workbook(book)
    lists = wb["lists"]
    header = [c.value for c in lists[1]]
    lists.delete_cols(header.index("stamp_cell") + 1, 2)
    _sheet_cell(wb, "responses", 3, "reaction_text").value = "A writer's line."
    _sheet_cell(wb, "scenarios", 2, "thought_text").value = "Still hungry."
    wb.save(book)
    set_cell(lib, "dialogue/responses.csv", 2, "reaction_text", "Dev fixed this line.")
    code, out = run(capsys, "import-xlsx", "--content", lib, book)
    assert code == 1, out
    assert "responses!E2 [reaction_text]" in out and "no per-cell export stamp" in out
    assert "scenarios!" not in out  # scenarios.csv did not change: its text would be taken
    set_cell(lib, "dialogue/responses.csv", 2, "reaction_text", "Bless you, {ruler}! We will not forget this.")
    code, out = run(capsys, "import-xlsx", "--content", lib, book)  # the CSV is as exported again
    assert code == 0, out
    assert read_rows(lib / "dialogue" / "responses.csv")[2][4] == "A writer's line."


@needs_openpyxl
@pytest.mark.parametrize("damage", ["cp1252", "extra column"])
def test_import_never_replaces_a_csv_it_cannot_read(lib, tmp_path, capsys, damage):
    """Verification of review finding M4: a CSV that became unreadable after the export (Excel's 'CSV (Comma
    delimited)' saves cp1252; or a column was added) was rewritten from the stale workbook, silently undoing the
    developer's edits. Now the import refuses it and leaves the bytes alone."""
    book = tmp_path / "book.xlsx"
    assert run(capsys, "export-xlsx", "--content", lib, book)[0] == 0
    path = lib / "dialogue" / "scenarios.csv"
    rows = read_rows(path)
    rows[2][rows[0].index("question_text")] = "Dev: the café is closed, {ruler}."
    rows[2][rows[0].index("priority")] = "99"
    if damage == "extra column":
        rows[0].append("reviewer")
        for r in rows[1:]:
            r.append("dev")
    write_rows(path, rows)
    if damage == "cp1252":
        path.write_bytes(path.read_text(encoding="utf-8").encode("cp1252"))
    before = csv_bytes(lib)
    for flags in ((), ("--text-only",)):
        code, out = run(capsys, "import-xlsx", "--content", lib, book, *flags)
        assert code == 1, out
        assert "no CSV was changed" in out
        assert csv_bytes(lib) == before
    code, out = run(capsys, "import-xlsx", "--content", lib, book)
    reason = "not UTF-8 text: byte 0xE9" if damage == "cp1252" else "unknown column 'reviewer'"
    assert "book.xlsx › scenarios: dialogue/scenarios.csv cannot be read (" in out and reason in out, out
    assert "it never replaces the file blindly" in out


@needs_openpyxl
def test_import_recreates_a_missing_csv_from_the_workbook(lib, tmp_path, capsys):
    book = tmp_path / "book.xlsx"
    assert run(capsys, "export-xlsx", "--content", lib, book)[0] == 0
    before = csv_bytes(lib)
    (lib / "dialogue" / "scenarios.csv").unlink()
    code, out = run(capsys, "import-xlsx", "--content", lib, book)
    assert code == 0, out
    assert "dialogue/scenarios.csv" in out
    assert read_rows(lib / "dialogue" / "scenarios.csv") == list(csv.reader(io.StringIO(
        before["dialogue/scenarios.csv"].decode("utf-8"), newline="")))


@needs_openpyxl
def test_import_refuses_spreadsheet_error_values(lib, tmp_path, capsys):
    """Verification of review finding L1: a pasted error value ('#REF!', data type 'e') was imported as text."""
    from openpyxl import load_workbook

    before = csv_bytes(lib)
    book = tmp_path / "book.xlsx"
    run(capsys, "export-xlsx", "--content", lib, book)
    wb = load_workbook(book)
    cell = _sheet_cell(wb, "responses", 3, "reaction_text")
    cell.value = "#REF!"
    cell.data_type = "e"
    wb.save(book)
    assert load_workbook(book)["responses"]["E3"].data_type == "e"
    for flags in ((), ("--text-only",)):
        code, out = run(capsys, "import-xlsx", "--content", lib, book, *flags)
        assert code == 1, out
        assert "book.xlsx: responses!E3: this cell shows a spreadsheet error ('#REF!'); type the value itself" in out
        assert csv_bytes(lib) == before


@needs_openpyxl
def test_import_refuses_a_formula_typed_as_text_in_a_notes_cell(lib, tmp_path, capsys):
    """Verification of review finding L1: notes and actions.direct took '=cmd|...' as plain strings."""
    from openpyxl import load_workbook

    before = csv_bytes(lib)
    book = tmp_path / "book.xlsx"
    run(capsys, "export-xlsx", "--content", lib, book)
    wb = load_workbook(book)
    for sheet, column, value in (("decisions", "notes", "=cmd|' /C calc'!A0"),
                                 ("actions", "direct", "@SUM(1+1)*cmd|' /C calc'!A0")):
        cell = _sheet_cell(wb, sheet, 2, column)
        cell.value = value
        cell.data_type = "s"
    wb.save(book)
    code, out = run(capsys, "import-xlsx", "--content", lib, book)
    assert code == 1, out
    assert "book.xlsx › decisions!Z2 [notes]: a spreadsheet would run this note as a formula" in out
    assert "book.xlsx › actions!D2 [direct]: a spreadsheet would run this note as a formula" in out
    assert csv_bytes(lib) == before


@needs_openpyxl
def test_editing_one_cell_twice_in_one_workbook_needs_a_fresh_export(lib, tmp_path, capsys):
    """Verification of review finding M4: after an import the workbook's stamps describe the CSV before that
    import, so a second edit of the same cell is refused (README step 6: export a fresh workbook)."""
    from openpyxl import load_workbook

    book = tmp_path / "book.xlsx"
    run(capsys, "export-xlsx", "--content", lib, book)
    wb = load_workbook(book)
    _sheet_cell(wb, "responses", 3, "reaction_text").value = "Writer first try."
    wb.save(book)
    assert run(capsys, "import-xlsx", "--content", lib, book, "--text-only")[0] == 0
    wb = load_workbook(book)
    _sheet_cell(wb, "responses", 3, "reaction_text").value = "Writer second try."
    wb.save(book)
    code, out = run(capsys, "import-xlsx", "--content", lib, book, "--text-only")
    assert code == 1, out
    assert "possibly by an earlier import of this same workbook" in out
    fresh = tmp_path / "fresh.xlsx"
    run(capsys, "export-xlsx", "--content", lib, fresh)
    wb = load_workbook(fresh)
    _sheet_cell(wb, "responses", 3, "reaction_text").value = "Writer second try."
    wb.save(fresh)
    assert run(capsys, "import-xlsx", "--content", lib, fresh, "--text-only")[0] == 0
    assert read_rows(lib / "dialogue" / "responses.csv")[2][4] == "Writer second try."


def test_export_names_the_cell_with_an_invisible_character(lib, tmp_path, capsys):
    pytest.importorskip("openpyxl")
    set_cell(lib, "dialogue/responses.csv", 2, "reaction_text", "Bless you!\x0bWe'll work those fields.")
    book = tmp_path / "x.xlsx"
    code, out = run(capsys, "export-xlsx", "--content", lib, book)
    assert code == 1
    assert "dialogue/responses.csv:2:5: [reaction_text] cannot export: invisible character U+000B" in out
    assert not book.exists()


@needs_openpyxl
def test_import_rejects_invisible_characters(lib, tmp_path, capsys):
    from openpyxl import load_workbook

    book = tmp_path / "book.xlsx"
    run(capsys, "export-xlsx", "--content", lib, book)
    wb = load_workbook(book)
    _sheet_cell(wb, "scenarios", 2, "question_text").value = "Ruler\u200bword, we starve."
    _sheet_cell(wb, "responses", 2, "reaction_text").value = "line\u2028sep"
    wb.save(book)
    before = csv_bytes(lib)
    code, out = run(capsys, "import-xlsx", "--content", lib, book)
    assert code == 1
    assert "scenarios!H2 [question_text]: invisible character U+200B (zero-width space) at position 6" in out
    assert "responses!E2 [reaction_text]: invisible character U+2028 (line separator)" in out
    assert csv_bytes(lib) == before


# ----------------------------------------------------------------------------- real content regressions
def test_real_content_keeps_the_balance_review_fixes(capsys):
    code, out = run(capsys, "lint")
    assert code == 0
    for rid in ("captain_drill_drill_myself", "miner_rockfall_inspect", "elder_fear_whispers_open_court"):
        assert f"free answer '{rid}'" not in out  # findings: free answers dominated (L11: elder open court)
    assert "free answer" not in out
    for tag in ("feast_promised", "forge_half_promise", "promised_harvest", "promised_pay", "inspected_face"):
        assert f"memory '{tag}' is set but no trigger reads it" not in out  # every promise comes due
    _, model, report = content.load_and_validate(REAL)
    assert not report.errors
    assert any(t["weights"]["courage"] < 0 for t in model.traits.values())  # someone dislikes courage
    assert any((a["rep"]["fear"] or 0) < 0 for a in model.actions)  # setbacks lower fear
    mercy = [d for d in model.decisions.values() if (d["rep"]["fear"] or 0) < 0 and d["farm_key"] == "mercy"]
    assert mercy  # a repeatable fear-lowering answer
    assert "KeepUpgraded" not in {(a["event"] or "").split(":")[0] for a in model.actions}
    for a in model.actions:  # construction reputation only through a measured need
        if (a["event"] or "").startswith(content.CONSTRUCTION_EVENTS) and any(v > 0 for v in a["rep"].values()):
            assert a["condition"], a


def test_real_content_risk_notices_describe_only_the_outcome():
    """FORGE-00 ruling #2(a): the panel shows the chance, so no notice repeats it ('Risk (50%): ...')."""
    _, model, report = content.load_and_validate(REAL)
    assert not report.errors
    risks = [c for c in model.consequences.values() if c["chance"] < 1]
    assert len(risks) >= 3
    for c in model.consequences.values():
        assert c["notice"] is None or "{chance}" not in c["notice"], c["id"]
    for c in risks:
        assert c["notice"] and not c["notice"].lower().startswith("risk"), (c["id"], c["notice"])
    assert model.consequences["speaker_walkout_risk"]["notice"] == "{name} may walk off the job for {duration}."


def test_real_content_tuning_rulings():
    """FORGE-00 ruling #7: raid_defended decays by 0.5 per repeat; phase_fed_* by 0.8."""
    _, model, report = content.load_and_validate(REAL)
    assert not report.errors
    by_id: dict[str, list[dict]] = {}
    for a in model.actions:
        by_id.setdefault(a["id"], []).append(a)
    assert {a["repeat_decay"] for a in by_id["raid_defended"]} == {0.5}
    phase = [a for aid, rows in by_id.items() if aid.startswith("phase_fed_") for a in rows]
    assert len(phase) == 3 and {a["repeat_decay"] for a in phase} == {0.8}


def test_spec_and_readme_follow_the_rulings():
    """Doc rulings #2, #5, #6, #7, #11 of the FORGE-05 review of PR #67 (wording pinned loosely)."""
    spec = (REAL / "SPEC.md").read_text(encoding="utf-8")
    readme = (REAL / "README.md").read_text(encoding="utf-8")
    assert "Open for FORGE-00: loyalty consequences" not in spec  # #5: ruled in §6.3
    assert "`{ id, text, cost?, risks? }`" in spec and "choiceRisks" in spec  # #6, #2
    assert "PhaseEndedWellFed` fires at most once per phase end" in spec  # #7
    assert "decayed as one value" in spec  # #11
    assert "creditedRulers" in spec  # #3
    # follow-ups: the cost paid comes from decisionEffects, choicePreview needs fillNames, the bound needs a full window
    assert "pay the cost (`choicePreview(...).cost`)" not in spec and "`decisionEffects(response.decision).cost`" in spec
    assert "**`fillNames` is required**" in spec
    assert "forever adds nothing" not in spec and "The bound needs a full window" in spec
    assert "default-ignorable" in spec
    assert "must show its odds with `{chance}`" not in readme and "never write `{chance}`" in readme
    assert "the numbers `{chance}`" not in readme


def test_real_content_player_text_is_13_plus_friendly():
    """Content review: no execution wording in player-facing text (the mechanics keep DeserterExecuted)."""
    tables, _, report = content.load_and_validate(REAL)
    assert not report.errors
    for spec in content.SPECS:
        texts = [c.name for c in spec.columns if c.kind == "text"]
        for row in tables[spec.key].rows:
            for col in texts:
                assert "execut" not in row.get(col).lower(), f"{spec.path}:{row.num} [{col}]: {row.get(col)}"


def test_real_content_rejects_a_strike_by_the_captains(tmp_path, capsys):
    lib = tmp_path / "content"
    for spec in content.SPECS:
        (lib / spec.path).parent.mkdir(parents=True, exist_ok=True)
        shutil.copyfile(REAL / spec.path, lib / spec.path)
    rows = read_rows(lib / "dialogue" / "consequences.csv")
    row = next(i for i, r in enumerate(rows, 1) if r[0] == "speaker_refuses_work")
    set_cell(lib, "dialogue/consequences.csv", row, "kind", "strike")
    set_cell(lib, "dialogue/consequences.csv", row, "magnitude", "0.3")
    set_cell(lib, "dialogue/consequences.csv", row, "target", "occupation:captain")
    code, out = run(capsys, "validate", "--content", lib)
    assert code == 1
    assert f"dialogue/consequences.csv:{row}:3: [target] 'captain' are soldiers" in out


def test_lint_and_contentquery_share_the_loyalty_clamp():
    """Lint scores answers with ContentQuery's formula, so the two copies of the M3 clamp must agree."""
    source = (REAL_OUT / "ContentQuery.luau").read_text(encoding="utf-8")
    m = re.search(r"ContentQuery\.LOYALTY_CLAMP = table\.freeze\(\{([^}]*)\}\)", source)
    assert m, "ContentQuery.LOYALTY_CLAMP not found"
    luau = {k.strip(): int(v) for k, v in (item.split("=") for item in m.group(1).split(","))}
    assert luau == content.LOYALTY_CLAMP


def test_cli_entry_point_runs():
    res = subprocess.run([sys.executable, str(ROOT / "tools" / "content" / "content.py"), "validate", "--content",
                          str(FIXTURE)], capture_output=True, text=True)
    assert res.returncode == 0, res.stdout + res.stderr
    assert "— 0 errors" in res.stdout
