#!/usr/bin/env python3
"""content — AI Kingdoms (game-001) content library: validate, lint, build Luau, round-trip a workbook.

The CSV files under games/game-001/content are the canonical source (schema: content/SPEC.md,
editor guide: content/README.md). This tool turns them into generated Luau modules under
games/game-001/src/shared/Content, so the game never parses text at runtime.

Commands (all take --content DIR; build/check also take --out DIR)
  validate                 check every CSV; prints 'N scenarios, ... — 0 errors' or file:row:column errors (exit 1)
  lint                     design warnings (never fails)
  build                    validate, then write the generated Luau (deterministic)
  check                    exit 1 if validation fails or the generated Luau is missing or stale (CI)
  export-xlsx [PATH]       write the editing workbook (default content/AI_Kingdoms_Content.xlsx), stamped with
                           the CSV versions it was made from
  import-xlsx PATH [--build] [--text-only]
                           read the workbook back into the CSVs: validates first, lists every change (TEXT /
                           MECHANICS / NOTES), refuses formulas and mechanics from a workbook older than the CSVs,
                           never reverts a CSV text edited after the export, restores the original row order;
                           --text-only takes only the text. Unchanged rows keep their bytes
  fmt                      rewrite every CSV in canonical form

Standard library + openpyxl (only for the two workbook commands).
"""
from __future__ import annotations

import argparse
import csv
import difflib
import io
import math
import re
import sys
from collections import deque
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Callable

ROOT = Path(__file__).resolve().parents[2]
CONTENT_DIR = ROOT / "games" / "game-001" / "content"
OUT_DIR = ROOT / "games" / "game-001" / "src" / "shared" / "Content"
XLSX_NAME = "AI_Kingdoms_Content.xlsx"
GENERATOR = "tools/content/content.py"

# ----------------------------------------------------------------------------- vocabulary
CATEGORIES = ("compassion", "honor", "competence", "courage", "authority", "fear")
RESOURCES = ("food", "wood", "stone", "gold")  # CONTRACTS G2 order
JOBS = ("farmer", "woodcutter", "miner", "builder", "hauler", "soldier")
KEY_ROLES = ("blacksmith", "healer", "captain", "provisioner", "elder", "merchant")
BUILDING_TYPES = (
    "keep", "house", "farm", "storehouse", "lumbercamp", "quarry", "barracks", "watchtower", "wall", "gate",
    # CONTRACTS G13.5: additive G2 BuildingTypes, built from VS-2 (key citizens appear with them)
    "mine", "blacksmith", "market", "infirmary", "granary", "shrine",
)
PHASES = ("Lobby", "Founding", "Expansion", "Conflict", "Endgame", "Arena", "Results")
EVENT_TAGS = (
    "raid", "shortage", "keep_upgraded", "building_completed", "objective_captured", "soldiers_recruited",
    "strike", "riot",
)
CONSEQUENCE_KINDS = (
    "productivity", "strike", "morale", "happiness", "stability", "loyalty", "resource_over_time",
    "unlock_scenario", "memory", "desertion_risk", "unrest_pressure", "cue",
)
SIMPLE_TARGETS = ("speaker", "family", "friends", "soldiers", "civilians", "kingdom")
LOCAL_TARGETS = ("speaker", "family", "friends")  # a handful of named citizens; everything else is a group
BROAD_TARGETS = ("kingdom", "civilians", "soldiers")  # kingdom-wide groups (model decision 6 caps)
DECISION_AUDIENCES = ("none", "occupation", "kingdom")
ACTION_AUDIENCES = ("soldiers", "civilians", "kingdom", "speaker")  # plus occupation:<id>
NAME_KINDS = ("given_f", "given_m", "given_n", "family")
THOUGHT_ICONS = ("request", "concern", "idea")  # CONTRACTS G13.3 thought.icon
REACTION_TONES = ("happy", "angry", "sad", "neutral")  # CONTRACTS G13.3 reaction.tone
# Feedback cue ids for consequences.cue_id (FORGE-04 vocabulary, SPEC §5.4). 'celebrate' is the G4 cue.
CUE_IDS = ("celebrate", "grumble", "salute", "bow", "cower", "strike_placard", "work_fast", "rest", "hammer")
# TEXT placeholders (SPEC §3). Which ones a column may use depends on whether it has a speaking citizen.
CITIZEN_PLACEHOLDERS = ("name", "ruler", "kingdom", "occupation", "family", "friend", "rival")
MECHANICS_PLACEHOLDERS = ("chance", "magnitude", "magnitude_pct", "duration")  # consequences.notice_text only
PLACEHOLDERS = CITIZEN_PLACEHOLDERS + MECHANICS_PLACEHOLDERS
RELATION_FLAGS = {"family": "citizen.has_family", "friend": "citizen.has_friend", "rival": "citizen.has_rival"}
SYSTEM_MEMORIES = ("rulerVisited", "raided")
SUGGESTED_NEEDS = ("food", "shelter", "rest", "safety", "tools", "health", "pay", "faith", "company", "respect")
# CONTRACTS G9: the game events on the framework bus (Contracts/GameEvents). Any other event name in
# actions.csv is a proposed reputation trigger: it needs a contract (FORGE-00) and FORGE-01 must fire it.
BUS_EVENTS = (
    "KingdomJoined", "MatchPhaseChanged", "BuildingPlaced", "BuildingCompleted", "BuildingUpgraded",
    "BuildingDestroyed", "ResourceGathered", "UnitsRecruited", "ObjectiveCaptured", "KingdomEliminated",
    "MatchEnded",
)
# CONTRACTS G13.4: the reputation action events FORGE-01 produces and appends to the GameEvents bus.
REPUTATION_EVENTS = (
    "FoodShortage", "RulerVisited", "ObjectiveLost", "PhaseEndedWellFed",  # VS-1
    "RaidDefended", "RaidAbandoned", "EnemyAttacked", "EnemyFarmsBurned", "TruceBroken", "DeserterExecuted",
    "DeserterForgiven",  # VS-2
)
CONTRACTED_EVENTS = BUS_EVENTS + REPUTATION_EVENTS  # lint warns on any other event name in actions.csv
CONSTRUCTION_EVENTS = ("BuildingPlaced", "BuildingCompleted", "BuildingUpgraded")
MIN_GIVEN_NAMES = 60
MIN_FAMILY_NAMES = 40

REP_RANGE = (-10, 10)
DELTA_RANGE = (-10, 10)  # loyalty / chemistry / kingdom deltas per decision or action row
CONSEQUENCE_MAGNITUDE = {  # kind -> (lo, hi) or None (ignored)
    "productivity": (-1, 1), "strike": (0, 1), "morale": (-20, 20), "happiness": (-20, 20),
    "stability": (-20, 20), "loyalty": (-20, 20), "resource_over_time": (-500, 500), "unlock_scenario": None,
    "memory": None, "desertion_risk": (0, 1), "unrest_pressure": (-20, 20), "cue": None,
}
# Model decision 6 guards: one decision may never stop or turn a whole kingdom (SPEC §5.4, §6.6).
MAX_DURATION_S = 600  # any consequence
STRIKE_MAX_SHARE = 0.5  # a strike stops at most half of one occupation ...
STRIKE_MAX_DURATION_S = 120  # ... for at most two minutes
GROUP_PRODUCTIVITY_FLOOR = -0.5  # productivity on any group target (not speaker/family/friends)
KINGDOM_LOYALTY_FLOOR = -5  # per decision: kingdom audience_loyalty + loyalty consequences on broad targets
KINGDOM_STABILITY_FLOOR = -5  # per decision: stability column + stability consequences
# Runtime clamp of one decision's applied loyalty change per listener (ContentQuery.LOYALTY_CLAMP, decision M3).
LOYALTY_CLAMP = {"speaker": 25, "family": 10, "friends": 10, "rivals": 10, "kingdom": 5}
# Characters that are invisible, reorder text or break a workbook (openpyxl refuses C0 controls).
_INVISIBLE_NAMES = {0x2028: "line separator", 0x2029: "paragraph separator", 0x200B: "zero-width space",
                    0x200C: "zero-width non-joiner", 0x200D: "zero-width joiner", 0x2060: "word joiner",
                    0xFEFF: "zero-width no-break space", 0x7F: "delete", 0x09: "tab", 0x0A: "line break",
                    0x0D: "carriage return", 0x0B: "vertical tab (Shift+Enter in Word)",
                    0x00AD: "soft hyphen",
                    # bidirectional controls: they reorder how the text is shown (review finding L2)
                    0x200E: "left-to-right mark", 0x200F: "right-to-left mark",
                    0x202A: "left-to-right embedding", 0x202B: "right-to-left embedding",
                    0x202C: "pop directional formatting", 0x202D: "left-to-right override",
                    0x202E: "right-to-left override", 0x2066: "left-to-right isolate",
                    0x2067: "right-to-left isolate", 0x2068: "first strong isolate",
                    0x2069: "pop directional isolate"}
# Player-facing text may not start with these: a spreadsheet would run the cell as a formula when the CSV is
# opened (CSV injection, review finding L1). Notes are not checked: they never reach the game, and they often
# start with '+' or '-' ("+housing capacity").
FORMULA_STARTS = ("=", "+", "-", "@")

ID_RE = re.compile(r"^[a-z][a-z0-9_]*$")
TAG_RE = re.compile(r"^[A-Za-z][A-Za-z0-9_]*$")
ID_MAX = 40
LUAU_KEYWORDS = frozenset(
    "and break do else elseif end false for function if in local nil not or repeat return then true until while".split()
)


# ----------------------------------------------------------------------------- schema
@dataclass(frozen=True)
class Column:
    name: str
    kind: str  # "key" | "mech" | "text" | "style" (presentation choice, edited with the text) | "notes"
    help: str
    limit: int | None = None
    required: bool = True  # must exist in the header
    choices: str | None = None  # workbook dropdown list (see LIST_SOURCES)


@dataclass(frozen=True)
class TableSpec:
    key: str  # also the workbook sheet name
    path: str  # relative to the content directory
    columns: tuple[Column, ...]
    noun: str  # for the summary line

    def col(self, name: str) -> Column | None:
        for c in self.columns:
            if c.name == name:
                return c
        return None


def _rep_cols(prefix: str, what: str) -> tuple[Column, ...]:
    return tuple(Column(f"{prefix}{c}", "mech", f"{what} {c}") for c in CATEGORIES)


SPECS: tuple[TableSpec, ...] = (
    TableSpec("scenarios", "dialogue/scenarios.csv", (
        Column("scenario_id", "key", "unique id of the situation"),
        Column("occupation", "mech", "occupation id that can raise it, or 'any'", choices="occupation_or_any"),
        Column("trigger", "mech", "condition expression (SPEC §4); empty = always"),
        Column("priority", "mech", "1-100; higher is offered first"),
        Column("cooldown_s", "mech", "seconds before the same citizen raises it again"),
        Column("once_per_citizen", "mech", "true/false", choices="bools"),
        Column("thought_icon", "style", "bubble icon: request | concern | idea (default request)", required=False,
               choices="thought_icons"),
        Column("thought_text", "text", "thought bubble above the citizen", 48),
        Column("question_text", "text", "what the citizen says when selected", 160),
        Column("notes", "notes", "designer notes (not in game)"),
    ), "scenarios"),
    TableSpec("responses", "dialogue/responses.csv", (
        Column("scenario_id", "key", "scenario this answer belongs to", choices="scenario_ids"),
        Column("response_id", "key", "unique id of the answer"),
        Column("order", "mech", "1-4 position in the menu"),
        Column("response_text", "text", "the ruler's answer", 70),
        Column("reaction_text", "text", "the citizen's reaction (speech bubble)", 80),
        Column("reaction_tone", "style", "how the reaction is shown: happy | angry | sad | neutral (default neutral)",
               required=False, choices="reaction_tones"),
        Column("decision_id", "mech", "the mechanics this answer applies (decisions sheet)", choices="decision_ids"),
        Column("notes", "notes", "designer notes (not in game)", required=False),
    ), "responses"),
    TableSpec("decisions", "dialogue/decisions.csv", (
        Column("decision_id", "key", "unique id of the decision"),
        Column("farm_key", "mech", "anti-farming group (default: decision_id)"),
        Column("repeat_decay", "mech", "(0,1] gain multiplier per repeat in the window"),
        Column("repeat_window_s", "mech", "seconds a repeat counts"),
        Column("cost", "mech", "paid now, e.g. gold:20;food:10"),
        Column("resource_effects", "mech", "signed resource changes, e.g. food:-20;gold:5"),
        *_rep_cols("rep_", "ruler reputation change, -10..+10:"),
        Column("speaker_loyalty", "mech", "-10..+10 loyalty of the citizen spoken to"),
        Column("speaker_chemistry", "mech", "-10..+10 chemistry change with the citizen spoken to"),
        Column("family_loyalty", "mech", "-10..+10 loyalty of the speaker's family"),
        Column("friends_loyalty", "mech", "-10..+10 loyalty of the speaker's friends"),
        Column("rivals_loyalty", "mech", "-10..+10 loyalty of the speaker's rivals"),
        Column("audience", "mech", "none | occupation | kingdom", choices="decision_audiences"),
        Column("audience_loyalty", "mech", "-10..+10 loyalty of the audience"),
        Column("happiness", "mech", "-10..+10 kingdom happiness"),
        Column("stability", "mech", "-10..+10 kingdom stability"),
        Column("morale", "mech", "-10..+10 army morale"),
        Column("trait_overrides", "mech", "extra loyalty by trait, e.g. greedy:+3;kind:-2"),
        Column("consequences", "mech", "consequence ids separated by ;"),
        Column("memory_tag", "mech", "memory stored on the speaker (usable in triggers)"),
        Column("notes", "notes", "designer notes (not in game)"),
    ), "decisions"),
    TableSpec("consequences", "dialogue/consequences.csv", (
        Column("consequence_id", "key", "unique id"),
        Column("kind", "mech", "what happens (SPEC §5.4)", choices="consequence_kinds"),
        Column("target", "mech", "who it affects", choices="targets"),
        Column("magnitude", "mech", "size; meaning depends on kind"),
        Column("param", "mech", "resource / scenario id / memory tag, per kind", required=False),
        Column("duration_s", "mech", "seconds it lasts (0 = instant)"),
        Column("delay_s", "mech", "seconds before it starts"),
        Column("chance", "mech", "(0,1], default 1; below 1 is shown to the player as a risk"),
        Column("cue_id", "mech", "feedback cue id (FORGE-04)"),
        Column("notice_text", "text", "notice shown to the player", 100),
        Column("notes", "notes", "designer notes (not in game)"),
    ), "consequences"),
    TableSpec("traits", "citizens/traits.csv", (
        Column("trait_id", "key", "unique id"),
        Column("name_text", "text", "trait name", 24),
        Column("description_text", "text", "one-line description", 100),
        Column("opposite", "mech", "trait that can never be combined with this one", choices="trait_ids"),
        *_rep_cols("weight_", "-2..+2 how much this trait likes the ruler gaining"),
        Column("obedience_bias", "mech", "-2..+2 how strongly authority/fear make them comply"),
        Column("productivity", "mech", "0.8-1.2 work-rate multiplier"),
        Column("hunger_sensitivity", "mech", "0.5-1.5"),
        Column("gossip", "mech", "0-2 how far news spreads from them"),
        Column("volatility", "mech", "0.5-1.5 how fast their mood swings"),
        Column("notes", "notes", "designer notes (not in game)", required=False),
    ), "traits"),
    TableSpec("names", "citizens/names.csv", (
        Column("kind", "mech", "given_f | given_m | given_n | family", choices="name_kinds"),
        Column("name", "text", "a name", 16),
        Column("notes", "notes", "designer notes (not in game)", required=False),
    ), "names"),
    TableSpec("occupations", "citizens/occupations.csv", (
        Column("occupation_id", "key", "unique id"),
        Column("name_text", "text", "occupation name", 24),
        Column("job_id", "mech", "simulation job (CONTRACTS G2)", choices="jobs"),
        Column("key_role", "mech", "Crown Arena key role, or empty", choices="key_roles"),
        Column("needs", "mech", "personal needs, most important first, e.g. food;tools"),
        Column("trait_affinity", "mech", "trait likelihood multipliers, e.g. brave:2;timid:0.5"),
        Column("notes", "notes", "designer notes (not in game)"),
    ), "occupations"),
    TableSpec("categories", "reputation/categories.csv", (
        Column("category_id", "key", "one of the six reputation categories"),
        Column("name_text", "text", "category name", 24),
        Column("description_text", "text", "what it means", 160),
        Column("affects_text", "text", "what raises or lowers it, in plain words", 300),
        Column("start", "mech", "0-100 starting value"),
        Column("notes", "notes", "designer notes (not in game)", required=False),
    ), "categories"),
    TableSpec("actions", "reputation/actions.csv", (
        Column("action_id", "key", "ruler action id (rows may repeat it, one per audience)"),
        Column("event", "mech", "game event that fires it, e.g. BuildingCompleted:barracks"),
        Column("condition", "mech", "kingdom condition (the measured need for construction)"),
        Column("direct", "notes", "what the mechanic already does (informational)"),
        Column("farm_key", "mech", "anti-farming group (default: action_id)"),
        Column("repeat_decay", "mech", "(0,1] gain multiplier per repeat in the window"),
        Column("repeat_window_s", "mech", "seconds a repeat counts"),
        *_rep_cols("rep_", "ruler reputation change, -10..+10:"),
        Column("audience", "mech", "who reacts", choices="action_audiences"),
        Column("audience_condition", "mech", "condition on each listener (citizen variables allowed)"),
        Column("audience_loyalty", "mech", "-10..+10 loyalty of matching listeners"),
        Column("morale", "mech", "-10..+10 army morale"),
        Column("happiness", "mech", "-10..+10 kingdom happiness"),
        Column("stability", "mech", "-10..+10 kingdom stability"),
        Column("notice_text", "text", "history-log explanation", 100),
        Column("notes", "notes", "designer notes (not in game)"),
    ), "actions"),
)
SPEC_BY_KEY = {s.key: s for s in SPECS}
KEY_COLUMN = {
    "scenarios": "scenario_id", "responses": "response_id", "decisions": "decision_id",
    "consequences": "consequence_id", "traits": "trait_id", "occupations": "occupation_id",
    "categories": "category_id",
}


# ----------------------------------------------------------------------------- CSV I/O (RFC 4180, raw-preserving)
class CsvSyntaxError(Exception):
    def __init__(self, row: int, message: str):
        super().__init__(message)
        self.row = row


@dataclass
class Record:
    fields: list[str]
    raw: str  # exact source text of the record, without its line terminator
    row: int  # 1-based record number (header = 1)


@dataclass
class CsvDoc:
    records: list[Record]
    newline: str = "\n"
    bom: bool = False
    trailing_newline: bool = True


def parse_csv(text: str) -> CsvDoc:
    bom = text.startswith("﻿")
    if bom:
        text = text[1:]
    newline = "\r\n" if re.search(r"\r\n", text.split("\n", 1)[0] + "\n") else "\n"
    records: list[Record] = []
    i, n, rownum = 0, len(text), 0
    while i < n:
        start = i
        rownum += 1
        fields: list[str] = []
        while True:
            if i < n and text[i] == '"':
                i += 1
                buf: list[str] = []
                while True:
                    if i >= n:
                        raise CsvSyntaxError(rownum, "a quoted cell is never closed (missing closing \")")
                    c = text[i]
                    if c == '"':
                        if i + 1 < n and text[i + 1] == '"':
                            buf.append('"')
                            i += 2
                            continue
                        i += 1
                        break
                    buf.append(c)
                    i += 1
                if i < n and text[i] not in ",\r\n":
                    raise CsvSyntaxError(rownum, "unexpected text after a closing quote (inner quotes must be doubled: \"\")")
                fields.append("".join(buf))
            else:
                j = i
                while j < n and text[j] not in ",\r\n":
                    j += 1
                fields.append(text[i:j])
                i = j
            if i < n and text[i] == ",":
                i += 1
                continue
            break
        raw = text[start:i]
        if i < n and text[i] == "\r":
            i += 1
        if i < n and text[i] == "\n":
            i += 1
        records.append(Record(fields, raw, rownum))
    trailing = text.endswith("\n") or text.endswith("\r")
    return CsvDoc(records, newline, bom, trailing or not text)


def serialize_row(fields: list[str]) -> str:
    buf = io.StringIO()
    csv.writer(buf, lineterminator="", quoting=csv.QUOTE_MINIMAL).writerow(fields)
    return buf.getvalue()


def render_csv(rows: list[list[str]], old: CsvDoc | None) -> str:
    """Rows -> CSV text. Rows identical to a row of `old` keep their exact original bytes."""
    pool: dict[tuple[str, ...], deque[str]] = {}
    newline, bom, trailing = "\n", False, True
    if old is not None:
        newline, bom, trailing = old.newline, old.bom, old.trailing_newline
        for rec in old.records:
            if rec.fields == [""] and rec.raw == "":
                continue
            pool.setdefault(tuple(rec.fields), deque()).append(rec.raw)
    lines = []
    for fields in rows:
        same = pool.get(tuple(fields))
        lines.append(same.popleft() if same else serialize_row(fields))
    text = newline.join(lines) + (newline if trailing and lines else "")
    return ("﻿" if bom else "") + text


# ----------------------------------------------------------------------------- in-memory tables and diagnostics
@dataclass
class Row:
    num: int  # spreadsheet row (header = 1)
    values: dict[str, str]

    def get(self, col: str) -> str:
        return self.values.get(col, "")


@dataclass
class Table:
    spec: TableSpec
    source: str  # display path (csv) or workbook file name (sheet)
    header: list[str]
    rows: list[Row]
    sheet: str | None = None  # set when the table came from a workbook
    usable: bool = True  # False when the header is broken (rows are not validated)


@dataclass
class Diag:
    level: str  # "error" | "warning"
    source: str
    message: str
    table: Table | None = None
    row: int | None = None
    column: str | None = None

    def render(self) -> str:
        t = self.table
        colnum = None
        if t is not None and self.column is not None and self.column in t.header:
            colnum = t.header.index(self.column) + 1
        tag = f" [{self.column}]" if self.column else ""
        if t is not None and t.sheet is not None:
            where = f"{self.source} › {t.sheet}"
            if self.row is not None:
                letter = _col_letter(colnum) if colnum else ""
                where += f"!{letter}{self.row}"
            return f"{where}{tag}: {self.message}"
        loc = self.source
        if self.row is not None:
            loc += f":{self.row}"
            if colnum is not None:
                loc += f":{colnum}"
        return f"{loc}:{tag} {self.message}" if tag else f"{loc}: {self.message}"


def _col_letter(n: int) -> str:
    s = ""
    while n > 0:
        n, r = divmod(n - 1, 26)
        s = chr(65 + r) + s
    return s


class Report:
    def __init__(self) -> None:
        self.items: list[Diag] = []

    def add(self, level: str, table: Table | None, row: int | None, column: str | None, message: str,
            source: str | None = None) -> None:
        src = source if source is not None else (table.source if table is not None else "content")
        self.items.append(Diag(level, src, message, table, row, column))

    def error(self, table: Table | None, row: int | None, column: str | None, message: str) -> None:
        self.add("error", table, row, column, message)

    def warn(self, table: Table | None, row: int | None, column: str | None, message: str) -> None:
        self.add("warning", table, row, column, message)

    @property
    def errors(self) -> list[Diag]:
        return [d for d in self.items if d.level == "error"]

    @property
    def warnings(self) -> list[Diag]:
        return [d for d in self.items if d.level == "warning"]


def display_path(path: Path) -> str:
    try:
        return path.resolve().relative_to(ROOT).as_posix()
    except ValueError:
        return path.resolve().as_posix()


def table_from_records(spec: TableSpec, source: str, records: list[tuple[int, list[str]]], report: Report,
                       sheet: str | None = None) -> Table:
    """records: (row number, cells) with the header first."""
    table = Table(spec, source, [], [], sheet)
    if not records:
        report.error(table, None, None, "the file is empty (it needs a header row)")
        table.usable = False
        return table
    header_row, header = records[0]
    header = [h.strip() for h in header]
    while header and header[-1] == "" and sheet is not None:
        header.pop()
    table.header = header
    known = {c.name for c in spec.columns}
    seen: set[str] = set()
    for h in header:
        if h == "":
            report.error(table, header_row, None, "a header cell is empty")
            table.usable = False
        elif h not in known:
            hint = _suggest(h, known)
            report.error(table, header_row, h, f"unknown column '{h}'{hint} (allowed: {', '.join(c.name for c in spec.columns)})")
            table.usable = False
        elif h in seen:
            report.error(table, header_row, h, f"column '{h}' appears twice")
            table.usable = False
        seen.add(h)
    for c in spec.columns:
        if c.required and c.name not in seen:
            report.error(table, header_row, None, f"missing column '{c.name}'")
            table.usable = False
    for num, cells in records[1:]:
        if all(v.strip() == "" for v in cells):
            continue
        if len(cells) > len(header) and any(v.strip() for v in cells[len(header):]):
            report.error(table, num, None, f"this row has {len(cells)} cells but the header has {len(header)} "
                                           "(an unquoted comma? wrap the cell in double quotes)")
        values = {h: (cells[k] if k < len(cells) else "") for k, h in enumerate(header)}
        table.rows.append(Row(num, values))
    return table


def load_csv_table(spec: TableSpec, content_dir: Path, report: Report) -> Table | None:
    path = content_dir / spec.path
    source = display_path(path)
    if not path.is_file():
        report.add("error", None, None, None, f"missing file (expected {spec.path})", source=source)
        return None
    data = path.read_bytes()
    try:
        doc = parse_csv(data.decode("utf-8"))
    except UnicodeDecodeError as e:
        row, col = _csv_position(data[:e.start].decode("utf-8"))
        report.add("error", None, row, None, f"not UTF-8 text: byte 0x{data[e.start]:02X} at byte offset {e.start} "
                                             f"(column {col}) is not valid UTF-8; save the CSV as 'CSV UTF-8'",
                   source=source)
        return None
    except CsvSyntaxError as e:
        report.add("error", None, e.row, None, str(e), source=source)
        return None
    return table_from_records(spec, source, [(r.row, r.fields) for r in doc.records], report)


def _csv_position(prefix: str) -> tuple[int, int]:
    """(spreadsheet row, 1-based column) where `prefix` ends: record breaks and commas outside quotes count."""
    row, col, quoted = 1, 1, False
    i, n = 0, len(prefix)
    while i < n:
        c = prefix[i]
        if c == '"':
            quoted = not quoted
        elif not quoted and c == ",":
            col += 1
        elif not quoted and c in "\r\n":
            if c == "\r" and i + 1 < n and prefix[i + 1] == "\n":
                i += 1
            row, col = row + 1, 1
        i += 1
    return row, col


def load_tables(content_dir: Path, report: Report, overrides: dict[str, Table] | None = None) -> dict[str, Table]:
    tables: dict[str, Table] = {}
    for spec in SPECS:
        if overrides and spec.key in overrides:
            tables[spec.key] = overrides[spec.key]
            continue
        t = load_csv_table(spec, content_dir, report)
        if t is not None:
            tables[spec.key] = t
    return tables


def _suggest(name: str, candidates) -> str:
    """' (did you mean 'x'?)' for the closest candidate; compares case-insensitively, answers in its spelling."""
    by_lower: dict[str, str] = {}
    for c in candidates:
        by_lower.setdefault(str(c).lower(), str(c))
    match = difflib.get_close_matches(name.lower(), list(by_lower), n=1, cutoff=0.7)
    return f" (did you mean '{by_lower[match[0]]}'?)" if match else ""


# ----------------------------------------------------------------------------- condition language
@dataclass(frozen=True)
class VarSpec:
    type: str  # "number" | "bool" | "enum"
    scope: str  # "kingdom" | "citizen"
    lo: float | None = None
    hi: float | None = None
    choices: tuple[str, ...] = ()


_N = VarSpec("number", "kingdom", 0, None)
STATIC_VARS: dict[str, VarSpec] = {
    "phase": VarSpec("enum", "kingdom", choices=PHASES),
    **{f"kingdom.{r}": _N for r in RESOURCES},
    "kingdom.fed_ratio": VarSpec("number", "kingdom", 0, 1),
    "kingdom.housing_ratio": VarSpec("number", "kingdom", 0, 10),
    "kingdom.happiness": VarSpec("number", "kingdom", 0, 100),
    "kingdom.stability": VarSpec("number", "kingdom", 0, 100),
    "kingdom.morale": VarSpec("number", "kingdom", 0, 100),
    "kingdom.population": _N,
    "kingdom.soldiers": _N,
    "kingdom.level.keep": VarSpec("number", "kingdom", 0, 10),
    **{f"kingdom.count.{b}": _N for b in BUILDING_TYPES},
    **{f"rep.{c}": VarSpec("number", "kingdom", 0, 100) for c in CATEGORIES},
    "citizen.loyalty": VarSpec("number", "citizen", 0, 100),
    "citizen.chemistry": VarSpec("number", "citizen", -20, 20),
    "citizen.mood": VarSpec("number", "citizen", 0, 100),
    "citizen.hunger": VarSpec("number", "citizen", 0, 1),
    "citizen.obedience": VarSpec("number", "citizen", 0, 100),
    "citizen.has_family": VarSpec("bool", "citizen"),
    "citizen.has_friend": VarSpec("bool", "citizen"),
    "citizen.has_rival": VarSpec("bool", "citizen"),
    **{f"event.{e}": VarSpec("bool", "kingdom") for e in EVENT_TAGS},
}
VAR_TOKEN = r"[A-Za-z_][A-Za-z0-9_.]*"
RE_BARE = re.compile(rf"(!?)\s*({VAR_TOKEN})")
RE_IN = re.compile(rf"({VAR_TOKEN})\s+in\s+(.+)")
RE_CMP = re.compile(rf"({VAR_TOKEN})\s*(<=|>=|!=|=|<|>)\s*([^<>=!\s].*)")


class CondError(Exception):
    pass


@dataclass
class Vocab:
    occupations: set[str]
    traits: set[str]
    decisions: set[str]
    check_occupations: bool = True
    check_traits: bool = True
    check_decisions: bool = True

    def resolve(self, name: str) -> VarSpec:
        if name in STATIC_VARS:
            return STATIC_VARS[name]
        if name == "occupation":
            return VarSpec("enum", "citizen", choices=tuple(sorted(self.occupations)))
        for prefix, ids, check, what in (
            ("citizen.trait.", self.traits, self.check_traits, "trait"),
            ("decision.", self.decisions, self.check_decisions, "decision"),
        ):
            if name.startswith(prefix):
                ident = name[len(prefix):]
                if check and ident not in ids:
                    raise CondError(f"unknown {what} '{ident}' in '{name}'{_suggest(ident, ids)}")
                return VarSpec("bool", "citizen" if what == "trait" else "kingdom")
        if name.startswith("citizen.memory."):
            tag = name[len("citizen.memory."):]
            if not TAG_RE.match(tag):
                raise CondError(f"bad memory tag '{tag}' (letters, digits and _ only)")
            return VarSpec("bool", "citizen")
        raise CondError(f"unknown variable '{name}'{_suggest(name, self.all_names())}")

    def all_names(self) -> list[str]:
        names = list(STATIC_VARS) + ["occupation"]
        names += [f"citizen.trait.{t}" for t in self.traits] + [f"decision.{d}" for d in self.decisions]
        return names


# ASCII decimal only (review finding L3): float() alone also takes '1_0', Arabic-Indic and full-width digits.
_NUM_RE = re.compile(r"[+-]?(?:[0-9]+(?:\.[0-9]*)?|\.[0-9]+)(?:[eE][+-]?[0-9]+)?")


def _num(text: str) -> float:
    if not _NUM_RE.fullmatch(text.strip(" ")):
        raise ValueError(text)
    v = float(text)
    if math.isnan(v) or math.isinf(v):
        raise ValueError(text)
    return v


def _clean_number(v: float) -> int | float:
    return int(v) if float(v).is_integer() else v


def parse_condition(text: str, vocab: Vocab, scope: str) -> tuple[list[dict], list[str]]:
    """Returns (clauses, memory tags used). Raises CondError with a precise message."""
    text = text.strip()
    if text == "":
        return [], []
    if re.search(r"\|\||\||\s(or|OR|Or)\s", text):
        raise CondError("only AND is supported (join clauses with ' & '); for OR, write two rows")
    if re.search(r"\s(and|AND|And)\s", text):
        raise CondError("join clauses with ' & ' (not 'and'), e.g. 'kingdom.food < 50 & citizen.hunger > 0.5'")
    if "&&" in text:
        raise CondError("use a single '&' between clauses")
    clauses: list[dict] = []
    memories: list[str] = []
    for k, part in enumerate(text.split("&"), 1):
        s = part.strip()
        where = f"clause {k} '{s}'"
        if s == "":
            raise CondError(f"clause {k} is empty (a doubled or trailing '&')")
        if "==" in s:
            raise CondError(f"{where}: use '=' instead of '=='")
        for wrong, right in (("=<", "<="), ("=>", ">="), ("≤", "<="), ("≥", ">="), ("≠", "!="), ("<>", "!=")):
            if wrong in s:
                raise CondError(f"{where}: write '{right}' instead of '{wrong}'")
        m_not = re.fullmatch(rf"(?:not|NOT|Not)\s+({VAR_TOKEN})", s)
        if m_not:
            raise CondError(f"{where}: write '!{m_not.group(1)}' instead of 'not {m_not.group(1)}'")
        m_in = RE_IN.fullmatch(s)
        m_cmp = None if m_in else RE_CMP.fullmatch(s)
        m_bare = None if (m_in or m_cmp) else RE_BARE.fullmatch(s)
        if not (m_in or m_cmp or m_bare):
            raise CondError(f"{where}: malformed; expected '<variable> <op> <value>' (op = != < <= > >=), "
                            "'<variable> in a,b', '<variable>' or '!<variable>'")
        var = (m_in or m_cmp or m_bare).group(2 if m_bare else 1)
        try:
            spec = vocab.resolve(var)
        except CondError as e:
            raise CondError(f"{where}: {e}") from None
        if spec.scope == "citizen" and scope == "kingdom":
            raise CondError(f"{where}: '{var}' describes a citizen, but this condition is checked for the kingdom")
        if var.startswith("citizen.memory."):
            memories.append(var[len("citizen.memory."):])
        if m_bare:
            if spec.type != "bool":
                raise CondError(f"{where}: '{var}' is a {spec.type}; compare it, e.g. '{var} {'>=' if spec.type == 'number' else '='} …'")
            clauses.append({"var": var, "op": "=", "value": m_bare.group(1) != "!"})
            continue
        if m_in:
            if spec.type != "enum":
                raise CondError(f"{where}: 'in' only works with {', '.join(['occupation', 'phase'])}")
            items = [x.strip() for x in m_in.group(2).split(",")]
            if any(x == "" for x in items):
                raise CondError(f"{where}: empty item in the 'in' list")
            for x in items:
                _check_enum(where, var, x, spec)
            if len(set(items)) != len(items):
                raise CondError(f"{where}: an item appears twice in the 'in' list")
            clauses.append({"var": var, "op": "in", "value": items})
            continue
        op, raw = m_cmp.group(2), m_cmp.group(3).strip()
        if spec.type == "bool":
            if op not in ("=", "!=") or raw.lower() not in ("true", "false"):
                raise CondError(f"{where}: '{var}' is true/false; write '{var}' or '!{var}'")
            clauses.append({"var": var, "op": "=", "value": (raw.lower() == "true") == (op == "=")})
        elif spec.type == "enum":
            if op not in ("=", "!="):
                raise CondError(f"{where}: '{var}' is a name; use = != or in")
            _check_enum(where, var, raw, spec)
            clauses.append({"var": var, "op": op, "value": raw})
        else:
            try:
                value = _num(raw)
            except ValueError:
                raise CondError(f"{where}: '{raw}' is not a number") from None
            if (spec.lo is not None and value < spec.lo) or (spec.hi is not None and value > spec.hi):
                rng = f"{_fmt(spec.lo)}..{_fmt(spec.hi) if spec.hi is not None else ''}"
                raise CondError(f"{where}: {_fmt(value)} is outside the range of '{var}' ({rng})")
            clauses.append({"var": var, "op": op, "value": _clean_number(value)})
    return clauses, memories


def _check_enum(where: str, var: str, value: str, spec: VarSpec) -> None:
    if var == "occupation" and not spec.choices:
        return  # occupations table unusable: skip the check
    if value not in spec.choices:
        raise CondError(f"{where}: unknown {var} '{value}'{_suggest(value, spec.choices)}")


def _fmt(v: float | None) -> str:
    if v is None:
        return ""
    return str(_clean_number(v))


# ----------------------------------------------------------------------------- validation -> model
@dataclass
class Model:
    scenarios: dict[str, dict] = field(default_factory=dict)
    scenario_order: list[str] = field(default_factory=list)
    responses: dict[str, list[dict]] = field(default_factory=dict)
    decisions: dict[str, dict] = field(default_factory=dict)
    consequences: dict[str, dict] = field(default_factory=dict)
    traits: dict[str, dict] = field(default_factory=dict)
    names: dict[str, list[str]] = field(default_factory=dict)
    occupations: dict[str, dict] = field(default_factory=dict)
    categories: dict[str, dict] = field(default_factory=dict)
    actions: list[dict] = field(default_factory=list)
    # bookkeeping for lint
    rows: dict[tuple[str, str], tuple[Table, Row]] = field(default_factory=dict)
    action_rows: list[tuple[Table, Row]] = field(default_factory=list)
    memory_uses: list[tuple[Table, int, str, str]] = field(default_factory=list)
    key_role_rows: dict[str, list[tuple[Table, Row]]] = field(default_factory=dict)
    referenced_decisions: set[str] = field(default_factory=set)
    referenced_consequences: set[str] = field(default_factory=set)
    counts: dict[str, int] = field(default_factory=dict)


class Validator:
    def __init__(self, tables: dict[str, Table], report: Report):
        self.t = tables
        self.r = report
        self.m = Model()
        self.ids: dict[str, set[str]] = {}
        self.checkable: dict[str, bool] = {}

    # --- cell helpers -------------------------------------------------------
    def err(self, table: Table, row: Row, col: str, msg: str) -> None:
        self.r.error(table, row.num, col, msg)

    def raw(self, row: Row, col: str) -> str:
        return row.get(col).strip()

    def p_id(self, table: Table, row: Row, col: str, required: bool = True, pattern=ID_RE) -> str | None:
        v = self.raw(row, col)
        if v == "":
            if required:
                self.err(table, row, col, "required")
            return None
        if not pattern.match(v) or len(v) > ID_MAX:
            self.err(table, row, col, f"'{v}' is not a valid id (lower-case letters, digits and _, starting with a letter, "
                                      f"at most {ID_MAX} characters)" if pattern is ID_RE else
                     f"'{v}' is not a valid tag (letters, digits and _, starting with a letter)")
            return None
        return v

    def p_int(self, table: Table, row: Row, col: str, lo: float | None, hi: float | None, default: int | None = 0):
        v = self.raw(row, col)
        if v == "":
            if default is None:
                self.err(table, row, col, "required")
            return default
        try:
            f = _num(v)
        except ValueError:
            self.err(table, row, col, f"'{v}' is not a whole number")
            return None
        if not f.is_integer():
            self.err(table, row, col, f"'{v}' must be a whole number")
            return None
        n = int(f)
        if (lo is not None and n < lo) or (hi is not None and n > hi):
            self.err(table, row, col, f"{n} is out of range ({_fmt(lo)}..{_fmt(hi)})")
            return None
        return n

    def p_num(self, table: Table, row: Row, col: str, lo: float | None, hi: float | None, default: float | None = 0,
              lo_open: bool = False):
        v = self.raw(row, col)
        if v == "":
            if default is None:
                self.err(table, row, col, "required")
            return default
        try:
            f = _num(v)
        except ValueError:
            self.err(table, row, col, f"'{v}' is not a number")
            return None
        if (lo is not None and (f < lo or (lo_open and f == lo))) or (hi is not None and f > hi):
            rng = f"{'>' if lo_open else ''}{_fmt(lo)}..{_fmt(hi)}"
            self.err(table, row, col, f"{_fmt(f)} is out of range ({rng})")
            return None
        return _clean_number(f)

    def p_bool(self, table: Table, row: Row, col: str, default: bool = False):
        v = self.raw(row, col).lower()
        if v == "":
            return default
        if v not in ("true", "false"):
            self.err(table, row, col, f"'{self.raw(row, col)}' must be true or false")
            return None
        return v == "true"

    def p_enum(self, table: Table, row: Row, col: str, choices, default: str | None = None):
        v = self.raw(row, col)
        if v == "":
            if default is None:
                self.err(table, row, col, "required")
            return default
        if v not in choices:
            self.err(table, row, col, f"'{v}' is not allowed{_suggest(v, choices)} (allowed: {', '.join(choices)})")
            return None
        return v

    def p_text(self, table: Table, row: Row, col: str, required: bool = True,
               placeholders: tuple[str, ...] = ()) -> str | None:
        """Player-facing text. `placeholders` lists the ones this column may use (SPEC §3); () = none."""
        spec = table.spec.col(col)
        v = row.get(col)
        if v.strip() == "":
            if required:
                self.err(table, row, col, "text is required")
            return None
        ok = True
        if v != v.strip():
            self.err(table, row, col, "remove the spaces at the start or end of the text")
            ok = False
        if "\n" in v or "\r" in v:
            self.err(table, row, col, "line breaks are not allowed in game text")
            ok = False
        if spec is not None and spec.limit is not None and len(v) > spec.limit:
            self.err(table, row, col, f"text is {len(v)} characters; the limit is {spec.limit}")
            ok = False
        if v[:1] in FORMULA_STARTS:
            self.err(table, row, col, f"text may not start with '{v[0]}': a spreadsheet would run it as a formula "
                                      "(CSV injection); reword it so it starts with a letter, a digit or a quote mark")
            ok = False
        for problem in _placeholder_problems(v, placeholders):
            self.err(table, row, col, problem)
            ok = False
        if _invisible_problem(v, "text"):
            ok = False  # reported once by v_cells
        return v if ok else None

    def relation_guard(self, table: Table, row: Row, cols: tuple[str, ...], trigger: list[dict] | None,
                       sid: str) -> None:
        """{family} / {friend} / {rival} only when the scenario's trigger guarantees that relation (SPEC §3)."""
        if trigger is None:
            return
        have = {c["var"] for c in trigger if c["op"] == "=" and c["value"] is True}
        for col in cols:
            for m in _PLACEHOLDER.finditer(row.get(col)):
                flag = RELATION_FLAGS.get(m.group(1))
                if flag and flag not in have:
                    self.err(table, row, col, f"'{{{m.group(1)}}}' needs a {m.group(1)}: add '{flag}' to the "
                                              f"trigger of scenario '{sid}', or the name could be missing in game")

    def p_list(self, table: Table, row: Row, col: str) -> list[str]:
        v = self.raw(row, col)
        return [x.strip() for x in v.split(";") if x.strip()] if v else []

    def p_kv(self, table: Table, row: Row, col: str, check_key: Callable[[str], str | None],
             parse_val: Callable[[str], Any], what: str) -> dict | None:
        out: dict[str, Any] = {}
        ok = True
        for item in self.p_list(table, row, col):
            if ":" not in item:
                self.err(table, row, col, f"'{item}' must be written key:value (e.g. {what})")
                ok = False
                continue
            k, val = (x.strip() for x in item.split(":", 1))
            problem = check_key(k)
            if problem:
                self.err(table, row, col, problem)
                ok = False
                continue
            if k in out:
                self.err(table, row, col, f"'{k}' is listed twice")
                ok = False
                continue
            try:
                out[k] = parse_val(val)
            except ValueError as e:
                self.err(table, row, col, f"'{item}': {e}")
                ok = False
        return out if ok else None

    def cond(self, table: Table, row: Row, col: str, scope: str) -> list[dict] | None:
        try:
            clauses, memories = parse_condition(row.get(col), self.vocab, scope)
        except CondError as e:
            self.err(table, row, col, str(e))
            return None
        for tag in memories:
            self.m.memory_uses.append((table, row.num, col, tag))
        return clauses

    # --- passes -------------------------------------------------------------
    def v_cells(self) -> None:
        """Invisible and control characters, in every cell of every table (they break the workbook)."""
        for t in self.t.values():
            if not t.usable:
                continue
            for row in t.rows:
                for col in t.header:
                    c = t.spec.col(col)
                    problem = _invisible_problem(row.values.get(col, ""), c.kind if c else "mech")
                    if problem:
                        self.err(t, row, col, problem)

    def run(self) -> Model:
        self.unlock_targets: set[str] = set()
        self.v_cells()
        for key, col in KEY_COLUMN.items():
            self.collect_ids(key, col)
        self.vocab = Vocab(
            self.ids.get("occupations", set()), self.ids.get("traits", set()), self.ids.get("decisions", set()),
            self.checkable.get("occupations", False), self.checkable.get("traits", False),
            self.checkable.get("decisions", False),
        )
        if not self.checkable.get("occupations", False):
            self.vocab.occupations = set()
        self.v_categories()
        self.v_traits()
        self.v_occupations()
        self.v_names()
        self.v_consequences()
        self.v_decisions()
        self.v_scenarios()
        self.v_responses()
        self.v_actions()
        self.m.counts = {
            "scenarios": len(self.rows_of("scenarios")), "responses": len(self.rows_of("responses")),
            "decisions": len(self.rows_of("decisions")), "consequences": len(self.rows_of("consequences")),
            "traits": len(self.rows_of("traits")), "names": len(self.rows_of("names")),
            "occupations": len(self.rows_of("occupations")), "categories": len(self.rows_of("categories")),
            "actions": len(self.rows_of("actions")),
        }
        return self.m

    def table(self, key: str) -> Table | None:
        t = self.t.get(key)
        return t if t is not None and t.usable else None

    def rows_of(self, key: str) -> list[Row]:
        t = self.table(key)
        return t.rows if t else []

    def collect_ids(self, key: str, col: str) -> None:
        t = self.table(key)
        self.ids[key] = set()
        self.checkable[key] = t is not None
        if t is None:
            return
        first: dict[str, int] = {}
        for row in t.rows:
            v = self.p_id(t, row, col)
            if v is None:
                continue
            if v in first:
                self.err(t, row, col, f"duplicate id '{v}' (first used on row {first[v]})")
                continue
            first[v] = row.num
            self.ids[key].add(v)
            self.m.rows[(key, v)] = (t, row)
        if key == "occupations" and "any" in self.ids[key]:
            row = self.m.rows[(key, "any")][1]
            self.err(t, row, col, "'any' is reserved (scenarios use it to mean every occupation)")

    def ref(self, table: Table, row: Row, col: str, target: str, value: str) -> bool:
        if not self.checkable.get(target, False):
            return True
        if value not in self.ids[target]:
            noun = target[:-1] if target.endswith("s") else target
            self.err(table, row, col, f"unknown {noun} '{value}'{_suggest(value, self.ids[target])}")
            return False
        return True

    def keyed_rows(self, key: str):
        """Rows whose id is valid and first-seen (duplicates and bad ids were reported by collect_ids)."""
        t = self.table(key)
        if t is None:
            return
        col = KEY_COLUMN[key]
        for row in t.rows:
            ident = self.raw(row, col)
            if (key, ident) in self.m.rows and self.m.rows[(key, ident)][1] is row:
                yield t, row, ident

    # --- tables -------------------------------------------------------------
    def v_categories(self) -> None:
        t = self.table("categories")
        if t is None:
            return
        for tt, row, cid in self.keyed_rows("categories"):
            if cid not in CATEGORIES:
                self.err(tt, row, "category_id", f"'{cid}' is not a reputation category (allowed: {', '.join(CATEGORIES)})")
                continue
            rec = {
                "id": cid,
                "name": self.p_text(tt, row, "name_text"),
                "description": self.p_text(tt, row, "description_text"),
                "affects": self.p_text(tt, row, "affects_text"),
                "start": self.p_int(tt, row, "start", 0, 100, None),
            }
            self.m.categories[cid] = rec
        missing = [c for c in CATEGORIES if c not in self.ids["categories"]]
        if missing:
            self.r.error(t, None, None, f"missing reputation categories: {', '.join(missing)} (exactly six rows)")

    def v_traits(self) -> None:
        t = self.table("traits")
        if t is None:
            return
        opposites: dict[str, tuple[Row, str | None]] = {}
        for tt, row, tid in self.keyed_rows("traits"):
            opp = self.p_id(tt, row, "opposite", required=False)
            if opp is not None:
                if opp == tid:
                    self.err(tt, row, "opposite", "a trait cannot be its own opposite")
                    opp = None
                elif not self.ref(tt, row, "opposite", "traits", opp):
                    opp = None
            opposites[tid] = (row, opp)
            rec = {
                "id": tid,
                "name": self.p_text(tt, row, "name_text"),
                "description": self.p_text(tt, row, "description_text"),
                "opposite": opp,
                "weights": {c: self.p_num(tt, row, f"weight_{c}", -2, 2) for c in CATEGORIES},
                "obedience_bias": self.p_num(tt, row, "obedience_bias", -2, 2, 0),
                "productivity": self.p_num(tt, row, "productivity", 0.8, 1.2, 1),
                "hunger_sensitivity": self.p_num(tt, row, "hunger_sensitivity", 0.5, 1.5, 1),
                "gossip": self.p_num(tt, row, "gossip", 0, 2, 1),
                "volatility": self.p_num(tt, row, "volatility", 0.5, 1.5, 1),
            }
            self.m.traits[tid] = rec
        for tid, (row, opp) in opposites.items():
            if opp is None or opp not in opposites:
                continue
            back = opposites[opp][1]
            if back != tid:
                self.err(t, row, "opposite", f"opposites must be symmetric: '{opp}' lists "
                                             f"{repr(back) if back else 'no opposite'}, expected '{tid}'")

    def v_occupations(self) -> None:
        roles: dict[str, list[tuple[Table, Row]]] = {}
        for tt, row, oid in self.keyed_rows("occupations"):
            key_role = self.p_enum(tt, row, "key_role", KEY_ROLES, default="")
            needs = []
            for need in self.p_list(tt, row, "needs"):
                if not ID_RE.match(need):
                    self.err(tt, row, "needs", f"'{need}' is not a valid need id (e.g. food;shelter;tools)")
                elif need in needs:
                    self.err(tt, row, "needs", f"'{need}' is listed twice")
                else:
                    needs.append(need)

            def trait_key(k: str) -> str | None:
                if self.checkable["traits"] and k not in self.ids["traits"]:
                    return f"unknown trait '{k}'{_suggest(k, self.ids['traits'])}"
                return None

            def weight(v: str):
                f = _num_or_raise(v)
                if f < 0 or f > 5:
                    raise ValueError("weight must be 0..5 (1 = normal, 2 = twice as likely, 0 = never)")
                return _clean_number(f)

            affinity = self.p_kv(tt, row, "trait_affinity", trait_key, weight, "brave:2;timid:0.5")
            rec = {
                "id": oid,
                "name": self.p_text(tt, row, "name_text"),
                "job": self.p_enum(tt, row, "job_id", JOBS),
                "key_role": key_role or None,
                "needs": needs,
                "trait_affinity": dict(sorted((affinity or {}).items())),
            }
            if key_role:
                roles.setdefault(key_role, []).append((tt, row))
            self.m.occupations[oid] = rec
        self.m.key_role_rows = roles

    def v_names(self) -> None:
        t = self.table("names")
        if t is None:
            return
        names: dict[str, list[str]] = {k: [] for k in NAME_KINDS}
        first: dict[tuple[str, str], int] = {}
        for row in t.rows:
            kind = self.p_enum(t, row, "kind", NAME_KINDS)
            name = self.p_text(t, row, "name")
            if kind is None or name is None:
                continue
            bad = [ch for ch in name if not (ch.isalpha() or ch in "'- ")]
            if bad:
                self.err(t, row, "name", f"names use letters, apostrophes, hyphens and spaces only (found '{bad[0]}')")
                continue
            if (kind, name.lower()) in first:
                self.err(t, row, "name", f"'{name}' is already listed as {kind} on row {first[(kind, name.lower())]}")
                continue
            first[(kind, name.lower())] = row.num
            names[kind].append(name)
        given = sum(len(names[k]) for k in ("given_f", "given_m", "given_n"))
        if given < MIN_GIVEN_NAMES:
            self.r.error(t, None, None, f"only {given} given names; at least {MIN_GIVEN_NAMES} are needed")
        if len(names["family"]) < MIN_FAMILY_NAMES:
            self.r.error(t, None, None, f"only {len(names['family'])} family names; at least {MIN_FAMILY_NAMES} are needed")
        self.m.names = names

    def target(self, table: Table, row: Row, col: str, simple: tuple[str, ...], required: bool = True) -> str | None:
        v = self.raw(row, col)
        if v == "":
            if required:
                self.err(table, row, col, "required")
            return None
        if v.startswith("occupation:"):
            occ = v.split(":", 1)[1].strip()
            if self.ref(table, row, col, "occupations", occ):
                return f"occupation:{occ}"
            return None
        if v not in simple:
            allowed = ", ".join((*simple, "occupation:<id>"))
            self.err(table, row, col, f"'{v}' is not allowed{_suggest(v, simple)} (allowed: {allowed})")
            return None
        return v

    def v_consequences(self) -> None:
        for tt, row, cid in self.keyed_rows("consequences"):
            kind = self.p_enum(tt, row, "kind", CONSEQUENCE_KINDS)
            target = self.target(tt, row, "target", SIMPLE_TARGETS)
            rng = CONSEQUENCE_MAGNITUDE.get(kind or "", None)
            magnitude = self.p_num(tt, row, "magnitude", rng[0] if rng else None, rng[1] if rng else None, 0)
            duration = self.p_int(tt, row, "duration_s", 0, None, 0)
            if duration is not None and duration > MAX_DURATION_S:
                self.err(tt, row, "duration_s", f"{duration} s is longer than the {MAX_DURATION_S} s limit (lasting "
                                                "kingdom trouble comes from the sustained unrest rules, model "
                                                "decision 6, never from one decision)")
                duration = None
            delay = self.p_int(tt, row, "delay_s", 0, None, 0)
            chance = self.p_num(tt, row, "chance", 0, 1, 1, lo_open=True)
            cue = self.p_id(tt, row, "cue_id", required=False)
            notice = self.p_text(tt, row, "notice_text", required=False,
                                 placeholders=("name", "ruler", "kingdom", "occupation") + MECHANICS_PLACEHOLDERS)
            param_raw = self.raw(row, "param")
            param: str | None = param_raw or None
            if kind == "resource_over_time":
                if param_raw not in RESOURCES:
                    self.err(tt, row, "param", f"resource_over_time needs param = a resource ({', '.join(RESOURCES)})")
            elif kind == "unlock_scenario":
                if param_raw == "":
                    self.err(tt, row, "param", "unlock_scenario needs param = the scenario id to unlock")
                elif self.ref(tt, row, "param", "scenarios", param_raw):
                    self.unlock_targets.add(param_raw)
            elif kind == "memory":
                param = param_raw or cid
                if not TAG_RE.match(param):
                    self.err(tt, row, "param", f"'{param}' is not a valid memory tag")
            elif kind == "cue":
                if cue is None and self.raw(row, "cue_id") == "":
                    self.err(tt, row, "cue_id", "kind 'cue' needs a cue_id")
            elif kind is not None and param_raw:
                self.err(tt, row, "param", f"kind '{kind}' takes no param (leave it empty)")
            if kind in ("productivity", "strike", "resource_over_time") and duration == 0:
                self.err(tt, row, "duration_s", f"kind '{kind}' needs duration_s > 0")
            if kind == "strike" and target is not None:
                self.strike_rules(tt, row, target, magnitude, duration)
            if (kind == "productivity" and target is not None and target not in LOCAL_TARGETS
                    and magnitude is not None and magnitude < GROUP_PRODUCTIVITY_FLOOR):
                self.err(tt, row, "magnitude", f"{_fmt(magnitude)} would stop a whole group: productivity on "
                                               f"'{target}' may fall at most to {_fmt(GROUP_PRODUCTIVITY_FLOOR)} "
                                               "(model decision 6; a full walk-out is only for speaker/family/friends)")
            if chance is not None and chance < 1:
                if self.raw(row, "notice_text") == "":
                    self.err(tt, row, "notice_text", f"chance {_fmt(chance)} is below 1, so the player must see the "
                                                     "risk before choosing (model decision 7): write a notice_text "
                                                     "with {chance}")
                elif "{chance}" not in row.get("notice_text"):
                    self.err(tt, row, "notice_text", "a risk must show its real odds: write {chance} in the text "
                                                     "(the game fills in e.g. 50%) instead of a typed number")
            rec = {
                "id": cid, "kind": kind, "target": target, "magnitude": magnitude, "duration_s": duration,
                "delay_s": delay, "chance": chance, "param": param, "cue_id": cue, "notice": notice,
            }
            self.m.consequences[cid] = rec

    def strike_rules(self, tt: Table, row: Row, target: str, magnitude, duration) -> None:
        """Model decision 6 + CONTRACTS G13.3 (`strike { occupation, on }`): a decision may only start a short,
        partial walk-out of one civilian occupation."""
        if not target.startswith("occupation:"):
            if target in BROAD_TARGETS:
                msg = ("a single decision may not stop a whole kingdom, all civilians or all soldiers (model decision "
                       "6); a strike targets one occupation:<id>, or use unrest_pressure")
            else:
                msg = ("a strike stops part of one occupation (CONTRACTS G13.3 sends strike { occupation }); for one "
                       "citizen's walk-out use kind productivity with magnitude -1")
            self.err(tt, row, "target", msg)
            return
        occ = target.split(":", 1)[1]
        if (self.m.occupations.get(occ) or {}).get("job") == "soldier":
            self.err(tt, row, "target", f"'{occ}' are soldiers: one decision may not make the army stop (model "
                                        "decision 6); use morale or desertion_risk")
        if magnitude is not None and magnitude > STRIKE_MAX_SHARE:
            self.err(tt, row, "magnitude", f"a strike from one decision stops at most {_fmt(STRIKE_MAX_SHARE)} of the "
                                           f"occupation (got {_fmt(magnitude)}; model decision 6)")
        if duration is not None and duration > STRIKE_MAX_DURATION_S:
            self.err(tt, row, "duration_s", f"a strike from one decision lasts at most {STRIKE_MAX_DURATION_S} s "
                                            f"(got {duration}; model decision 6)")

    def anti_farm(self, table: Table, row: Row) -> tuple[str | None, Any, Any]:
        farm = self.p_id(table, row, "farm_key", required=False)
        decay = self.p_num(table, row, "repeat_decay", 0, 1, 1, lo_open=True)
        window = self.p_int(table, row, "repeat_window_s", 0, None, 0)
        return farm, decay, window

    def resources(self, table: Table, row: Row, col: str, positive: bool) -> dict | None:
        def key(k: str) -> str | None:
            return None if k in RESOURCES else f"unknown resource '{k}'{_suggest(k, RESOURCES)} (allowed: {', '.join(RESOURCES)})"

        def val(v: str):
            f = _num_or_raise(v)
            if not f.is_integer():
                raise ValueError("amounts are whole numbers")
            if positive and f <= 0:
                raise ValueError("cost amounts are the positive amount paid (use resource_effects for gains)")
            if f == 0:
                raise ValueError("0 has no effect; remove the item")
            return int(f)

        bag = self.p_kv(table, row, col, key, val, "food:-20;gold:5" if not positive else "gold:20;food:10")
        if bag is None:
            return None
        return {r: bag[r] for r in RESOURCES if r in bag}

    def v_decisions(self) -> None:
        for tt, row, did in self.keyed_rows("decisions"):
            farm, decay, window = self.anti_farm(tt, row)
            audience = self.p_enum(tt, row, "audience", DECISION_AUDIENCES, default="none")
            audience_loyalty = self.p_int(tt, row, "audience_loyalty", *DELTA_RANGE)
            if audience == "none" and audience_loyalty:
                self.err(tt, row, "audience_loyalty", "audience_loyalty needs an audience (occupation or kingdom)")

            def trait_key(k: str) -> str | None:
                if self.checkable["traits"] and k not in self.ids["traits"]:
                    return f"unknown trait '{k}'{_suggest(k, self.ids['traits'])}"
                return None

            def delta(v: str):
                f = _num_or_raise(v)
                if not f.is_integer() or not DELTA_RANGE[0] <= f <= DELTA_RANGE[1]:
                    raise ValueError(f"must be a whole number {DELTA_RANGE[0]}..+{DELTA_RANGE[1]}")
                return int(f)

            overrides = self.p_kv(tt, row, "trait_overrides", trait_key, delta, "greedy:+3;kind:-2")
            consequences = []
            for cid in self.p_list(tt, row, "consequences"):
                if self.ref(tt, row, "consequences", "consequences", cid):
                    if cid in consequences:
                        self.err(tt, row, "consequences", f"'{cid}' is listed twice")
                    else:
                        consequences.append(cid)
                        self.m.referenced_consequences.add(cid)
            memory = self.p_id(tt, row, "memory_tag", required=False, pattern=TAG_RE)
            rec = {
                "id": did,
                "farm_key": farm or did,
                "repeat_decay": decay,
                "repeat_window_s": window,
                "cost": self.resources(tt, row, "cost", True) or {},
                "resources": self.resources(tt, row, "resource_effects", False) or {},
                "rep": {c: self.p_int(tt, row, f"rep_{c}", *REP_RANGE) for c in CATEGORIES},
                "speaker": {
                    "loyalty": self.p_int(tt, row, "speaker_loyalty", *DELTA_RANGE),
                    "chemistry": self.p_int(tt, row, "speaker_chemistry", *DELTA_RANGE),
                },
                "family_loyalty": self.p_int(tt, row, "family_loyalty", *DELTA_RANGE),
                "friends_loyalty": self.p_int(tt, row, "friends_loyalty", *DELTA_RANGE),
                "rivals_loyalty": self.p_int(tt, row, "rivals_loyalty", *DELTA_RANGE),
                "audience": audience,
                "audience_loyalty": audience_loyalty,
                "kingdom": {
                    "happiness": self.p_int(tt, row, "happiness", *DELTA_RANGE),
                    "stability": self.p_int(tt, row, "stability", *DELTA_RANGE),
                    "morale": self.p_int(tt, row, "morale", *DELTA_RANGE),
                },
                "trait_overrides": dict(sorted((overrides or {}).items())),
                "consequences": consequences,
                "memory_tag": memory,
                "_farm_missing": farm is None,
            }
            self.kingdom_caps(tt, row, rec)
            self.m.decisions[did] = rec

    def kingdom_caps(self, tt: Table, row: Row, d: dict) -> None:
        """Model decision 6: one answer may not push the whole kingdom towards rebellion (SPEC §6.6)."""
        loyalty = (d["audience_loyalty"] or 0) if d["audience"] == "kingdom" else 0
        stability = d["kingdom"]["stability"] or 0
        from_consequences = {"loyalty": False, "stability": False}
        for cid in d["consequences"]:
            c = self.m.consequences.get(cid)
            if c is None or c["magnitude"] is None:
                continue
            if c["kind"] == "loyalty" and c["target"] in BROAD_TARGETS:
                loyalty += c["magnitude"]
                from_consequences["loyalty"] = True
            elif c["kind"] == "stability":
                stability += c["magnitude"]
                from_consequences["stability"] = True
        if loyalty < KINGDOM_LOYALTY_FLOOR:
            col = "consequences" if from_consequences["loyalty"] else "audience_loyalty"
            self.err(tt, row, col, f"kingdom-wide loyalty {_fmt(loyalty)} is below {KINGDOM_LOYALTY_FLOOR} "
                                   "(audience_loyalty with audience kingdom, plus loyalty consequences on kingdom, "
                                   "civilians or soldiers): one decision may not turn the kingdom (model decision 6)")
        if stability < KINGDOM_STABILITY_FLOOR:
            col = "consequences" if from_consequences["stability"] else "stability"
            self.err(tt, row, col, f"stability {_fmt(stability)} is below {KINGDOM_STABILITY_FLOOR} (the stability "
                                   "column plus stability consequences): one decision may not push the kingdom "
                                   "towards rebellion (model decision 6)")

    def v_scenarios(self) -> None:
        for tt, row, sid in self.keyed_rows("scenarios"):
            occ = self.raw(row, "occupation")
            if occ == "":
                self.err(tt, row, "occupation", "required (an occupation id, or 'any')")
                occ_v = None
            elif occ == "any" or self.ref(tt, row, "occupation", "occupations", occ):
                occ_v = occ
            else:
                occ_v = None
            trigger = self.cond(tt, row, "trigger", "citizen")
            rec = {
                "id": sid,
                "occupation": occ_v,
                "trigger": trigger,
                "priority": self.p_int(tt, row, "priority", 1, 100, None),
                "cooldown_s": self.p_int(tt, row, "cooldown_s", 0, None, 0),
                "once_per_citizen": self.p_bool(tt, row, "once_per_citizen", False),
                "locked": sid in self.unlock_targets,  # offered only after an unlock_scenario consequence
                "icon": self.p_enum(tt, row, "thought_icon", THOUGHT_ICONS, default="request"),
                "thought": self.p_text(tt, row, "thought_text", placeholders=CITIZEN_PLACEHOLDERS),
                "question": self.p_text(tt, row, "question_text", placeholders=CITIZEN_PLACEHOLDERS),
            }
            self.relation_guard(tt, row, ("thought_text", "question_text"), trigger, sid)
            self.m.scenarios[sid] = rec
            self.m.scenario_order.append(sid)

    def v_responses(self) -> None:
        t = self.table("responses")
        by_scenario: dict[str, list[tuple[Row, dict]]] = {}
        for tt, row, rid in self.keyed_rows("responses"):
            sid = self.raw(row, "scenario_id")
            if sid == "":
                self.err(tt, row, "scenario_id", "required")
                continue
            if not self.ref(tt, row, "scenario_id", "scenarios", sid):
                continue
            did = self.raw(row, "decision_id")
            if did == "":
                self.err(tt, row, "decision_id", "required")
            elif self.ref(tt, row, "decision_id", "decisions", did):
                self.m.referenced_decisions.add(did)
            rec = {
                "id": rid,
                "order": self.p_int(tt, row, "order", 1, 4, None),
                "text": self.p_text(tt, row, "response_text", placeholders=CITIZEN_PLACEHOLDERS),
                "reaction": self.p_text(tt, row, "reaction_text", placeholders=CITIZEN_PLACEHOLDERS),
                "tone": self.p_enum(tt, row, "reaction_tone", REACTION_TONES, default="neutral"),
                "decision": did or None,
            }
            scenario = self.m.scenarios.get(sid)
            if scenario is not None:
                self.relation_guard(tt, row, ("response_text", "reaction_text"), scenario["trigger"], sid)
            by_scenario.setdefault(sid, []).append((row, rec))
        for sid, items in by_scenario.items():
            seen: dict[int, int] = {}
            for row, rec in items:
                o = rec["order"]
                if o is None:
                    continue
                if o in seen:
                    self.err(t, row, "order", f"order {o} is already used in scenario '{sid}' (row {seen[o]})")
                seen[o] = row.num
            self.m.responses[sid] = [rec for _, rec in sorted(items, key=lambda x: (x[1]["order"] or 0, x[0].num))]
        st = self.table("scenarios")
        if t is not None and st is not None:
            for sid in self.m.scenario_order:
                n = len(by_scenario.get(sid, []))
                if not 2 <= n <= 4:
                    srow = self.m.rows[("scenarios", sid)][1]
                    self.r.error(st, srow.num, "scenario_id", f"scenario '{sid}' has {n} responses; it needs 2-4 "
                                                              "(responses.csv)")

    def v_actions(self) -> None:
        t = self.table("actions")
        if t is None:
            return
        groups: dict[str, dict] = {}
        for row in t.rows:
            aid = self.p_id(t, row, "action_id")
            event = self.raw(row, "event")
            event_ok = self.check_event(t, row, event)
            farm, decay, window = self.anti_farm(t, row)
            audience = self.target(t, row, "audience", ACTION_AUDIENCES, required=False)
            rec = {
                "id": aid,
                "event": event if event_ok else None,
                "condition": self.cond(t, row, "condition", "kingdom"),
                "farm_key": farm,
                "repeat_decay": decay,
                "repeat_window_s": window,
                "rep": {c: self.p_int(t, row, f"rep_{c}", *REP_RANGE) for c in CATEGORIES},
                "audience": audience,
                "audience_condition": self.cond(t, row, "audience_condition", "citizen"),
                "audience_loyalty": self.p_int(t, row, "audience_loyalty", *DELTA_RANGE),
                "kingdom": {
                    "happiness": self.p_int(t, row, "happiness", *DELTA_RANGE),
                    "stability": self.p_int(t, row, "stability", *DELTA_RANGE),
                    "morale": self.p_int(t, row, "morale", *DELTA_RANGE),
                },
                # history-log line: {ruler}/{kingdom} always; a citizen's name only when one citizen is the audience
                "notice": self.p_text(t, row, "notice_text", required=False, placeholders=(
                    ("ruler", "kingdom") + (("name", "occupation") if audience == "speaker" else ()))),
            }
            if audience is None and rec["audience_loyalty"] and self.raw(row, "audience") == "":
                self.err(t, row, "audience_loyalty", "audience_loyalty needs an audience")
            if aid is None:
                continue
            g = groups.get(aid)
            explicit = {k: self.raw(row, k) != "" for k in ("farm_key", "repeat_decay", "repeat_window_s")}
            if g is None:
                groups[aid] = {"row": row, "event": event, "farm_key": farm, "repeat_decay": decay,
                               "repeat_window_s": window, "explicit": explicit, "_farm_missing": farm is None}
            else:
                if event != g["event"]:
                    self.err(t, row, "event", f"action '{aid}' uses event '{g['event']}' on row {g['row'].num}; "
                                              "all rows of one action share the event")
                for k in ("farm_key", "repeat_decay", "repeat_window_s"):
                    if explicit[k] and g["explicit"][k] and rec[k] != g[k]:
                        self.err(t, row, k, f"differs from row {g['row'].num} of action '{aid}' (rows of one action "
                                            "share the anti-farming settings; leave it empty here)")
                    elif explicit[k] and not g["explicit"][k]:
                        self.err(t, row, k, f"set it on the first row of action '{aid}' (row {g['row'].num}) instead")
                rec["farm_key"], rec["repeat_decay"], rec["repeat_window_s"] = (
                    g["farm_key"], g["repeat_decay"], g["repeat_window_s"])
            rec["_farm_missing"] = groups[aid]["_farm_missing"]
            rec["farm_key"] = rec["farm_key"] or aid
            self.m.actions.append(rec)
            self.m.action_rows.append((t, row))

    def check_event(self, table: Table, row: Row, event: str) -> bool:
        if event == "":
            self.err(table, row, "event", "required (e.g. BuildingCompleted:barracks)")
            return False
        m = re.fullmatch(r"([A-Z][A-Za-z]*)(?::([A-Za-z_][A-Za-z0-9_]*))?", event)
        if not m:
            self.err(table, row, "event", f"'{event}' must look like EventName or EventName:arg")
            return False
        name, arg = m.group(1), m.group(2)
        if name.startswith("Building") and arg is not None and arg not in BUILDING_TYPES:
            self.err(table, row, "event", f"unknown building type '{arg}'{_suggest(arg, BUILDING_TYPES)}")
            return False
        if name == "ResourceGathered" and arg is not None and arg not in RESOURCES:
            self.err(table, row, "event", f"unknown resource '{arg}'")
            return False
        if name == "MatchPhaseChanged" and arg is not None and arg not in PHASES:
            self.err(table, row, "event", f"unknown phase '{arg}'{_suggest(arg, PHASES)}")
            return False
        return True


def _num_or_raise(v: str) -> float:
    try:
        return _num(v)
    except ValueError:
        raise ValueError(f"'{v}' is not a number") from None


_PLACEHOLDER = re.compile(r"\{([^{}]*)\}")


def _placeholder_list(allowed: tuple[str, ...]) -> str:
    if not allowed:
        return "none: this text is not about one citizen"
    return ", ".join("{" + p + "}" for p in allowed)


def _placeholder_problems(text: str, allowed: tuple[str, ...] = ()) -> list[str]:
    problems = []
    for m in _PLACEHOLDER.finditer(text):
        name = m.group(1)
        if name not in PLACEHOLDERS:
            problems.append(f"unknown placeholder '{{{name}}}'{_suggest(name, allowed or PLACEHOLDERS)} "
                            f"(allowed here: {_placeholder_list(allowed)})")
        elif name not in allowed:
            problems.append(f"placeholder '{{{name}}}' cannot be used in this column "
                            f"(allowed here: {_placeholder_list(allowed)})")
    stripped = _PLACEHOLDER.sub("", text)
    if "{" in stripped or "}" in stripped:
        problems.append("unbalanced '{' or '}' (placeholders look like {name})")
    return problems


def _invisible_problem(value: str, kind: str) -> str | None:
    """First invisible/control character in a cell, or None. Notes may hold tabs and line breaks; text reports
    line breaks itself (p_text)."""
    for i, ch in enumerate(value):
        o = ord(ch)
        if ch in "\n\r" and kind in ("notes", "text"):
            continue
        if ch == "\t" and kind == "notes":
            continue
        if o < 0x20 or o == 0x7F or 0x80 <= o <= 0x9F or o in _INVISIBLE_NAMES:
            name = _INVISIBLE_NAMES.get(o, "control character")
            return (f"invisible character U+{o:04X} ({name}) at position {i + 1}; delete it (it usually comes "
                    "from copy-paste)")
    return None


def validate_tables(tables: dict[str, Table], report: Report) -> Model:
    return Validator(tables, report).run()


def summary_line(model: Model, errors: int) -> str:
    c = model.counts
    parts = [f"{c.get(s.key, 0)} {s.noun}" for s in SPECS]
    return f"{', '.join(parts)} — {errors} error{'s' if errors != 1 else ''}"


# ----------------------------------------------------------------------------- lint
def _clause_ok(clause: dict, actual) -> bool:
    op, value = clause["op"], clause["value"]
    if isinstance(value, bool):
        return bool(actual) == value if op == "=" else bool(actual) != value
    if op == "in":
        return actual in value
    if op == "=":
        return actual == value
    if op == "!=":
        return actual != value
    if not isinstance(actual, (int, float)) or isinstance(actual, bool):
        return False
    return {"<": actual < value, "<=": actual <= value, ">": actual > value, ">=": actual >= value}.get(op, False)


def sample_population(model: Model, occupation: str, n: int, rng) -> list[tuple[list[str], int]]:
    """A stand-in for FORGE-01's generator (CONTRACTS G13.2): 2-3 traits weighted by the occupation's
    trait_affinity (unlisted = 1, 0 = never), never an opposite pair, chemistry -20..+20. Used by lint only."""
    traits = sorted(model.traits)
    affinity = (model.occupations.get(occupation) or {}).get("trait_affinity", {})
    out = []
    for _ in range(n):
        want = rng.choice((2, 3))
        chosen: list[str] = []
        while len(chosen) < want:
            pool = [t for t in traits if t not in chosen and model.traits[t]["opposite"] not in chosen
                    and affinity.get(t, 1) > 0]
            if not pool:
                break
            chosen.append(rng.choices(pool, [affinity.get(t, 1) for t in pool])[0])
        out.append((chosen, rng.randint(-20, 20)))
    return out


def speaker_loyalty(model: Model, d: dict, traits: list[str], chemistry: float) -> float:
    """ContentQuery.decisionLoyalty(d, traits, "speaker", chemistry) at multiplier 1 (model decisions 3-4),
    clamped like the runtime (LOYALTY_CLAMP)."""
    total = sum((model.traits[t]["weights"][c] or 0) * (d["rep"][c] or 0) for t in traits for c in CATEGORIES)
    total += (d["speaker"]["loyalty"] or 0) + sum(d["trait_overrides"].get(t, 0) for t in traits)
    value = total * (1 + chemistry / 40) if total > 0 else total * (1 - chemistry / 40)
    return max(-LOYALTY_CLAMP["speaker"], min(LOYALTY_CLAMP["speaker"], value))


def _scenario_sample(model: Model, sid: str, n: int = 400) -> list[tuple[list[str], int]]:
    """Seeded sample of the citizens who can raise scenario `sid` (trait, occupation and chemistry clauses)."""
    import random

    s = model.scenarios[sid]
    rng = random.Random(f"lint:{sid}")
    occupations = sorted(model.occupations) if s["occupation"] == "any" else [s["occupation"]]
    trigger = s["trigger"] or []
    occ_clauses = [c for c in trigger if c["var"] == "occupation"]
    occupations = [o for o in occupations if all(_clause_ok(c, o) for c in occ_clauses)]
    out: list[tuple[list[str], int]] = []
    for attempt in range(20):
        for occ in occupations:
            for traits, chem in sample_population(model, occ, max(1, n // max(1, len(occupations))), rng):
                ok = True
                for c in trigger:
                    if c["var"].startswith("citizen.trait."):
                        ok = ok and _clause_ok(c, c["var"][len("citizen.trait."):] in traits)
                    elif c["var"] == "citizen.chemistry":
                        ok = ok and _clause_ok(c, chem)
                if ok:
                    out.append((traits, chem))
        if len(out) >= n // 2:
            break
    return out


def _is_free(d: dict) -> bool:
    return not d["cost"] and not any(v < 0 for v in d["resources"].values())


def _rep_without_fear(d: dict) -> int:
    return sum(d["rep"][c] or 0 for c in CATEGORIES if c != "fear")


def lint(model: Model, tables: dict[str, Table], report: Report) -> None:
    rows = model.rows

    def vector(d: dict) -> list[float]:
        # Fear is neither good nor bad for the ruler (model rules), so it is left out of "better".
        net = [d["resources"].get(r, 0) - d["cost"].get(r, 0) for r in RESOURCES]
        loyal = [d["speaker"]["loyalty"], d["speaker"]["chemistry"], d["family_loyalty"], d["friends_loyalty"],
                 d["rivals_loyalty"], d["audience_loyalty"] if d["audience"] != "none" else 0]
        kingdom = [d["kingdom"][k] for k in ("happiness", "stability", "morale")]
        return [d["rep"][c] for c in CATEGORIES if c != "fear"] + loyal + kingdom + net

    st = tables.get("scenarios")
    for sid, responses in model.responses.items():
        srow = rows.get(("scenarios", sid))
        srownum = srow[1].num if srow else None
        decided = [(r, model.decisions.get(r["decision"] or "")) for r in responses]
        if any(d is None for _, d in decided) or len(decided) < 2:
            continue
        seen_dec: dict[str, str] = {}
        for r, _ in decided:
            if r["decision"] in seen_dec:
                report.warn(st, srownum, "scenario_id",
                            f"responses '{seen_dec[r['decision']]}' and '{r['id']}' of '{sid}' share decision "
                            f"'{r['decision']}' (the choice has no mechanical difference)")
            seen_dec[r["decision"]] = r["id"]
        vecs = [(r, vector(d)) for r, d in decided]
        for r, v in vecs:
            others = [w for o, w in vecs if o is not r]
            if all(all(a >= b for a, b in zip(v, w)) for w in others) and any(v != w for w in others):
                report.warn(st, srownum, "scenario_id",
                            f"response '{r['id']}' ({r['decision']}) is at least as good as every other answer of "
                            f"'{sid}' on every reputation (fear aside), loyalty, chemistry, kingdom and resource "
                            "column (no real trade-off)")
        # The personality reaction is most of a speaker's loyalty change: score each answer over a seeded sample
        # of the citizens who can raise this scenario, with ContentQuery's formula.
        sample = _scenario_sample(model, sid) if model.scenarios.get(sid) else []
        if sample:
            best = [0] * len(decided)
            for traits, chem in sample:
                scores = [speaker_loyalty(model, d, traits, chem) for _, d in decided]
                best[scores.index(max(scores))] += 1
            top_rep = max(_rep_without_fear(d) for _, d in decided)
            for k, (r, d) in enumerate(decided):
                share = best[k] / len(sample)
                if _is_free(d) and share > 0.8 and _rep_without_fear(d) >= top_rep:
                    report.warn(st, srownum, "scenario_id",
                                f"free answer '{r['id']}' ({r['decision']}) is the best speaker-loyalty choice for "
                                f"{share:.0%} of sampled citizens of '{sid}' and also has the highest reputation sum "
                                "(fear aside): the paid answers are rarely worth it")
    for did, d in model.decisions.items():
        t, row = rows[("decisions", did)]
        rep_sum = _rep_without_fear(d)
        if not d["cost"] and rep_sum > 6:
            report.warn(t, row.num, "cost", f"free decision '{did}' gives +{rep_sum} reputation in total, fear "
                                            "aside (> 6); add a cost or a downside")
        if d.get("_farm_missing"):
            report.warn(t, row.num, "farm_key", f"no farm_key: repeats of '{did}' only decay as their own group")
        if did not in model.referenced_decisions:
            report.warn(t, row.num, "decision_id", f"decision '{did}' is not used by any response")
    used_unlocks: set[str] = set()
    for d in model.decisions.values():
        for cid in d["consequences"]:
            c = model.consequences.get(cid)
            if c and c["kind"] == "unlock_scenario" and c["param"]:
                used_unlocks.add(c["param"])
    for cid, c in model.consequences.items():
        t, row = rows[("consequences", cid)]
        if cid not in model.referenced_consequences:
            report.warn(t, row.num, "consequence_id", f"consequence '{cid}' is not used by any decision")
        if c["cue_id"] and c["cue_id"] not in CUE_IDS:
            report.warn(t, row.num, "cue_id", f"cue '{c['cue_id']}' is not in the cue vocabulary "
                                              f"({', '.join(CUE_IDS)}; SPEC §5.4): FORGE-04 must add it")
        if c["notice"] and re.search(r"\d", c["notice"]):
            report.warn(t, row.num, "notice_text", "the notice types a number: write {chance}, {magnitude}, "
                                                   "{magnitude_pct} or {duration} so it stays right when the "
                                                   "mechanics change")
    for sid, s in model.scenarios.items():
        t, row = rows[("scenarios", sid)]
        if s["once_per_citizen"] is False and s["cooldown_s"] == 0:
            report.warn(t, row.num, "cooldown_s", f"scenario '{sid}' repeats with no cooldown (dialogue farming)")
        if s["locked"] and sid not in used_unlocks:
            report.warn(t, row.num, "scenario_id", f"scenario '{sid}' is locked, but no decision uses a consequence "
                                                   "that unlocks it: it can never be offered")
    farm_warned: set[str] = set()
    events: dict[str, list[int]] = {}
    event_rows: dict[str, tuple[Table, Row]] = {}
    for (t, row), a in zip(model.action_rows, model.actions):
        name = (a["event"] or "").split(":")[0]
        if a["event"] and name not in CONTRACTED_EVENTS:
            events.setdefault(name, []).append(row.num)
            event_rows.setdefault(name, (t, row))
        if name in CONSTRUCTION_EVENTS and not a["condition"] and any((a["rep"][c] or 0) > 0 for c in CATEGORIES):
            report.warn(t, row.num, "condition", "construction gives reputation without a measured need "
                                                 "(model decision 5: set a condition such as kingdom.housing_ratio > 0.9)")
        if a.get("_farm_missing") and a["id"] not in farm_warned:
            farm_warned.add(a["id"])
            report.warn(t, row.num, "farm_key", f"no farm_key on action '{a['id']}'")
        if a["notice"] and re.search(r"\d", a["notice"]):
            report.warn(t, row.num, "notice_text", "the notice types a number; numbers belong in the mechanics "
                                                   "columns, so the text cannot go stale")
    for name, nums in events.items():
        t, row = event_rows[name]
        report.warn(t, row.num, "event", f"event '{name}' (rows {', '.join(map(str, nums))}) is not a contracted event "
                                         "(CONTRACTS G9 bus or G13.4 reputation events): nothing fires it until "
                                         "FORGE-00 contracts it and FORGE-01 produces it, or point the rows at a "
                                         "contracted event (SPEC §8)")
    known_memories = set(SYSTEM_MEMORIES)
    setters: dict[str, tuple[Table, int, str]] = {}
    for did, d in model.decisions.items():
        if d["memory_tag"]:
            known_memories.add(d["memory_tag"])
            t, row = rows[("decisions", did)]
            setters.setdefault(d["memory_tag"], (t, row.num, "memory_tag"))
    for cid, c in model.consequences.items():
        if c["kind"] == "memory" and c["param"]:
            known_memories.add(c["param"])
            t, row = rows[("consequences", cid)]
            setters.setdefault(c["param"], (t, row.num, "param"))
    read = {tag for _, _, _, tag in model.memory_uses}
    for t, rownum, col, tag in model.memory_uses:
        if tag not in known_memories:
            report.warn(t, rownum, col, f"memory '{tag}' is never set by a decision memory_tag or a memory consequence")
    for tag, (t, rownum, col) in setters.items():
        if tag not in read:
            report.warn(t, rownum, col, f"memory '{tag}' is set but no trigger reads it (fine as flavour for the "
                                        "citizen card; a promise needs a follow-up scenario that reads it)")
    for role, items in model.key_role_rows.items():
        if len(items) > 1:
            for t, row in items[1:]:
                report.warn(t, row.num, "key_role", f"key role '{role}' is used by more than one occupation")
    ot = tables.get("occupations")
    for oid, o in model.occupations.items():
        odd = [n for n in o["needs"] if n not in SUGGESTED_NEEDS]
        if odd and ot is not None:
            row = rows[("occupations", oid)][1]
            report.warn(ot, row.num, "needs", f"need(s) {', '.join(odd)} not in the suggested list "
                                              f"({', '.join(SUGGESTED_NEEDS)}); the runtime must know them")


# ----------------------------------------------------------------------------- Luau generation
class Arr(list):
    """A Luau array (lists are arrays; dicts are maps)."""


LINE_LIMIT = 110  # stay under stylua column_width 120 so StyLua never re-wraps generated tables


def lua_string(s: str) -> str:
    quote = "'" if s.count('"') > s.count("'") else '"'
    out = []
    for ch in s:
        if ch == "\\":
            out.append("\\\\")
        elif ch == quote:
            out.append("\\" + ch)
        elif ch == "\n":
            out.append("\\n")
        elif ch == "\r":
            out.append("\\r")
        elif ch == "\t":
            out.append("\\t")
        elif ord(ch) < 32 or ord(ch) == 127:
            out.append(f"\\{ord(ch):03d}")
        else:
            out.append(ch)
    return quote + "".join(out) + quote


def lua_key(k: str) -> str:
    if re.fullmatch(r"[A-Za-z_][A-Za-z0-9_]*", k) and k not in LUAU_KEYWORDS:
        return k
    return f"[{lua_string(k)}]"


def lua_scalar(v: Any) -> str:
    if isinstance(v, bool):
        return "true" if v else "false"
    if isinstance(v, int):
        return str(v)
    if isinstance(v, float):
        return str(int(v)) if v.is_integer() else repr(v)
    if isinstance(v, str):
        return lua_string(v)
    raise TypeError(f"cannot emit {v!r}")


def _items(v: dict | list) -> list[tuple[str | None, Any]]:
    if isinstance(v, dict):
        return [(k, x) for k, x in v.items() if x is not None]
    return [(None, x) for x in v]


def lua_inline(v: Any) -> str:
    if isinstance(v, (dict, list)):
        items = _items(v)
        if not items:
            return "{}"
        parts = [(f"{lua_key(k)} = " if k is not None else "") + lua_inline(x) for k, x in items]
        return "{ " + ", ".join(parts) + " }"
    return lua_scalar(v)


def _width(line: str) -> int:
    tabs = len(line) - len(line.lstrip("\t"))
    return tabs * 4 + len(line.lstrip("\t").encode("utf-8"))


def lua_emit(v: Any, depth: int, prefix: str, suffix: str) -> list[str]:
    ind = "\t" * depth
    is_table = isinstance(v, (dict, list))
    force = is_table and (depth <= 1 or (isinstance(v, dict) and "id" in v)) and len(_items(v)) > 0
    one = f"{ind}{prefix}{lua_inline(v)}{suffix}"
    if not is_table or (not force and _width(one) <= LINE_LIMIT) or not _items(v):
        return [one]
    lines = [f"{ind}{prefix}{{"]
    for k, x in _items(v):
        lines += lua_emit(x, depth + 1, f"{lua_key(k)} = " if k is not None else "", ",")
    lines.append(f"{ind}}}{suffix}")
    return lines


# Every generated module is deeply frozen (review finding L6): the tables are shared by every consumer, so a
# caller that writes into one gets an error instead of silently changing the content for everybody.
FREEZE_HELPER = [
    "-- Deeply frozen: the content is shared by every consumer, so nobody may change it.",
    "local function deepFreeze<T>(value: T): T",
    "\tfor _, child in value :: any do",
    '\t\tif type(child) == "table" and not table.isfrozen(child) then',
    "\t\t\tdeepFreeze(child)",
    "\t\tend",
    "\tend",
    "\treturn table.freeze(value :: any) :: any",
    "end",
    "",
]


def luau_module(data: Any, sources: list[str], content_label: str) -> str:
    header = [
        "--!strict",
        f"-- GENERATED by {GENERATOR} from {content_label} — do not edit",
        f"-- Source: {', '.join(sources)}. Edit the CSV files (or the workbook, see content/README.md),",
        f"-- then run: python3 {GENERATOR} build",
        "",
    ]
    body = lua_emit(data, 0, "return deepFreeze(", ")")
    return "\n".join(header + FREEZE_HELPER + body) + "\n"


def _clean(rec: dict) -> dict:
    return {k: v for k, v in rec.items() if not k.startswith("_")}


def generate(model: Model, content_label: str) -> dict[str, str]:
    files: dict[str, str] = {}
    scenarios = {}
    for sid in sorted(model.scenarios):
        s = dict(model.scenarios[sid])
        s["trigger"] = Arr(s["trigger"] or [])
        s["responses"] = Arr(dict(r) for r in model.responses.get(sid, []))
        scenarios[sid] = s
    files["Dialogue.luau"] = luau_module(
        {"scenarios": scenarios, "order": Arr(model.scenario_order)},
        ["dialogue/scenarios.csv", "dialogue/responses.csv"], content_label)

    decisions = {}
    for did in sorted(model.decisions):
        d = _clean(model.decisions[did])
        d["consequences"] = Arr(d["consequences"])
        decisions[did] = d
    files["Decisions.luau"] = luau_module(decisions, ["dialogue/decisions.csv"], content_label)

    files["Consequences.luau"] = luau_module(
        {cid: model.consequences[cid] for cid in sorted(model.consequences)}, ["dialogue/consequences.csv"],
        content_label)
    files["Traits.luau"] = luau_module(
        {tid: model.traits[tid] for tid in sorted(model.traits)}, ["citizens/traits.csv"], content_label)
    files["Names.luau"] = luau_module(
        {k: Arr(model.names.get(k, [])) for k in NAME_KINDS}, ["citizens/names.csv"], content_label)
    occupations = {}
    for oid in sorted(model.occupations):
        o = dict(model.occupations[oid])
        o["needs"] = Arr(o["needs"])
        occupations[oid] = o
    files["Occupations.luau"] = luau_module(occupations, ["citizens/occupations.csv"], content_label)

    actions = []
    by_event: dict[str, list[int]] = {}
    for i, a in enumerate(model.actions, 1):
        a = _clean(a)
        a["condition"] = Arr(a["condition"] or [])
        a["audience_condition"] = Arr(a["audience_condition"] or [])
        actions.append(a)
        by_event.setdefault(a["event"], []).append(i)
    files["Reputation.luau"] = luau_module(
        {
            "categories": {c: model.categories[c] for c in CATEGORIES if c in model.categories},
            "order": Arr(CATEGORIES),
            "actions": Arr(actions),
            "byEvent": {e: Arr(by_event[e]) for e in sorted(by_event)},
        },
        ["reputation/categories.csv", "reputation/actions.csv"], content_label)
    return files


# ----------------------------------------------------------------------------- commands
def load_and_validate(content_dir: Path, overrides: dict[str, Table] | None = None) -> tuple[dict, Model, Report]:
    report = Report()
    tables = load_tables(content_dir, report, overrides)
    model = validate_tables(tables, report)
    return tables, model, report


def print_diags(diags: list[Diag], out=None) -> None:
    out = out or sys.stdout
    for d in diags:
        prefix = "" if d.level == "error" else "warning: "
        print(prefix + d.render(), file=out)


def content_label(content_dir: Path) -> str:
    return display_path(content_dir)


def cmd_validate(args) -> int:
    _, model, report = load_and_validate(args.content)
    print_diags(report.errors)
    print(summary_line(model, len(report.errors)))
    return 1 if report.errors else 0


def cmd_lint(args) -> int:
    tables, model, report = load_and_validate(args.content)
    if report.errors:
        print_diags(report.errors)
        print(f"lint: {len(report.errors)} validation error(s) first; fix them, then lint again")
        return 0
    lint(model, tables, report)
    print_diags(report.warnings)
    print(f"lint: {len(report.warnings)} warning{'s' if len(report.warnings) != 1 else ''}")
    return 0


def _built(args) -> tuple[dict[str, str] | None, Report, Model]:
    _, model, report = load_and_validate(args.content)
    if report.errors:
        return None, report, model
    return generate(model, content_label(args.content)), report, model


def cmd_build(args) -> int:
    files, report, model = _built(args)
    if files is None:
        print_diags(report.errors)
        print(summary_line(model, len(report.errors)))
        print("build: not written (fix the errors above)")
        return 1
    out: Path = args.out
    out.mkdir(parents=True, exist_ok=True)
    changed = []
    for name, text in files.items():
        path = out / name
        if not path.is_file() or path.read_bytes() != text.encode("utf-8"):
            path.write_text(text, encoding="utf-8", newline="\n")
            changed.append(name)
    print(summary_line(model, 0))
    print(f"build: {len(files)} modules in {display_path(out)} ({len(changed)} changed"
          f"{': ' + ', '.join(changed) if changed else ''})")
    return 0


def cmd_check(args) -> int:
    files, report, model = _built(args)
    if files is None:
        print_diags(report.errors)
        print(summary_line(model, len(report.errors)))
        return 1
    stale = []
    for name, text in files.items():
        path = args.out / name
        if not path.is_file():
            stale.append(f"{name} (missing)")
        elif path.read_bytes() != text.encode("utf-8"):
            stale.append(f"{name} (differs)")
    if stale:
        print(f"check: generated Luau is out of date in {display_path(args.out)}: {', '.join(stale)}")
        print(f"check: run python3 {GENERATOR} build")
        return 1
    print(summary_line(model, 0))
    print(f"check: {len(files)} generated modules up to date")
    return 0


def cmd_fmt(args) -> int:
    changed = []
    for spec in SPECS:
        path = args.content / spec.path
        if not path.is_file():
            continue
        try:
            doc = parse_csv(path.read_bytes().decode("utf-8"))
        except (CsvSyntaxError, UnicodeDecodeError) as e:
            print(f"{display_path(path)}: cannot format: {e}")
            return 1
        rows = [r.fields for r in doc.records if not (r.fields == [""] and r.raw == "")]
        width = len(rows[0]) if rows else 0
        rows = [r + [""] * (width - len(r)) for r in rows]
        text = render_csv(rows, None)
        if path.read_bytes() != text.encode("utf-8"):
            path.write_text(text, encoding="utf-8", newline="")
            changed.append(spec.path)
    print(f"fmt: {len(changed)} file(s) rewritten{': ' + ', '.join(changed) if changed else ''}")
    return 0


# ----------------------------------------------------------------------------- workbook
FILL = {"text_head": "FFE699", "text": "FFF2CC", "style_head": "FFE699", "style": "FFF2CC", "mech_head": "BFBFBF",
        "mech": "EDEDED", "key_head": "A6A6A6", "key": "E0E0E0", "notes_head": "FFFFFF", "info_head": "BDD7EE",
        "info": "DDEBF7"}
CONDITION_COLUMNS = ("trigger", "condition", "audience_condition")
LIST_SHEET = "lists"
README_SHEET = "README"
MAX_ROWS = 2000  # validations cover this many rows so new rows get dropdowns too
STAMP_HEADERS = ("stamp_file", "stamp_sha256")  # export stamp in the hidden lists sheet (stale-workbook check)
# Per-cell export stamp of every editable cell (text, choice, notes): "table|row label|column" -> short SHA-256 of
# the value at export. import-xlsx uses it to never revert a CSV cell edited after the export (finding M4).
CELL_STAMP_HEADERS = ("stamp_cell", "stamp_cell_sha256")
# Read-only helper columns: export-xlsx fills them in, import-xlsx ignores them.
INFO_COLUMNS = {"responses": ("info: question", "info: effects")}
EDITABLE_KINDS = ("text", "style", "notes")  # what import-xlsx --text-only takes from the workbook
# Tables without a unique key are matched by position; these columns must agree for --text-only.
POSITION_IDENTITY = {"names": ("kind",), "actions": ("action_id", "event", "audience")}
CHANGE_LINES = 200  # import-xlsx prints at most this many change lines


def _require_openpyxl():
    try:
        import openpyxl  # noqa: F401
    except ImportError:
        print("this command needs openpyxl: python3 -m pip install openpyxl")
        raise SystemExit(2)


def _raw_tables(content_dir: Path) -> tuple[dict[str, Table], Report]:
    report = Report()
    tables = load_tables(content_dir, report)
    return tables, report


def _sha256(path: Path) -> str | None:
    import hashlib

    return hashlib.sha256(path.read_bytes()).hexdigest() if path.is_file() else None


def _column_comment(spec: TableSpec, col: Column) -> str:
    if col.kind == "text":
        limit = f" Max {col.limit} characters." if col.limit else ""
        allowed = _column_placeholders(spec.key, col.name)
        ph = f" Placeholders: {_placeholder_list(allowed)}." if allowed else " No placeholders here."
        return (f"EDIT — player-facing text: {col.help}.{limit}{ph} Do not type numbers that the mechanics decide.")
    if col.kind == "style":
        return f"EDIT — presentation, pick from the dropdown: {col.help}."
    if col.kind == "notes":
        return f"NOTES — {col.help}. Free text, never shown in game."
    return f"MECHANICS — {col.help}. Changes gameplay; see content/SPEC.md before editing."


def _column_placeholders(table: str, col: str) -> tuple[str, ...]:
    """The placeholders a TEXT column accepts (SPEC §3); mirrors the Validator."""
    if table in ("scenarios", "responses"):
        return CITIZEN_PLACEHOLDERS
    if table == "consequences" and col == "notice_text":
        return ("name", "ruler", "kingdom", "occupation") + MECHANICS_PLACEHOLDERS
    if table == "actions" and col == "notice_text":
        return ("ruler", "kingdom", "name", "occupation")  # name/occupation only when the audience is speaker
    return ()


def _signed(v: str) -> str:
    v = v.strip()
    return v if v[:1] in "+-" else f"+{v}"


def _decision_summary(r: Row) -> str:
    """One line for the read-only 'info: effects' column (raw cells, so it works on an invalid library too)."""
    parts = [f"cost {r.get('cost').strip()}" if r.get("cost").strip() else "free"]
    if r.get("resource_effects").strip():
        parts.append(f"resources {r.get('resource_effects').strip()}")
    rep = [f"{c} {_signed(r.get(f'rep_{c}'))}" for c in CATEGORIES
           if r.get(f"rep_{c}").strip() not in ("", "0", "+0", "-0")]
    if rep:
        parts.append("rep " + ", ".join(rep))
    for col, label in (("speaker_loyalty", "speaker"), ("family_loyalty", "family"), ("audience_loyalty", "audience")):
        if r.get(col).strip() not in ("", "0", "+0", "-0"):
            parts.append(f"{label} {_signed(r.get(col))}")
    if r.get("consequences").strip():
        parts.append("then " + ", ".join(x.strip() for x in r.get("consequences").split(";") if x.strip()))
    return " | ".join(parts)


def _info_columns(tables: dict[str, Table]) -> dict[str, Callable[[Row], list[str]]]:
    questions = {r.get("scenario_id").strip(): r.get("question_text") for r in
                 (tables["scenarios"].rows if tables.get("scenarios") else [])}
    effects = {r.get("decision_id").strip(): _decision_summary(r) for r in
               (tables["decisions"].rows if tables.get("decisions") else [])}

    def responses(row: Row) -> list[str]:
        did = row.get("decision_id").strip()
        return [questions.get(row.get("scenario_id").strip(), ""),
                effects.get(did, f"(unknown decision '{did}')" if did else "")]

    return {"responses": responses}


def export_xlsx(content_dir: Path, path: Path) -> tuple[int, Report]:
    _require_openpyxl()
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils.exceptions import IllegalCharacterError
    from openpyxl.worksheet.datavalidation import DataValidation

    tables, report = _raw_tables(content_dir)
    if report.errors and len(tables) < len(SPECS):
        return 1, report
    wb = Workbook()
    readme = wb.active
    readme.title = README_SHEET
    occ_ids = _column_values(tables.get("occupations"), "occupation_id")
    lists = {
        "occupation_or_any": ["any", *occ_ids],
        "consequence_kinds": list(CONSEQUENCE_KINDS),
        "targets": [*SIMPLE_TARGETS, *(f"occupation:{o}" for o in occ_ids)],
        "decision_audiences": list(DECISION_AUDIENCES),
        "action_audiences": [*ACTION_AUDIENCES, *(f"occupation:{o}" for o in occ_ids)],
        "jobs": list(JOBS),
        "key_roles": list(KEY_ROLES),
        "name_kinds": list(NAME_KINDS),
        "bools": ["true", "false"],
        "thought_icons": list(THOUGHT_ICONS),
        "reaction_tones": list(REACTION_TONES),
    }
    live = {"decision_ids": ("decisions", "decision_id"), "scenario_ids": ("scenarios", "scenario_id"),
            "trait_ids": ("traits", "trait_id")}
    headers: dict[str, list[str]] = {}
    for spec in SPECS:
        t = tables.get(spec.key)
        headers[spec.key] = t.header if t is not None and t.header else [c.name for c in spec.columns if c.required]
    list_ws = wb.create_sheet(LIST_SHEET)
    list_ref: dict[str, str] = {}
    for i, (name, values) in enumerate(lists.items(), 1):
        letter = _col_letter(i)
        list_ws.cell(row=1, column=i, value=name)
        for j, v in enumerate(values, 2):
            list_ws.cell(row=j, column=i, value=v).number_format = "@"
        list_ref[name] = f"{LIST_SHEET}!${letter}$2:${letter}${max(2, len(values) + 1)}"
    # export stamp: which CSV bytes this workbook was made from (import-xlsx refuses stale mechanics)
    stamp_col = len(lists) + 1
    list_ws.cell(row=1, column=stamp_col, value=STAMP_HEADERS[0])
    list_ws.cell(row=1, column=stamp_col + 1, value=STAMP_HEADERS[1])
    for j, spec in enumerate(SPECS, 2):
        list_ws.cell(row=j, column=stamp_col, value=spec.path)
        list_ws.cell(row=j, column=stamp_col + 1, value=_sha256(content_dir / spec.path) or "")
    list_ws.cell(row=1, column=stamp_col + 2, value=CELL_STAMP_HEADERS[0])
    list_ws.cell(row=1, column=stamp_col + 3, value=CELL_STAMP_HEADERS[1])
    for j, (cell_key, digest) in enumerate(_cell_stamps(tables).items(), 2):
        list_ws.cell(row=j, column=stamp_col + 2, value=cell_key).number_format = "@"
        list_ws.cell(row=j, column=stamp_col + 3, value=digest).number_format = "@"
    for name, (sheet, col) in live.items():
        if col in headers[sheet]:
            letter = _col_letter(headers[sheet].index(col) + 1)
            list_ref[name] = f"{sheet}!${letter}$2:${letter}${MAX_ROWS}"
    list_ws.sheet_state = "hidden"

    bold = Font(bold=True)
    top_wrap = Alignment(wrap_text=True, vertical="top")
    info_of = _info_columns(tables)
    illegal = 0
    for spec in SPECS:
        ws = wb.create_sheet(spec.key)
        t = tables.get(spec.key)
        header = headers[spec.key]
        rows = t.rows if t is not None else []
        info = INFO_COLUMNS.get(spec.key, ()) if t is not None else ()
        for ci, name in enumerate(header, 1):
            col = spec.col(name)
            kind = col.kind if col else "mech"
            cell = ws.cell(row=1, column=ci, value=name)
            cell.font = bold
            cell.fill = PatternFill("solid", fgColor=FILL[f"{kind}_head"])
            cell.number_format = "@"
            if col is not None:
                comment = Comment(_column_comment(spec, col), "content.py")
                comment.width, comment.height = 320, 120
                cell.comment = comment
        for k, name in enumerate(info, len(header) + 1):
            cell = ws.cell(row=1, column=k, value=name)
            cell.font = bold
            cell.fill = PatternFill("solid", fgColor=FILL["info_head"])
            comment = Comment("READ ONLY — filled in by export-xlsx to help you edit; import-xlsx ignores this "
                              "column.", "content.py")
            comment.width, comment.height = 320, 60
            cell.comment = comment
        for ri, row in enumerate(rows, 2):
            for ci, name in enumerate(header, 1):
                value = row.values.get(name, "")
                col = spec.col(name)
                kind = col.kind if col else "mech"
                try:
                    cell = ws.cell(row=ri, column=ci, value=value if value != "" else None)
                except IllegalCharacterError:
                    problem = _invisible_problem(value, "mech") or "a character a workbook cannot hold"
                    report.error(t, row.num, name, f"cannot export: {problem}")
                    illegal += 1
                    continue
                if value.startswith("="):
                    cell.data_type = "s"  # never a formula
                cell.number_format = "@"
                if kind in FILL:
                    cell.fill = PatternFill("solid", fgColor=FILL[kind])
                cell.alignment = top_wrap
            if info:
                for k, value in enumerate(info_of[spec.key](row), len(header) + 1):
                    try:
                        cell = ws.cell(row=ri, column=k, value=value or None)
                    except IllegalCharacterError:
                        continue  # the source cell is reported above
                    cell.data_type = "s"
                    cell.number_format = "@"
                    cell.fill = PatternFill("solid", fgColor=FILL["info"])
                    cell.alignment = top_wrap
        # widths
        for ci, name in enumerate(header, 1):
            col = spec.col(name)
            longest = max([len(name)] + [len(r.values.get(name, "")) for r in rows])
            if col is not None and (col.kind in ("text", "notes") or name in CONDITION_COLUMNS):
                width = min(max(longest * 0.9, 18), 60)
            else:
                width = min(max(longest + 2, 8), 40)
            ws.column_dimensions[_col_letter(ci)].width = round(width, 1)
        for k in range(len(header) + 1, len(header) + len(info) + 1):
            ws.column_dimensions[_col_letter(k)].width = 50
        ws.freeze_panes = "B2"
        if rows:
            ws.auto_filter.ref = f"A1:{_col_letter(len(header) + len(info))}{len(rows) + 1}"
        # validations
        for ci, name in enumerate(header, 1):
            col = spec.col(name)
            if col is None:
                continue
            rng = f"{_col_letter(ci)}2:{_col_letter(ci)}{MAX_ROWS}"
            if col.choices and col.choices in list_ref:
                strict = col.choices in ("consequence_kinds", "jobs", "key_roles", "name_kinds", "bools",
                                         "decision_audiences", "thought_icons", "reaction_tones")
                dv = DataValidation(type="list", formula1=list_ref[col.choices], allow_blank=True,
                                    errorStyle="stop" if strict else "warning",
                                    error=f"Pick a value from the list ({name}).", errorTitle=name,
                                    showErrorMessage=True)
                dv.add(rng)
                ws.add_data_validation(dv)
            if col.kind == "text" and col.limit:
                dv = DataValidation(type="textLength", operator="lessThanOrEqual", formula1=str(col.limit),
                                    allow_blank=True, errorStyle="stop", showErrorMessage=True,
                                    errorTitle="Too long", error=f"{name}: at most {col.limit} characters.",
                                    promptTitle="EDIT", prompt=f"Max {col.limit} characters", showInputMessage=True)
                dv.add(rng)
                ws.add_data_validation(dv)
    if illegal:
        return 1, report
    _write_readme(readme, headers)
    wb.active = 0
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return 0, report


def _cell_hash(value: str) -> str:
    import hashlib

    return hashlib.sha256(value.encode("utf-8")).hexdigest()[:16]


def _row_labels(spec: TableSpec, rows: list[Row]) -> list[str | None]:
    """The label _match_rows gives each CSV row: its key (first occurrence only), or 'row N' without a key."""
    key = KEY_COLUMN.get(spec.key)
    labels: list[str | None] = []
    seen: set[str] = set()
    for i, r in enumerate(rows):
        if not key:
            labels.append(f"row {i + 2}")
            continue
        k = _row_key(spec, r)
        labels.append(k if k is not None and k not in seen else None)
        if k is not None:
            seen.add(k)
    return labels


def _cell_stamps(tables: dict[str, Table]) -> dict[str, str]:
    """{'table|row label|column': hash of the value} for every editable cell, absent optional columns included."""
    out: dict[str, str] = {}
    for spec in SPECS:
        t = tables.get(spec.key)
        if t is None:
            continue
        cols = [c.name for c in spec.columns if c.kind in EDITABLE_KINDS]
        for label, r in zip(_row_labels(spec, t.rows), t.rows):
            if label is None:
                continue
            for col in cols:
                out[f"{spec.key}|{label}|{col}"] = _cell_hash(r.values.get(col, ""))
    return out


def _column_values(t: Table | None, col: str) -> list[str]:
    if t is None:
        return []
    out = []
    for r in t.rows:
        v = r.get(col).strip()
        if v and v not in out:
            out.append(v)
    return out


def _write_readme(ws, headers: dict[str, list[str]]) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill

    lines = [
        ("AI Kingdoms — content workbook (generated by tools/content/content.py export-xlsx)", "title"),
        ("The CSV files in games/game-001/content are the master copy. This workbook is a working copy for editing.", ""),
        ("", ""),
        ("HOW TO EDIT", "head"),
        ("1. Change text only in YELLOW columns (marked EDIT in the header comment). Stay within the character limit.", ""),
        ("   Text may not start with = + - or @ (a spreadsheet would run it as a formula). Never type formulas.", ""),
        ("2. GREY columns are mechanics (ids, numbers, conditions). Changing them changes gameplay — see SPEC.md first.", ""),
        ("3. WHITE columns are notes for designers; they never reach the game. BLUE columns are read-only help.", ""),
        ("4. Never rename sheets or header cells, never add or remove columns or sheets. Add rows at the bottom.", ""),
        ("5. Sorting and filtering are fine: the import puts rows back in their original order.", ""),
        ("6. Ids link sheets together: responses.decision_id must exist in decisions, and so on. Use the dropdowns.", ""),
        ("7. Placeholders the game fills in: {name} {ruler} {kingdom} {occupation}; {family} {friend} {rival} only "
         "when the scenario's trigger requires that relation. Hover a header to see which ones a column accepts.", ""),
        ("8. Never type numbers the mechanics decide (odds, durations, amounts). In consequence notices write "
         "{chance}, {magnitude}, {magnitude_pct} or {duration}.", ""),
        ("9. Keep cells as Text. Write true/false in lower case and lists like food:-20;gold:5.", ""),
        ("", ""),
        ("HAND IT BACK", "head"),
        ("Save as .xlsx (same name is fine) and give the file back, or run:", ""),
        ("    python3 tools/content/content.py import-xlsx <file.xlsx> --text-only --build", "code"),
        ("--text-only takes only the yellow and white columns, so mechanics can never change by accident (safest).", ""),
        ("Without it, grey columns are imported too, but only if nobody changed the CSV files since this export.", ""),
        ("Text a developer changed in the CSV after this export is never undone: cells you did not edit keep the newer "
         "text; a cell you both changed is refused (export a fresh workbook).", ""),
        ("The import checks everything first and changes nothing if it finds an error; errors name the sheet and cell.", ""),
        ("It always lists every change it made, split into TEXT, MECHANICS and NOTES.", ""),
        ("", ""),
        ("COLUMNS", "head"),
    ]
    head_fill = PatternFill("solid", fgColor="D9E1F2")
    r = 1
    for text, style in lines:
        c = ws.cell(row=r, column=1, value=text or None)
        if style == "title":
            c.font = Font(bold=True, size=14)
        elif style == "head":
            c.font = Font(bold=True)
            c.fill = head_fill
        elif style == "code":
            c.font = Font(name="Consolas")
        r += 1
    for ci, title in enumerate(("sheet", "column", "kind", "limit", "meaning"), 1):
        c = ws.cell(row=r, column=ci, value=title)
        c.font = Font(bold=True)
        c.fill = head_fill
    r += 1
    kind_label = {"text": "EDIT (text)", "style": "EDIT (choice)", "mech": "mechanics", "key": "id (mechanics)",
                  "notes": "notes"}
    for spec in SPECS:
        for name in headers[spec.key]:
            col = spec.col(name)
            if col is None:
                continue
            values = (spec.key, col.name, kind_label[col.kind], col.limit, col.help)
            for ci, v in enumerate(values, 1):
                cell = ws.cell(row=r, column=ci, value=v)
                if col.kind in ("text", "style"):
                    cell.fill = PatternFill("solid", fgColor=FILL["text"])
            r += 1
    for letter, width in zip("ABCDE", (16, 20, 16, 7, 90)):
        ws.column_dimensions[letter].width = width
    for row in ws.iter_rows(min_row=1, max_row=len(lines)):
        row[0].alignment = Alignment(wrap_text=False)


def _cell_text(value: Any) -> str:
    import datetime as _dt

    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else repr(value)
    if isinstance(value, (_dt.datetime, _dt.date, _dt.time)):
        raise ValueError("looks like a date (format the cell as Text and type the value again)")
    return str(value)


def _read_stamps(ws, headers: tuple[str, str] = STAMP_HEADERS) -> dict[str, str] | None:
    """{key: digest} from two columns of the hidden lists sheet; None when the columns are absent."""
    header = [str(c.value).strip() if c.value is not None else "" for c in ws[1]]
    if headers[0] not in header or headers[1] not in header:
        return None
    fi, si = header.index(headers[0]), header.index(headers[1])
    stamps: dict[str, str] = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if fi < len(row) and row[fi]:
            stamps[str(row[fi]).strip()] = str(row[si]).strip() if si < len(row) and row[si] else ""
    return stamps


@dataclass
class Stamps:
    files: dict[str, str] | None = None  # {csv path: sha256 of the file at export}
    cells: dict[str, str] | None = None  # {'table|row label|column': hash of the editable cell at export}


def read_workbook(path: Path, report: Report) -> tuple[dict[str, Table], Stamps]:
    """The table sheets (info columns dropped) and the export stamps (files and editable cells; None when absent)."""
    _require_openpyxl()
    from openpyxl import load_workbook

    source = path.name
    try:
        # data_only=False: a formula is read as a formula and refused below. data_only=True would blank a formula
        # without a cached value and silently take the cached value of any other (review finding L1).
        wb = load_workbook(path, data_only=False)
    except Exception as e:  # noqa: BLE001 - any unreadable file
        report.add("error", None, None, None, f"cannot open workbook: {e}", source=source)
        return {}, Stamps()
    tables: dict[str, Table] = {}
    stamps = Stamps()
    for ws in wb.worksheets:
        if ws.title == LIST_SHEET:
            stamps = Stamps(_read_stamps(ws), _read_stamps(ws, CELL_STAMP_HEADERS))
            continue
        if ws.title == README_SHEET:
            continue
        spec = SPEC_BY_KEY.get(ws.title)
        if spec is None:
            known = ", ".join(s.key for s in SPECS)
            report.add("error", None, None, None, f"unknown sheet '{ws.title}'{_suggest(ws.title, SPEC_BY_KEY)} "
                                                  f"(known sheets: {known}, {README_SHEET}, {LIST_SHEET})", source=source)
            continue
        records: list[tuple[int, list[str]]] = []
        for r_idx, row in enumerate(ws.iter_rows(), 1):
            cells = []
            for c_idx, cell in enumerate(row, 1):
                where = f"{ws.title}!{_col_letter(c_idx)}{r_idx}"
                if cell.data_type == "f":
                    formula = getattr(cell.value, "text", None) or str(cell.value)
                    report.add("error", None, None, None, f"{where}: this cell is a formula ({_short(formula, 40)}); "
                                                          "formulas are never imported: type the value itself, as "
                                                          "text", source=source)
                    cells.append("")
                    continue
                try:
                    cells.append(_cell_text(cell.value))
                except ValueError as e:
                    report.add("error", None, None, None, f"{where}: {e}", source=source)
                    cells.append("")
            while cells and cells[-1] == "" and records:
                cells.pop()
            records.append((r_idx, cells))
        if not records:
            report.add("error", None, None, None, f"sheet '{ws.title}' is empty (it needs its header row)", source=source)
            continue
        info = INFO_COLUMNS.get(ws.title, ())
        drop = {i for i, h in enumerate(records[0][1]) if h.strip() in info}
        if drop:
            records = [(n, [v for i, v in enumerate(cells) if i not in drop]) for n, cells in records]
        table = table_from_records(spec, source, records, report, sheet=ws.title)
        tables[spec.key] = table
    return tables, stamps


# --- import: order, diff, merge
@dataclass
class Change:
    kind: str  # "TEXT" | "MECHANICS" | "NOTES"
    sheet: str
    label: str  # row key, or "row N" for tables without a key
    column: str | None
    old: str | None  # None: the row is new
    new: str | None  # None: the row was removed
    row: int | None = None  # workbook row, for error locations
    applied: bool = True

    def describe(self) -> str:
        where = f"{self.sheet} › {self.label}" + (f" [{self.column}]" if self.column else "")
        if self.old is None:
            return f"{where}: new row"
        if self.new is None:
            return f"{where}: row removed"
        return f"{where}: {_short(self.old)} → {_short(self.new)}"


def _short(s: str, limit: int = 70) -> str:
    s = s if len(s) <= limit else s[: limit - 1] + "…"
    return '""' if s == "" else repr(s)


def _change_kind(spec: TableSpec, col: str) -> str:
    c = spec.col(col)
    if c is not None and c.kind in ("text", "style"):
        return "TEXT"
    if c is not None and c.kind == "notes":
        return "NOTES"
    return "MECHANICS"


def _row_key(spec: TableSpec, row: Row) -> str | None:
    key = KEY_COLUMN.get(spec.key)
    return row.get(key).strip() or None if key else None


def _restore_order(spec: TableSpec, cur: Table, rows: list[Row]) -> list[Row]:
    """Undo a sort in the workbook: keyed rows go back to their CSV order (new ids at the end, in workbook order).
    A table without a key is restored only when its rows are a pure permutation of the CSV rows."""
    key = KEY_COLUMN.get(spec.key)
    if key:
        position = {}
        for i, r in enumerate(cur.rows):
            position.setdefault(r.get(key).strip(), i)
        placed: dict[str, Row] = {}
        rest: list[Row] = []
        for r in rows:
            k = r.get(key).strip()
            if k in position and k not in placed:
                placed[k] = r
            else:
                rest.append(r)
        return sorted(placed.values(), key=lambda r: position[r.get(key).strip()]) + rest
    cols = cur.header

    def sig(r: Row) -> tuple[str, ...]:
        return tuple(r.values.get(c, "") for c in cols)

    from collections import Counter

    if Counter(sig(r) for r in rows) != Counter(sig(r) for r in cur.rows):
        return rows
    pool: dict[tuple[str, ...], deque[Row]] = {}
    for r in rows:
        pool.setdefault(sig(r), deque()).append(r)
    return [pool[sig(r)].popleft() for r in cur.rows]


def _match_rows(spec: TableSpec, cur_rows: list[Row], wb_rows: list[Row]):
    """(pairs [(cur, wb, label)], added [(wb, label)], removed [(cur, label)]): by key, else by position."""
    pairs, added, removed = [], [], []
    if KEY_COLUMN.get(spec.key):
        by_key: dict[str, Row] = {}
        for r in cur_rows:
            k = _row_key(spec, r)
            if k is not None:
                by_key.setdefault(k, r)
        matched: set[str] = set()
        for w in wb_rows:
            k = _row_key(spec, w)
            if k is not None and k in by_key and k not in matched:
                matched.add(k)
                pairs.append((by_key[k], w, k))
            else:
                added.append((w, k or f"row {w.num}"))
        removed = [(r, k) for k, r in by_key.items() if k not in matched]
        return pairs, added, removed
    for i in range(max(len(cur_rows), len(wb_rows))):
        if i < len(cur_rows) and i < len(wb_rows):
            pairs.append((cur_rows[i], wb_rows[i], f"row {i + 2}"))
        elif i < len(wb_rows):
            added.append((wb_rows[i], f"row {i + 2}"))
        else:
            removed.append((cur_rows[i], f"row {i + 2}"))
    return pairs, added, removed


def _diff_table(spec: TableSpec, cur: Table, wb: Table) -> list[Change]:
    changes: list[Change] = []
    if set(cur.header) != set(wb.header):
        changes.append(Change("MECHANICS", spec.key, "(header)", None, ", ".join(cur.header), ", ".join(wb.header)))
    cols = list(dict.fromkeys([*cur.header, *wb.header]))
    pairs, added, removed = _match_rows(spec, cur.rows, wb.rows)
    for c_row, w_row, label in pairs:
        for col in cols:
            a, b = c_row.values.get(col, ""), w_row.values.get(col, "")
            if a != b:
                changes.append(Change(_change_kind(spec, col), spec.key, label, col, a, b, w_row.num))
    changes += [Change("MECHANICS", spec.key, label, None, None, "", w.num) for w, label in added]
    changes += [Change("MECHANICS", spec.key, label, None, "", None) for _, label in removed]
    return changes


def _guard_editable(spec: TableSpec, cur: Table, wbt: Table, stamps: Stamps, file_unchanged: bool,
                    report: Report) -> list[Change]:
    """Review finding M4: never revert a text, choice or notes cell of the CSV that changed after the export.
    For each editable cell where the workbook and the current CSV differ:
      - the CSV cell is as exported (file stamp unchanged, or cell stamp = CSV value): take the workbook value;
      - the workbook cell is as exported (cell stamp = workbook value): only the CSV changed, so keep the CSV
        value (the cell is reset in `wbt` and returned, so the import can list it);
      - both changed, or no stamp can tell: refuse it, naming the cell (export a fresh workbook)."""
    kept: list[Change] = []
    if file_unchanged:
        return kept
    pairs, _, _ = _match_rows(spec, cur.rows, wbt.rows)
    cols = [c.name for c in spec.columns if c.kind in EDITABLE_KINDS and c.name in wbt.header]
    for c_row, w_row, label in pairs:
        for col in cols:
            now, theirs = c_row.values.get(col, ""), w_row.values.get(col, "")
            if now == theirs:
                continue
            stamp = stamps.cells.get(f"{spec.key}|{label}|{col}") if stamps.cells is not None else None
            if stamp is not None and stamp == _cell_hash(now):
                continue  # only the workbook changed it
            if stamp is not None and stamp == _cell_hash(theirs):
                w_row.values[col] = now
                kept.append(Change(_change_kind(spec, col), spec.key, label, col, theirs, now, w_row.num, False))
                continue
            if stamp is not None:
                why = f"it was also changed in {spec.path} after this workbook was exported"
            elif stamps.cells is None:
                why = (f"{spec.path} changed after the export, and this workbook has no per-cell export stamp (made "
                       "by an older tool), so the import cannot tell whether this would undo a newer edit")
            else:
                why = f"this row was not in {spec.path} when the workbook was exported"
            report.add("error", wbt, w_row.num, col,
                       f"{spec.key} › {label} [{col}]: refused: {why} (CSV now {_short(now)}, workbook "
                       f"{_short(theirs)}). Export a fresh workbook and redo this edit there")
    return kept


def _text_only_rows(spec: TableSpec, cur: Table, wb: Table) -> tuple[list[Row] | None, str | None]:
    """The current CSV rows with only the TEXT, choice and notes cells taken from the workbook."""
    pairs, added, removed = _match_rows(spec, cur.rows, wb.rows)
    if not KEY_COLUMN.get(spec.key):
        ident = POSITION_IDENTITY.get(spec.key, ())
        moved = any(c.values.get(k, "").strip() != w.values.get(k, "").strip() for c, w, _ in pairs for k in ident)
        if added or removed or moved:
            return None, (f"rows were added, removed or moved in sheet '{spec.key}', so --text-only cannot match "
                          "them to the CSV: undo that, or import without --text-only")
    take = {c.name for c in spec.columns if c.kind in EDITABLE_KINDS} & set(wb.header)
    by_cur = {id(c): w for c, w, _ in pairs}
    out = []
    for c in cur.rows:
        w = by_cur.get(id(c))
        if w is None:
            out.append(c)
            continue
        values = {col: (w.values.get(col, "") if col in take else c.values.get(col, "")) for col in cur.header}
        out.append(Row(w.num, values))
    return out, None


@dataclass
class ImportResult:
    code: int
    report: Report
    written: list[str]
    model: Model | None
    changes: list[Change]
    kept: list[Change] = field(default_factory=list)  # editable cells kept as in the CSV (newer than the export)


def import_xlsx(path: Path, content_dir: Path, text_only: bool = False) -> ImportResult:
    report = Report()
    sheets, stamps = read_workbook(path, report)
    if report.errors:
        return ImportResult(1, report, [], None, [])
    changes: list[Change] = []
    kept: list[Change] = []
    merged: dict[str, Table] = {}
    for key, wbt in sheets.items():
        spec = SPEC_BY_KEY[key]
        cur = load_csv_table(spec, content_dir, Report())
        if cur is None or not cur.usable or not wbt.usable:
            if text_only:
                report.add("error", wbt, None, None, f"--text-only needs a readable {spec.path} to merge into")
                continue
            merged[key] = wbt  # nothing to compare with: validation below has the last word
            continue
        wbt.rows = _restore_order(spec, cur, wbt.rows)
        file_unchanged = stamps.files is not None and stamps.files.get(spec.path) == _sha256(content_dir / spec.path)
        kept += _guard_editable(spec, cur, wbt, stamps, file_unchanged, report)
        diff = _diff_table(spec, cur, wbt)
        changes += diff
        if text_only:
            rows, problem = _text_only_rows(spec, cur, wbt)
            if rows is None:
                report.add("error", wbt, None, None, problem or "cannot merge")
                continue
            for c in diff:
                c.applied = c.kind != "MECHANICS"
            merged[key] = Table(spec, wbt.source, list(cur.header), rows, sheet=wbt.sheet)
            continue
        stale = None
        if stamps.files is None:
            stale = "this workbook has no export stamp (made by an older tool or by hand)"
        elif not file_unchanged:
            stale = f"{spec.path} changed after this workbook was exported"
        mechanics = [c for c in diff if c.kind == "MECHANICS"]
        if stale and mechanics:
            for c in mechanics:
                c.applied = False
                report.add("error", wbt, c.row, c.column,
                           f"{c.describe()} — refused: {stale}, so this would overwrite newer mechanics. Export a "
                           "fresh workbook and redo the change, or import with --text-only to take only the text")
            continue
        merged[key] = wbt
    if report.errors:
        return ImportResult(1, report, [], None, changes, kept)
    tables = load_tables(content_dir, report, overrides=merged)
    model = validate_tables(tables, report)
    if report.errors:
        return ImportResult(1, report, [], model, changes, kept)
    written = []
    for key, table in merged.items():
        spec = SPEC_BY_KEY[key]
        csv_path = content_dir / spec.path
        old = None
        old_header: list[str] | None = None
        if csv_path.is_file():
            try:
                old = parse_csv(csv_path.read_bytes().decode("utf-8"))
            except (UnicodeDecodeError, CsvSyntaxError):
                old = None  # an unreadable file is replaced as a whole
            if old is not None and old.records:
                old_header = [h.strip() for h in old.records[0].fields]
        if old_header is not None and set(old_header) == set(table.header):
            order = old_header
        else:
            order = [c.name for c in spec.columns if c.name in table.header]
        rows = [order] + [[r.values.get(c, "") for c in order] for r in table.rows]
        text = render_csv(rows, old)
        if not csv_path.is_file() or csv_path.read_bytes() != text.encode("utf-8"):
            csv_path.parent.mkdir(parents=True, exist_ok=True)
            csv_path.write_text(text, encoding="utf-8", newline="")
            written.append(spec.path)
    return ImportResult(0, report, written, model, changes, kept)


def print_changes(changes: list[Change], out=None, kept: list[Change] | None = None) -> None:
    out = out or sys.stdout
    applied = [c for c in changes if c.applied]
    ignored = [c for c in changes if not c.applied]
    counts = {k: sum(1 for c in applied if c.kind == k) for k in ("TEXT", "MECHANICS", "NOTES")}
    print(f"import-xlsx: {len(applied)} change(s) from the workbook: TEXT {counts['TEXT']}, "
          f"MECHANICS {counts['MECHANICS']}, NOTES {counts['NOTES']}", file=out)
    for c in applied[:CHANGE_LINES]:
        print(f"  {c.kind:<9} {c.describe()}", file=out)
    if len(applied) > CHANGE_LINES:
        print(f"  … and {len(applied) - CHANGE_LINES} more", file=out)
    if ignored:
        print(f"import-xlsx: {len(ignored)} mechanics change(s) in the workbook were NOT imported:", file=out)
        for c in ignored[:CHANGE_LINES]:
            print(f"  ignored   {c.describe()}", file=out)
    if kept:
        print(f"import-xlsx: {len(kept)} cell(s) kept as in the CSV (changed there after your export; you had not "
              "edited them in the workbook):", file=out)
        for c in kept[:CHANGE_LINES]:
            print(f"  kept      {c.describe()}", file=out)


def cmd_export(args) -> int:
    path = Path(args.path) if args.path else args.content / XLSX_NAME
    code, report = export_xlsx(args.content, path)
    if code:
        print_diags(report.errors)
        print("export-xlsx: cannot export (fix the files above)")
        return code
    if report.errors:
        print(f"export-xlsx: note: the library has {len(report.errors)} validation error(s); run validate")
    print(f"export-xlsx: wrote {display_path(path)}")
    return 0


def cmd_import(args) -> int:
    res = import_xlsx(Path(args.path), args.content, text_only=args.text_only)
    if res.code:
        print_diags(res.report.errors)
        print(f"import-xlsx: {len(res.report.errors)} error(s); no CSV was changed")
        return 1
    print_changes(res.changes, kept=res.kept)
    print(f"import-xlsx: {len(res.written)} CSV file(s) updated{': ' + ', '.join(res.written) if res.written else ''}")
    print(summary_line(res.model, 0))
    if args.build:
        return cmd_build(args)
    return 0


# ----------------------------------------------------------------------------- CLI
def build_parser() -> argparse.ArgumentParser:
    common = argparse.ArgumentParser(add_help=False)
    common.add_argument("--content", type=Path, default=CONTENT_DIR, help="content directory (default: %(default)s)")
    with_out = argparse.ArgumentParser(add_help=False)
    with_out.add_argument("--out", type=Path, default=OUT_DIR, help="generated Luau directory (default: %(default)s)")
    p = argparse.ArgumentParser(prog="content.py", description=__doc__.split("\n\n")[0])
    sub = p.add_subparsers(dest="command", required=True)
    sub.add_parser("validate", parents=[common], help="check every CSV").set_defaults(func=cmd_validate)
    sub.add_parser("lint", parents=[common], help="design warnings (never fails)").set_defaults(func=cmd_lint)
    sub.add_parser("build", parents=[common, with_out], help="validate and write Luau").set_defaults(func=cmd_build)
    sub.add_parser("check", parents=[common, with_out], help="fail if Luau is stale").set_defaults(func=cmd_check)
    sub.add_parser("fmt", parents=[common], help="rewrite CSVs canonically").set_defaults(func=cmd_fmt)
    ex = sub.add_parser("export-xlsx", parents=[common], help="write the editing workbook")
    ex.add_argument("path", nargs="?", help=f"workbook path (default: <content>/{XLSX_NAME})")
    ex.set_defaults(func=cmd_export)
    im = sub.add_parser("import-xlsx", parents=[common, with_out], help="read the workbook back into the CSVs")
    im.add_argument("path", help="workbook path")
    im.add_argument("--build", action="store_true", help="also regenerate the Luau modules")
    im.add_argument("--text-only", action="store_true",
                    help="take only TEXT, choice and notes columns from the workbook (mechanics stay as in the CSV)")
    im.set_defaults(func=cmd_import)
    return p


def main(argv: list[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    return args.func(args)


if __name__ == "__main__":
    sys.exit(main())
